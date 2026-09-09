"""Build the preliminary user manual for implemented UIABO features only.

Uses the prior editable installation section and actual Android captures.
Produces Word documents only. Preserve manual edits before rebuilding.
"""
from pathlib import Path
from shutil import copy2
import re

from docx import Document
from docxcompose.composer import Composer
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
WORKSPACE = ROOT.parents[2]
REPO = WORKSPACE if (WORKSPACE/'backend').is_dir() else WORKSPACE/'uiabo'
PARTS = ROOT/'PUMParts_Implemented'
ASSETS = ROOT/'assets_implemented'
MASTER = ROOT/'FYP-26-S3-30_PrelimUserManual_DRAFT.docx'
DATE = '8 September 2026'
CODE = 'FYP-26-S3-30'


def p(d, value, style=None):
    return d.add_paragraph(value, style)


def h(d, value, level=2):
    return d.add_heading(value, level)


def steps(d, values):
    for i, value in enumerate(values, 1):
        p(d, f'{i}. {value}')


def bullets(d, values):
    for value in values:
        p(d, value, 'List Bullet')


def table(d, headers, rows):
    t = d.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for cell, value in zip(t.rows[0].cells, headers): cell.text = value
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in rows:
        for cell, value in zip(t.add_row().cells, row): cell.text = value
    return t


def field(p, instruction):
    r = p.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    r._r.extend([begin, instr, separate, end])


def new_doc():
    return Document()


def figure(d, filename, caption):
    p(d, caption, 'Caption').paragraph_format.keep_with_next = True
    para = d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(ASSETS/filename), height=Cm(15))
    para.paragraph_format.space_after = Pt(8)


def cover():
    d = new_doc()
    for value in ['School of Computing and Information Technology', 'CSIT321 — Project']:
        p(d, value).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for value in ['uiabo.', 'Preliminary User Manual']:
        p(d, value, 'Title').alignment = WD_ALIGN_PARAGRAPH.CENTER
    for value in ['Project Topic: CSIT-26-S3-30', 'AI-Based Misinformation Detection for Short-Form Content',
                  'Group Number: FYP-26-S3-30', 'Supervisor: Mr Liaw Chun Huei',
                  'Assessor: Mr Terrence Chew', f'Version 0.2 — {DATE}']:
        p(d, value).alignment = WD_ALIGN_PARAGRAPH.CENTER
    previous = Document(ROOT/'PUMParts/00_Cover_Document_Control_and_Contents.docx')
    roster = [[cell.text for cell in row.cells] for row in previous.tables[0].rows]
    table(d, roster[0], roster[1:])
    p(d, 'A practical guide to installing and using the implemented Android text-checking prototype.').alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()
    h(d, 'Document Control', 1)
    table(d, ['Field', 'Record'], [
        ('Title', 'uiabo — Preliminary User Manual'), ('Document name', CODE + '_PrelimUserManual'),
        ('Version', '0.2 — working draft for team review'), ('Revision date', DATE),
        ('Scope', 'Implemented account, text-checking, result and allowance functions'),
        ('Environment', 'Windows development computer and Android Studio emulator'),
        ('Evidence basis', 'Current application code and existing Android/Firestore integration records'),
    ])
    p(d, 'Record of Revision', 'Subtitle')
    table(d, ['Date', 'Version', 'Description', 'Sections'], [
        ('7 September 2026', '0.1', 'Initial manual assembled from project documents.', 'All'),
        (DATE, '0.2', 'Aligned with the sample manual and limited to implemented functions; retained actual app captures.', 'All'),
    ])
    p(d, 'Prepared for team review. The screenshots are existing Android captures dated 7 September 2026; names, times and remaining checks are example account values.')
    d.add_page_break()
    p(d, 'Table of Contents', 'Title')
    field(d.add_paragraph(), ' TOC \\o "1-3" \\h \\z \\u ')
    return d


def introduction():
    d = new_doc(); h(d, '1. Introduction', 1)
    p(d, 'This Preliminary User Manual explains how to run UIABO and complete a first text check. UIABO helps users examine short English messages or captions by identifying a checkable claim, finding published evidence and presenting a readable result with citations and uncertainty.')
    h(d, '1.1 What This Manual Covers')
    bullets(d, ['Preparing the local Windows backend and Android emulator application.',
        'Creating a free account, verifying an email address, signing in and resetting a password.',
        'Checking text, reading the result, opening sources and sharing a result summary.',
        'Reopening saved results, checking the remaining allowance, changing a profile name and signing out.',
        'Using Help and responding to the current application’s error messages.'])
    h(d, '1.2 What We Assume About the Readers')
    p(d, 'The person setting up the prototype should be comfortable opening PowerShell, obtaining the project files and entering configuration values supplied by the team. Chapter 2 is for that person.')
    p(d, 'A person using an already configured emulator can start with Chapters 3 and 4. They need an email address they can access, an internet connection and basic familiarity with tapping buttons and entering text. UIABO’s intended audience includes Singapore citizens aged 65 and above, so the operating instructions use short steps and the labels shown in the app.')
    h(d, '1.3 Scope and Purpose')
    p(d, 'This edition covers the working Android text-checking prototype. It guides a user from account access to a saved assessment and its source links. The instructions describe actions available in the current application.')
    p(d, 'An automated assessment can be wrong. Read the explanation, source passages and uncertainty before deciding what to believe or share. A risk indicator is not a percentage probability that a claim is false. Do not submit passwords or private information in the text-checking field.')
    return d


def installation():
    # Reuse the earlier setup instructions, already checked against this checkout.
    d = Document(ROOT/'PUMParts/02_Initial_Installation_Instructions.docx')
    for para in list(d.paragraphs):
        value = para.text
        if value.startswith('Source:'):
            para._element.getparent().remove(para._element); continue
        if value.startswith('These instructions prepare'):
            para.text = 'These instructions prepare the existing Windows and Android-emulator demonstration. If the team has already configured the application, continue to Chapter 3. The setup uses the team’s configured Firebase project and privately supplied backend credentials.'
        elif value.startswith('The backend verifies the Firebase identity token'):
            para.text = 'The backend controls account access, allowance and saved results. Use the team’s configured Firebase project and authorised service account. Do not change database access rules to public access when troubleshooting.'
        elif value.startswith('For the standard Android Studio emulator'):
            para.text = 'For the standard Android Studio emulator, 10.0.2.2 reaches the host computer. Using localhost inside the emulator refers to the emulator itself. The addresses in this section are for the local development setup.'
        elif value.startswith('The API documentation is available locally'):
            para.text = 'The API documentation is available locally at http://127.0.0.1:8000/docs. The signed-in app supplies the identity token for account and analysis requests. Text analysis and saved-result access require a verified, active account.'
        elif value == '2. Initial Installation Instructions':
            para.text = '2. The Initial Installation Instructions'
        elif value == '2.3 Frontend / Android setup':
            para.text = '2.3 Frontend / Android Setup'
    return d


def features():
    d = new_doc(); h(d, '3. Key Features of uiabo', 1)
    table(d, ['Implemented feature', 'What the user can do'], [
        ('Account access', 'Create a free account, verify an email address, sign in, request a password-reset email and sign out.'),
        ('Text checking', 'Paste up to 5,000 characters of English text and submit it for claim and evidence analysis.'),
        ('Result explanation', 'Read the concern label, risk indicator when available, uncertainty, identified claim and recommended action.'),
        ('Source evidence', 'Read the retrieved passages and open source links in the browser.'),
        ('Saved results', 'View, refresh and reopen the signed-in account’s result history.'),
        ('Result sharing', 'Open the Android share sheet with an assessment summary and its citation URLs.'),
        ('Profile', 'View the account email and type, and save a changed display name.'),
        ('Allowance and Help', 'View remaining checks and read the in-app text-checking and recovery guidance.'),
    ])
    h(d, '3.1 Account Types and Allowances')
    table(d, ['Account type', 'Text-check allowance', 'Reset'], [
        ('Free', 'One successful check per day', 'Midnight Singapore time each day'),
        ('Premium account already enabled for the prototype', 'Up to 60 successful checks per calendar month', 'Midnight Singapore time on the first day of the month'),
    ])
    p(d, 'A completed Not Enough Information result uses one check. A failed analysis does not use the allowance. Reading an existing result or opening its sources does not submit a new check. The Premium instructions apply when the team has already enabled that account type; there is no payment step in this manual.')
    h(d, '3.2 Understanding the Result')
    table(d, ['Displayed item', 'How to read it'], [
        ('Low Concern', 'The current assessment indicates lower concern. Review its supporting evidence.'),
        ('Needs Caution', 'Read the explanation and evidence carefully before sharing.'),
        ('High Concern', 'The assessment indicates greater concern; check the cited evidence and recommended action.'),
        ('Not Enough Information', 'A suitable claim or sufficient evidence could not be established. This is not proof that the content is false.'),
        ('Risk indicator', 'Higher values mean greater concern in the current scoring method. A dash means no score is available.'),
        ('Uncertainty', 'Read the level and accompanying reasons to understand limitations of the assessment.'),
    ])
    return d


def interfaces():
    d = new_doc(); h(d, '4. Initial GUIs of uiabo', 1)
    p(d, 'This chapter follows the screens used in the current application. Button names match the app. The figures are actual Android captures; only the text-checking workflow described below should be used. Cards marked LOCK or SOON in the Home captures do not perform an analysis.')
    h(d, '4.1 Sign-in Page')
    p(d, 'When no account is signed in, UIABO opens the Sign in screen. It contains Email address, Password, Sign in, Create a free account and Forgot password?. If a verified session is already saved on the emulator, the app can open Home directly.')
    steps(d, ['Enter the email address used for your account.', 'Enter your password.', 'Select Sign in.', 'If Verify your email appears, follow Section 4.2.1. Otherwise, wait for Home to open.'])
    h(d, '4.2 Create a Free Account')
    steps(d, ['From Sign in, select Create a free account.', 'Enter your name and an email address you can access.', 'Enter a password with at least eight characters, then enter the same password in Confirm password.', 'Read the account/content storage statement and select its checkbox if you agree.', 'Select Create account and wait for Verify your email.'])
    p(d, 'The registration screen contains a storage-consent checkbox. It does not open a separate privacy-policy page. If you need clarification about the statement, ask the person running the demonstration before registering.')
    h(d, '4.2.1 Email Verification', 3)
    steps(d, ['Open the verification email in your email application or browser.', 'Follow the verification link.', 'Return to UIABO and select I have verified my email.', 'If verification is complete, Home opens. If the app says the email is not verified yet, finish the email-link step and try again.'])
    p(d, 'Select Send the email again if you need another verification message. Check the inbox and spam/junk folder. Select Use a different account to leave this screen and return to account access.')
    h(d, '4.2.2 Registration Validation', 3)
    table(d, ['Message or situation', 'Action'], [
        ('Enter your name / email address', 'Complete the indicated field.'),
        ('Use a password with at least 8 characters', 'Enter a password of at least eight characters.'),
        ('The passwords do not match', 'Make Password and Confirm password identical.'),
        ('Please agree to the privacy notice to create an account', 'Review the displayed statement. The checkbox is required to submit registration.'),
        ('An account already uses this email address', 'Return to sign-in or use password recovery for your existing account.'),
    ])
    h(d, '4.3 Incorrect Sign-in Details')
    p(d, 'If UIABO displays The email or password is incorrect, check the email address and re-enter the password. Select Sign in again. If you cannot remember the password, use Forgot password?. A message about too many attempts means you should wait before retrying. An unavailable account cannot proceed until its access issue is resolved with the demonstration organiser.')
    h(d, '4.4 Password Recovery')
    steps(d, ['From Sign in, select Forgot password?.', 'Enter your account email address.', 'Select Send reset email.', 'Check your email and follow the password-reset link.', 'Select Back to sign in, then sign in using the new password.'])
    p(d, 'The app displays a neutral confirmation: If an account uses that email, Firebase has sent reset instructions. This message does not confirm whether a particular email has an account.')
    h(d, '4.5 Free-user Home')
    p(d, 'Home shows your greeting, Free badge and Today’s allowance. Select Check text to start a check. Use the navigation at the bottom to open Results, Help or Profile. You can also open Profile using the initials near the top of Home.')
    figure(d, 'home-free.png', 'Figure 4.1. Actual free-user Home, with one check available.')
    p(d, 'The full Home capture includes inactive cards and a promotional banner. Their presence is not an available checking or purchase workflow. For this prototype, select Check text.')
    h(d, '4.5.1 Refreshing and Using the Allowance', 3)
    steps(d, ['Read the number of checks available on Home.', 'Pull down on Home to refresh account and allowance information.', 'If the allowance could not be loaded, select the retry message.', 'When no checks remain, read the reset time and use Results to review checks you have already saved.'])
    h(d, '4.6 Premium-user Home')
    p(d, 'An account already enabled as Premium shows a Premium badge and This month’s allowance. The working text-check flow is the same as for a free account, with up to 60 successful checks in a calendar month.')
    figure(d, 'home-premium.png', 'Figure 4.2. Actual Premium Home and monthly text-check allowance.')
    p(d, 'The number shown is the remaining allowance at the time of the capture. Refresh Home to see your account’s current value. The additional Home cards remain unavailable in this prototype.')
    h(d, '4.7 Text Analysis')
    steps(d, ['On Home, select Check text.', 'In Message or caption, paste or enter a short English message, claim or caption.', 'Keep the text within the displayed 5,000-character limit. Do not include passwords or private information.', 'Review the text, then select Check this text.', 'Keep the screen open while Checking the claim and its sources… is displayed.', 'When the check completes, read the Analysis result screen.'])
    p(d, 'Example text for learning the input flow: Singapore became independent on 9 August 1965. A completed practice check uses one allowance; the result depends on the evidence returned during that check.')
    p(d, 'The Webpage link tab currently directs users to copied-text checking. If you arrive there, select Check copied text, then paste the article’s relevant wording into the text field. A URL alone is not a direct webpage analysis.')
    h(d, '4.7.1 While a Check Is Running', 3)
    p(d, 'UIABO shows a progress indicator and prevents navigation within the app while the check runs. Processing may take a couple of minutes. Avoid repeatedly pressing the submit button or closing the app during this period.')
    h(d, '4.7.2 Text-check Errors and Retry', 3)
    table(d, ['Situation', 'What to do'], [
        ('Empty field', 'Enter the message when Paste a message to check is shown.'),
        ('No remaining checks', 'Read saved results or wait for the account’s next reset.'),
        ('Allowance not loaded', 'Select Refresh allowance or return to Home and refresh.'),
        ('Invalid or unsupported input', 'Read the error and supply usable English text within the length limit.'),
        ('Connection failure or a long-running request', 'Select View saved results or open Results before submitting again. The first check may already have completed.'),
        ('A failed check is shown', 'Read the error and resolve the cause before starting a new check. A failed server analysis does not use the allowance.'),
    ])
    h(d, '4.8 Analysis Result')
    p(d, 'The Analysis result screen displays the saved assessment. Scroll down to read the complete explanation and evidence; the first screen does not show every source.')
    figure(d, 'result.png', 'Figure 4.3. Actual saved result showing a risk indicator, uncertainty and the identified claim.')
    p(d, 'This capture illustrates the screen layout, not a verified conclusion about the example message.')
    steps(d, ['Read the concern label and risk indicator, if present.', 'Read Uncertainty and its reasons.', 'Check Claim identified, or Submitted text if no claim was extracted.', 'Read What the evidence suggests and Recommended action.', 'Continue down to Evidence sources and review the passages and source links.'])
    h(d, '4.8.1 Opening Evidence Sources', 3)
    steps(d, ['Find an evidence source on the result screen.', 'Read its publisher, title, passage and stance label.', 'Select the source title or Open source to open the cited page.', 'Use Android’s Back control to return to UIABO when you have finished reading.'])
    p(d, 'Publication date unavailable means the app did not establish a publication date for that item. It should not be read as an estimated date. If a link cannot open, check the connection and try another available source.')
    h(d, '4.8.2 Not Enough Information and Failed Checks', 3)
    p(d, 'Not Enough Information is a completed outcome when the app cannot identify a suitable checkable claim or establish sufficient evidence. Read the explanation and uncertainty. The screen can show a dash and No risk score rather than a number. This completed outcome uses one check.')
    p(d, 'A Failed result instead says This check could not finish, shows the available error message and states that the allowance was not used. Select Start a new check after addressing the cause.')
    h(d, '4.9 Saved Result History')
    steps(d, ['Select Results in the bottom navigation.', 'Review Your saved checks. Each row shows the claim or submitted text, saved time and concern label.', 'Select a row to open its explanation and sources.', 'Pull down to refresh the list. Select Load more results when that button appears.', 'If the list could not load, select Retry.'])
    figure(d, 'history.png', 'Figure 4.4. Actual result history for the signed-in account.')
    p(d, 'If no results are saved, the screen offers Check some text. The list belongs to the signed-in account. Reopening an existing result does not start another analysis or use another check.')
    h(d, '4.10 Profile and Sign-out')
    steps(d, ['Select Profile, or select the initials on Home.', 'Read My profile. The screen shows Name, Email address and Account type.', 'To change the displayed name, edit Name and select Save name.', 'Wait for Your profile has been updated.', 'Select Back to home to return.'])
    p(d, 'Email address and Account type are read-only on this screen. Save name changes the name only.')
    h(d, '4.10.1 Signing Out', 3)
    steps(d, ['On My profile, select Sign out.', 'Read the confirmation that you will need your email and password to sign in again.', 'Select Sign out to confirm, or Cancel to stay signed in.'])
    h(d, '4.11 Help')
    steps(d, ['Select Help in the bottom navigation.', 'Read the guidance on checking text, interpreting concern scores, allowances and saved results.', 'Select Check some text to open the text-checking screen.'])
    table(d, ['Question', 'Available action'], [
        ('The result seems wrong', 'Review the claim, evidence passages and source pages before relying on the result.'),
        ('The app says it cannot reach UIABO', 'Check the connection. Ask the person running the demonstration to check the backend and Expo terminals.'),
        ('The request took too long', 'Open Results first to see whether the original check was saved.'),
        ('The app shows Finish app setup', 'Ask the installer to check the mobile Firebase configuration using Chapter 2.'),
    ])
    h(d, '4.12 Sharing a Result')
    steps(d, ['Open a completed result and review the explanation and sources.', 'Select Share result.', 'Choose a receiving application from the Android share sheet.', 'Review the prepared summary and citation URLs in the receiving application before sending.'])
    p(d, 'Sharing uses the Android share sheet to pass summary text and source URLs. The shared text includes the automated-assessment caveat. Cancelling the share sheet does not change the saved result.')
    return d


def format_doc(d, combined=False):
    for name, size in [('Normal', 11), ('Title', 23), ('Subtitle', 13), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
        style = d.styles[name]; style.font.name = 'Arial'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.12
        if name.startswith('Heading'):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(12)
    for level in range(1, 4):
        name = f'TOC {level}'
        if name not in d.styles:
            from docx.enum.style import WD_STYLE_TYPE
            style = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = d.styles['Normal']
        style = d.styles[name]
        style.font.name = 'Arial'; style.font.size = Pt(10)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2)
    for para in d.paragraphs:
        if para.style.name.startswith('Heading'):
            for run in para.runs:
                run.font.name = 'Arial'; run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.size = d.styles[para.style.name].font.size
            if para.style.name == 'Heading 1' and re.match(r'^\d+\.', para.text):
                para.paragraph_format.page_break_before = True
        if para.style.name == 'Caption':
            for run in para.runs: run.font.size = Pt(10)
    for t in d.tables:
        if len(t.columns) == 2:
            t.autofit = False
            t.columns[0].width = Cm(4.5); t.columns[1].width = Cm(12.5)
            for row in t.rows:
                row.cells[0].width = Cm(4.5); row.cells[1].width = Cm(12.5)
        for index, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(3)
                    para.paragraph_format.line_spacing = 1.0
                    para.paragraph_format.keep_with_next = False
                    for run in para.runs:
                        run.font.name = 'Arial'; run.font.size = Pt(10)
                        if index == 0: run.bold = True
    for section in d.sections:
        section.page_width = Cm(21); section.page_height = Cm(29.7)
        section.top_margin = Cm(2); section.bottom_margin = Cm(2)
        section.left_margin = Cm(2); section.right_margin = Cm(2)
        section.header_distance = Cm(.8); section.footer_distance = Cm(.8)
        section.different_first_page_header_footer = combined
        header = section.header.paragraphs[0]
        header.text = 'uiabo  |  Preliminary User Manual  |  FYP-26-S3-30'
        for run in header.runs: run.font.name = 'Arial'; run.font.size = Pt(9)
        footer = section.footer.paragraphs[0]; footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('Page '); field(footer, ' PAGE ')
    d.core_properties.title = 'uiabo — Preliminary User Manual'
    d.core_properties.author = 'FYP-26-S3-30 project team'
    d.core_properties.subject = 'Implemented Android text-checking prototype — version 0.2'


def main():
    PARTS.mkdir(exist_ok=True); ASSETS.mkdir(exist_ok=True)
    for source_name, name in [('home-before-check.png', 'home-free.png'), ('home-premium.png', 'home-premium.png'), ('result.png', 'result.png'), ('history.png', 'history.png')]:
        copy2(REPO/'evaluation/reports/app-integration'/source_name, ASSETS/name)
    builders = [('00_Cover_Document_Control_and_Contents.docx', cover),
                ('01_Introduction.docx', introduction), ('02_Initial_Installation_Instructions.docx', installation),
                ('03_Key_Features.docx', features), ('04_Initial_GUIs.docx', interfaces)]
    for name, builder in builders:
        d = builder(); format_doc(d, name.startswith('00_')); d.save(PARTS/name)
        print('Created', name)
    master = Document(PARTS/builders[0][0]); composer = Composer(master)
    for name, _ in builders[1:]: composer.append(Document(PARTS/name))
    format_doc(master, True); composer.save(MASTER)
    print('Saved', MASTER.name, 'with', len(master.inline_shapes), 'actual app captures. No PDF generated.')


if __name__ == '__main__':
    main()
