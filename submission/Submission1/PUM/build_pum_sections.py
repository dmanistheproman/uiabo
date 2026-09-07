"""Build editable UIABO PUM sections and a combined draft from existing sources.

Run with the workspace Python: python build_pum_sections.py
Only generated outputs in this PUM folder are overwritten; source PDFs stay intact.
"""

from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
WORKSPACE = ROOT.parent
REPO = WORKSPACE if (WORKSPACE / 'backend').is_dir() else WORKSPACE / 'uiabo'
PARTS = HERE / 'PUMParts'
ASSETS = HERE / 'assets'
CODE = 'FYP-26-S3-30'
DATE = '7 September 2026'
TEAM = [
    ('Ho Sze Wei, Donovan', '8949724', 'swdho001@mymail.sim.edu.sg'),
    ('Wong Yi Da', '1037274', 'ydwong003@mymail.sim.edu.sg'),
    ('Matthew Alexander Peeris', '1052585', 'peeris001@mymail.sim.edu.sg'),
    ('Chu Wai Chung', '8949256', 'wcchu001@mymail.sim.edu.sg'),
    ('Poon Chun Ping', '8865747', 'cppoon001@mymail.sim.edu.sg'),
]


def p(d, text, style=None):
    return d.add_paragraph(text, style)


def h(d, text, level=1):
    return d.add_heading(text, level)


def bullets(d, lines):
    for line in lines:
        p(d, line, 'List Bullet')


def steps(d, lines):
    for i, line in enumerate(lines, 1):
        p(d, f'{i}. {line}')


def source(d, text):
    para = p(d, 'Source: ' + text)
    for run in para.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string('526477')


def table(d, headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for c, value in zip(t.rows[0].cells, headers):
        c.text = value
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c, value in zip(t.add_row().cells, row):
            c.text = value
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                cell.width = Cm(widths[i])
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(5)
                for run in para.runs:
                    run.font.size = Pt(10)
    return t


def code(d, text):
    for line in text.strip().splitlines():
        para = p(d, line)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.keep_with_next = False
        for run in para.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(9)


def field(para, instruction):
    run = para.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t'); placeholder.text = 'Update field in Word.'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, text, separate, placeholder, end])


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width = Cm(21); s.page_height = Cm(29.7)
    s.top_margin = Cm(1.9); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2); s.right_margin = Cm(2)
    s.header_distance = Cm(.8); s.footer_distance = Cm(.8)
    for name, size in [('Normal', 11), ('Title', 25), ('Heading 1', 17), ('Heading 2', 14), ('Heading 3', 12)]:
        st = d.styles[name]
        st.font.name = 'Arial'; st.font.size = Pt(size)
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.line_spacing = 1.08
        if name.startswith('Heading'):
            st.font.color.rgb = RGBColor.from_string('173F56')
            st.paragraph_format.keep_with_next = True
    from docx.enum.style import WD_STYLE_TYPE
    front = d.styles.add_style('Frontmatter Heading', WD_STYLE_TYPE.PARAGRAPH)
    front.base_style = d.styles['Heading 2']
    outline = OxmlElement('w:outlineLvl'); outline.set(qn('w:val'), '9')
    front.element.get_or_add_pPr().append(outline)
    header = s.header.paragraphs[0]
    header.text = f'uiabo.   |   Preliminary User Manual   |   {CODE}'
    for run in header.runs:
        run.font.name = 'Arial'; run.font.size = Pt(9)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Draft 0.1  •  ' + DATE + '  |  Page ')
    field(footer, ' PAGE ')
    for run in footer.runs:
        run.font.size = Pt(9)
    d.core_properties.title = 'uiabo — Preliminary User Manual'
    d.core_properties.author = 'FYP-26-S3-30 project team'
    d.core_properties.subject = 'Editable preliminary manual adapted from PRD, URS, TDM and PTD'
    return d


def cover(d, combined=False):
    p(d, 'School of Computing and Information Technology')
    p(d, 'CSIT321 — Project')
    h(d, 'uiabo.', 0)
    h(d, 'Preliminary User Manual', 0)
    p(d, 'Project topic: CSIT-26-S3-30')
    p(d, 'AI-Based Misinformation Detection for Short-Form Content')
    p(d, f'Group number: {CODE}')
    p(d, 'Supervisor: Mr Liaw Chun Huei')
    p(d, 'Assessor: Mr Terrence Chew')
    p(d, f'Version 0.1 — Working draft — {DATE}')
    p(d, 'Project team', 'Frontmatter Heading')
    table(d, ['Student name', 'UOW ID', 'Student email'], TEAM)
    p(d, 'Prepared for the preliminary prototype submission. This draft contains both current prototype instructions and clearly identified proposed interfaces from the submitted Technical Design Manual.')
    source(d, 'PTD cover; PRD cover; URS cover and Section 1.2. The previous team’s PUM is used only for document structure.')
    d.add_page_break()
    h(d, 'Document Control')
    table(d, ['Field', 'Value'], [
        ('Title', 'uiabo — Preliminary User Manual'),
        ('Document identifier', CODE + '_PrelimUserManual'),
        ('Version / status', '0.1 / Working draft for team review'),
        ('Prepared for', 'CSIT321 preliminary prototype submission'),
        ('Scope baseline', 'Submitted PRD, URS, TDM and adapted PTD sections'),
        ('Implementation snapshot', DATE + '; text-check Android prototype'),
    ])
    p(d, 'Record of Revision', 'Frontmatter Heading')
    table(d, ['Date', 'Description', 'Sections', 'Version'], [
        (DATE, 'Initial UIABO draft compiled from existing team documents; current prototype differences recorded.', 'All', '0.1'),
    ])
    p(d, 'How to read the draft', 'Frontmatter Heading')
    bullets(d, [
        'Current prototype: a function present in the Android implementation reviewed on the date above.',
        'Proposed design: a screen or workflow documented in the URS/TDM that is not yet available in the working prototype.',
        'TDM figures are design mock-ups. Names, example claims, scores, usage totals and billing details shown in them are illustrative.',
        'Missing information and unresolved differences are recorded in PUM_MISSING_INFORMATION.md. Detailed source mapping is in PUMParts/PUM_SOURCES_AND_GAPS.md.',
    ])
    d.add_page_break()
    p(d, 'Table of Contents', 'Frontmatter Heading')
    if combined:
        field(p(d, ''), ' TOC \\o "1-3" \\h \\z \\u ')
    else:
        for line in ['1. Introduction', '2. Initial Installation Instructions', '3. Key Features of uiabo', '4. Initial GUIs of uiabo']:
            p(d, line)
        p(d, 'The combined draft contains the automatic page-numbered contents. These individual files are provided for section-by-section editing.')


def introduction(d):
    h(d, '1. Introduction')
    p(d, 'uiabo is an AI-assisted Android application that helps users assess potentially misleading short-form online content before believing or forwarding it. The application is designed primarily for Singapore citizens aged 65 and above. It presents an identified claim, a concern label, uncertainty, a plain-language explanation and links to relevant evidence in a guided interface.')
    p(d, 'This preliminary manual explains how a project demonstrator can start the Android prototype and how a first-time user can register, submit English text and read a saved result. It also introduces the proposed screens for functions that remain under development.')
    h(d, '1.1 What this manual covers', 2)
    bullets(d, [
        'Local Windows setup of the FastAPI backend, Firebase configuration and the Expo Android application.',
        'Account creation, email verification, sign-in, password reset and profile use.',
        'Home navigation, free and premium allowances, text submission, result interpretation, evidence links and saved history.',
        'Proposed webpage, image, OCR, caption-image context and deepfake/AI-image interfaces.',
        'Proposed subscription, system-administrator, data-engineer and feedback workflows.',
    ])
    h(d, '1.2 What we assume about the readers', 2)
    p(d, 'End users can use an Android smartphone, enter or paste text, open email and follow an internet link. No artificial-intelligence or fact-checking expertise is assumed. A family member or community helper may assist with onboarding.')
    p(d, 'Section 2 is intended for the project team, assessor or demonstrator preparing the local prototype. These readers need basic familiarity with PowerShell, project folders, Python, Node.js, Android Studio and Firebase. End users on a prepared device can proceed directly to Section 4.')
    h(d, '1.3 Scope and purpose', 2)
    p(d, 'The purpose of uiabo is decision support: it helps users inspect evidence and make a more informed judgement. A concern score is not a probability proven to be correct, and the application does not replace professional fact-checkers. Not Enough Information means that an assessment cannot be established from the available claim and evidence; it is not a declaration that the content is false.')
    p(d, 'The current prototype supports English text checks through an Android emulator, together with authentication, allowances and saved results. Public webpage checking and the required premium static-image functions remain planned. Audio and video analysis are outside the selected scope. Deepfake/AI-image detection remains part of the required project scope even though it is not implemented in this snapshot.')
    p(d, 'Two individual account tiers are defined: free users receive one successful submission per day, and premium users receive up to 60 per month. Operational roles are system administrator and data engineer. Their separate portal designs are introduced later in the manual.')
    source(d, 'PTD Sections 1, 2.1 and 2.3; PRD Sections 1.2 and 1.4; URS Sections 1, 2.1, 2.3–2.7; current prototype scope confirmed in the project conversation and mobile README.')


def installation(d):
    h(d, '2. Initial Installation Instructions')
    p(d, 'These instructions prepare the local Windows demonstration environment. They are for the person hosting the prototype, not tasks that every older end user must perform. A public download, production API address and final installation package have not been supplied.')
    h(d, '2.1 Prerequisites', 2)
    table(d, ['Item', 'Requirement for this draft'], [
        ('Project files', 'A team-provided copy of uiabo containing backend/ and mobile/. The examples use C:\\Dev\\uiabo; change that path for another checkout.'),
        ('Backend', 'Python and pip. The repository README specifies Python 3.11 or newer; the existing working environment reports Python 3.14.3. Install the pinned backend requirements.'),
        ('Mobile tools', 'Node.js with npm, Android Studio with an Android Virtual Device, and Expo Go. The working Node.js version reports v24.14.0; minimum compatibility has not been independently established.'),
        ('Application dependencies', 'Use mobile/package-lock.json. The current package file uses Expo 57, React Native 0.86.3 and React 19.2.3.'),
        ('Firebase', 'One Firebase project with Email/Password Authentication and Cloud Firestore. The mobile web configuration and backend Admin credential must refer to that project.'),
        ('Analysis services', 'Backend-only Ollama Cloud, Google Fact Check Tools API and Tavily credentials, supplied privately by the authorised project members.'),
        ('Connectivity', 'Internet access for Firebase, cloud models, evidence search and cited sources. Local-only/offline fact checking is not provided.'),
    ])
    p(d, 'The current configuration uses Ollama Cloud, so the demonstration computer does not need a local model installation.')
    source(d, 'PTD Section 6.2; TDM Section 3.1; URS Sections 2.4–2.5. Exact commands and dependency details supplement those documents using the current repository READMEs, requirements and package files.')
    d.add_page_break()
    h(d, '2.2 Backend setup', 2)
    h(d, '2.2.1 Install the Python dependencies', 3)
    p(d, 'Open PowerShell in the project checkout. If a configured virtual environment already exists, reuse it rather than recreating it.')
    code(d, r'''
cd C:\Dev\uiabo\backend
python -m venv .venv
.\.venv\Scripts\python.exe -m pip install -r requirements.txt
''')
    p(d, 'Calling the virtual environment’s Python directly avoids a dependency on PowerShell activation scripts.')
    h(d, '2.2.2 Configure Firebase', 3)
    steps(d, [
        'Use the team’s Firebase project. Confirm that a Cloud Firestore database is available.',
        'In Firebase Authentication, enable the Email/Password sign-in method.',
        'Obtain an authorised Firebase Admin service-account JSON file through the team’s private setup process. Store it outside the repository.',
        'Use Firebase Console → Project settings → Your apps to obtain the Firebase Web app configuration for the Expo client. Section 2.3 explains where to place it.',
    ])
    p(d, 'The backend verifies the Firebase identity token and controls account access, allowance and saved results. Do not make the Firestore database publicly readable or writable to bypass a configuration error. The complete new-project rules/IAM provisioning procedure still needs to be supplied; the current flow has been exercised against the existing configured project.')
    h(d, '2.2.3 Configure backend environment variables', 3)
    p(d, 'For a new checkout, copy uiabo/.env.instructions to uiabo/.env and replace the placeholders. Keep an existing configured .env. The example below contains no working credentials.')
    code(d, '''
OLLAMA_API_KEY="YOUR_OLLAMA_CLOUD_KEY"
GOOGLE_FACT_CHECK_API_KEY="YOUR_FACT_CHECK_TOOLS_KEY"
TAVILY_API_KEY="YOUR_TAVILY_KEY"
GOOGLE_APPLICATION_CREDENTIALS="C:/private/uiabo-service-account.json"
''')
    p(d, 'Enable Fact Check Tools API in the Google Cloud project associated with its key. Supply the Tavily key itself, not the MCP URL. Keep all provider keys and the service-account JSON out of Git and out of mobile EXPO_PUBLIC_ variables.')
    source(d, 'Current uiabo README, .env.instructions, backend/app/firebase.py and mobile README; TDM Sections 3.1 and 7.10.')
    d.add_page_break()
    h(d, '2.2.4 Start and check the backend', 3)
    code(d, r'''
cd C:\Dev\uiabo\backend
.\.venv\Scripts\python.exe -m uvicorn app.main:app --host 127.0.0.1 --port 8000
''')
    p(d, 'Keep this terminal running. Open http://127.0.0.1:8000/health on the host computer. A successful health response establishes that the API process is reachable; it does not by itself prove that Firebase and every analysis provider are configured correctly.')
    p(d, 'The API documentation is available locally at http://127.0.0.1:8000/docs. Protected account and analysis operations require a verified Firebase identity token, which the signed-in app supplies automatically.')
    p(d, 'The backend loads the project-root .env automatically. Existing shell environment variables take precedence. Restart the backend after changing keys or the credential path.')
    h(d, '2.3 Frontend / Android setup', 2)
    h(d, '2.3.1 Configure the mobile environment', 3)
    p(d, 'For a new checkout, copy mobile/.env.example to mobile/.env. Keep an existing configured file. Replace the Firebase placeholders with the Web app configuration from the same Firebase project as the backend.')
    code(d, '''
EXPO_PUBLIC_FIREBASE_API_KEY=YOUR_FIREBASE_WEB_API_KEY
EXPO_PUBLIC_FIREBASE_AUTH_DOMAIN=YOUR_PROJECT.firebaseapp.com
EXPO_PUBLIC_FIREBASE_PROJECT_ID=YOUR_PROJECT
EXPO_PUBLIC_FIREBASE_STORAGE_BUCKET=YOUR_PROJECT.firebasestorage.app
EXPO_PUBLIC_FIREBASE_MESSAGING_SENDER_ID=YOUR_SENDER_ID
EXPO_PUBLIC_FIREBASE_APP_ID=YOUR_WEB_APP_ID
EXPO_PUBLIC_API_URL=http://10.0.2.2:8000
''')
    p(d, 'Copy the storage-bucket value exactly from Firebase rather than assuming its suffix. Firebase Web configuration identifies the client project; it is different from the secret Admin service-account credential.')
    p(d, 'For the standard Android Studio emulator, 10.0.2.2 reaches the host computer. Using localhost inside the emulator would refer to the emulator itself. The HTTP address above is for the local demonstration; the submitted design calls for HTTPS in a deployed environment.')
    source(d, 'mobile/.env.example, mobile README, backend README; URS Section 2.5 and TDM Section 3.1.')
    d.add_page_break()
    h(d, '2.3.2 Install and run the Android application', 3)
    steps(d, [
        'Start the team’s Android Virtual Device in Android Studio Device Manager. The existing demonstration uses a Pixel 7 emulator.',
        'Open a second PowerShell terminal and run the commands below.',
        'Press a in the Expo terminal to open Android. If needed, open Expo Go and the development URL shown by Expo.',
        'Keep both the backend and Expo terminals running while using the prototype.',
    ])
    code(d, r'''
cd C:\Dev\uiabo\mobile
npm.cmd ci
npx.cmd expo start --lan
''')
    p(d, 'On the existing Windows setup, npx.cmd avoids PowerShell script execution-policy errors. If the emulator cannot download the bundle through the generated address, open exp://10.0.2.2:8081 in Expo Go while Metro is running on port 8081. Restart Expo after changing mobile/.env.')
    h(d, '2.3.3 First successful run', 3)
    steps(d, [
        'Select Create a free account, enter the requested details and create the account.',
        'Open the verification link in the email, return to uiabo and select I have verified my email.',
        'Confirm that Home shows the account name and remaining allowance.',
        'Open Check text, enter a short English claim and select Check this text.',
        'Read the result and its evidence, then open Results and reopen the saved item.',
    ])
    p(d, 'A completed demonstration check uses one allowance. Use a dedicated demonstration account. A returned result demonstrates the workflow; it is not by itself an accuracy test.')
    h(d, '2.3.4 Installation troubleshooting', 3)
    table(d, ['Problem', 'Action'], [
        ('Expo cannot load the app', 'Confirm the emulator and Metro are running. Check the development URL and port; use the emulator host address above if necessary.'),
        ('App cannot contact the backend', 'Check the host /health endpoint, EXPO_PUBLIC_API_URL and that port 8000 is in use by the backend. Restart Expo after configuration changes.'),
        ('Sign-in succeeds but profile cannot load', 'Check that frontend Firebase configuration and backend credentials refer to the same project, and that the account is verified and active.'),
        ('Analysis fails or times out', 'Check provider availability and backend configuration. Open Results before resubmitting because a result may have saved during a connection interruption.'),
    ])
    source(d, 'mobile README; AuthScreen, VerifyEmailScreen and TextCheckScreen; app-integration report dated 7 September 2026. A clean-machine installation matrix is not yet available.')


def features(d):
    h(d, '3. Key Features of uiabo')
    p(d, 'uiabo combines account access, guided content submission, claim identification, evidence retrieval and understandable results. The following table separates the agreed product design from the current text prototype.')
    table(d, ['Feature', 'Purpose', 'Status at ' + DATE], [
        ('Account access', 'Create a free account, verify email, sign in, reset a password and sign out.', 'Available'),
        ('Profile', 'View account details and update the display name.', 'Available; email and role are read-only. Account deletion is planned.'),
        ('Text checking', 'Identify a factual claim and search published fact-checks and relevant evidence.', 'Available for English short-form text; assessment accuracy remains under evaluation.'),
        ('Results and citations', 'Show concern, risk indicator, uncertainty, claim, explanation, evidence and next action.', 'Available; scores are not proof.'),
        ('Personal history', 'Reopen the signed-in user’s saved analyses.', 'Viewing available; deleting history is planned.'),
        ('Usage allowance', 'Display and enforce one successful check per day for free users or 60 per month for premium users.', 'Available; Premium accounts can be granted by the project team. Payment flow is not available.'),
        ('Webpage links', 'Screen a supported public URL and analyse extracted article text.', 'Proposed design; current link tab offers copied-text checking only.'),
        ('Premium image functions', 'Image + caption, OCR correction, caption-image context and deepfake/AI-image likelihood.', 'Required/planned scope; not implemented.'),
        ('Sharing and feedback', 'Share useful findings, report incorrect results and review the service.', 'Native text/source sharing available; public result links, reports and reviews are planned.'),
        ('Operational portals', 'Manage authorised accounts and feedback; investigate reports and monitor ingestion.', 'TDM designs available; portals not implemented.'),
    ])
    source(d, 'PTD Sections 1 and 2; PRD Sections 1.2, 1.4 and 1.5; URS Sections 2.2–2.3. Current status is supplemented from the app integration report and source code.')
    d.add_page_break()
    h(d, '3.1 User roles and allowances', 2)
    table(d, ['Role', 'Access described by the project'], [
        ('Free user', 'Registered individual user. One successful submission per day. Text is available now; public webpage checking is planned. Premium image tools are locked.'),
        ('Premium user', 'Individual account with up to 60 successful submissions per month. Text checks are available now; image tools are marked as coming soon.'),
        ('System administrator', 'Proposed operational role for account queries, account status changes, authorised operational-account creation and feedback review.'),
        ('Data engineer', 'Proposed operational role for incorrect-result investigations, provenance review and data-ingestion monitoring.'),
    ])
    p(d, 'In the current prototype, the free allowance resets at midnight Singapore time each day. The premium allowance resets at midnight Singapore time on the first day of each calendar month. Completed Not Enough Information results consume one check. Technical failures do not. If the connection breaks, check saved history before starting a new submission because the server may already have completed the original check.')
    h(d, '3.2 Understanding the outputs', 2)
    table(d, ['Output', 'Meaning for the user'], [
        ('Low Concern', 'The available assessment supports the claim. Read the cited evidence; the label is not a guarantee.'),
        ('Needs Caution', 'The assessment identifies a reason for caution, such as conflicting evidence. Review the explanation and uncertainty before sharing.'),
        ('High Concern', 'The assessment identifies evidence against the claim. Review that evidence before forwarding the content.'),
        ('Not Enough Information', 'No suitable factual claim or sufficient decisive evidence was available. It does not automatically mean false.'),
        ('Misinformation-risk indicator', 'A numerical indicator of concern about the claim, rather than a proven probability of falsehood.'),
        ('AI-generation/manipulation likelihood', 'A separate planned image-authenticity output. An authentic image can have a misleading caption, and a manipulated image does not establish the truth of an unrelated claim.'),
    ])
    source(d, 'PTD Sections 1 and 2.1; URS Sections 1.1, 2.5–2.6 and UC-10/11/24. Calendar reset and completion behavior: current account/analysis implementation.')


def gui(d, title, status, description, actions=(), image=None, caption=None, refs='', note='', level=2, new_page=True):
    if new_page:
        d.add_page_break()
    h(d, title, level)
    para = p(d, status)
    for run in para.runs:
        run.bold = True
    p(d, description)
    if image:
        from PIL import Image
        path = ASSETS / image
        width, height = Image.open(path).size
        if height > width:
            t = d.add_table(rows=1, cols=2)
            t.autofit = False
            t.columns[0].width = Cm(6.5); t.columns[1].width = Cm(10.3)
            a, b = t.rows[0].cells
            a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            a.paragraphs[0].add_run().add_picture(str(path), width=Cm(6.1))
            b.paragraphs[0].text = 'How to use' if status.startswith('Current') else 'Proposed user flow'
            for run in b.paragraphs[0].runs:
                run.bold = True
            steps(b, actions)
            if note:
                p(b, note)
            cp = p(d, caption, 'Caption')
            cp.paragraph_format.keep_with_next = False
        else:
            para = p(d, '')
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run().add_picture(str(path), width=Cm(15.4))
            p(d, caption, 'Caption')
            steps(d, actions)
            if note:
                p(d, note)
    else:
        steps(d, actions)
        if note:
            p(d, note)
    source(d, refs)


def interfaces(d):
    h(d, '4. Initial GUIs of uiabo')
    p(d, 'The interfaces below introduce account access, content analysis, result history, premium image tools and the two operational roles defined in the submitted URS and TDM.')
    p(d, 'Every TDM image is captioned as a proposed design. For implemented functions, the instructions use the current app’s controls, which may differ from the mock-up. Proposed workflows are retained to show the agreed design, not to imply that the screens are usable in the current build.')
    table(d, ['Area', 'Sections'], [
        ('Account access', '4.1–4.4: sign-in, registration, verification, errors and password recovery'),
        ('Individual user workflow', '4.5–4.11: Home, allowances, text/link checking, result/history and profile'),
        ('Premium design', '4.12–4.13: subscription and the four image tools'),
        ('Operational design', '4.14–4.15: system administrator and data engineer'),
        ('Guidance and feedback', '4.16: help, sharing, incorrect-result reports and reviews'),
    ])
    p(d, 'Use Home, Results, Help and Profile in the bottom navigation. The active tab is highlighted. An authenticated session may return directly to Home when the app is reopened.')
    source(d, 'TDM Section 6, printed pages 53–61 (PDF pages 56–64); URS Section 3; current mobile App.js.')

    gui(d, '4.1 Welcome and sign-in', 'Current prototype — TDM design shown',
        'A user without a saved session begins at the sign-in screen. Existing users sign in with their registered email address and password.',
        ['Enter the registered email address.', 'Enter the password.', 'Select Sign in.', 'If prompted, complete email verification. An active, verified account opens Home.', 'Select Create a free account to register, or Forgot password? to recover access.'],
        'TDM_pdf056_figure01.png', 'Figure 4.1. Proposed login design from TDM 6.1. The current button reads “Sign in”.',
        'URS UC-01 and UC-04; TDM 6.1, PDF p. 56; mobile AuthScreen and VerifyEmailScreen.',
        'The sign-in form is the entry point for users without a saved session.')

    gui(d, '4.2 Create a free account', 'Current prototype — TDM design shown',
        'Registration creates an individual free account. The TDM also illustrates premium selection during registration, but the current app does not provide paid registration.',
        ['Select Create a free account from Sign in.', 'Enter your name and email address.', 'Create a password with at least eight characters, then enter it again in Confirm password.', 'Read the consent wording and select the checkbox.', 'Select Create account and continue with email verification.'],
        'TDM_pdf056_figure02.png', 'Figure 4.2. Proposed account-creation design. Premium choice and payment in this mock-up are not implemented.',
        'URS UC-16 and UC-18; TDM 6.1, PDF p. 56; current AuthScreen.',
        'The current app includes password confirmation. The final linked privacy notice and terms still need to be supplied; see PUM_MISSING_INFORMATION.md.')
    gui(d, '4.2.1 Email verification', 'Current prototype',
        'The account must be verified before content checking is available. Verification uses a link sent to the registered email address.',
        ['Open the inbox for the registered email address.', 'Open the verification message and follow its link.', 'Return to uiabo and select I have verified my email.', 'If the message is missing, check the spam folder or select Send the email again.', 'Use Use a different account to leave this verification session if the wrong email was entered.'],
        refs='URS UC-16; mobile VerifyEmailScreen and AuthContext.', level=3)
    h(d, '4.2.2 Password and consent checks', 3)
    p(d, 'If the two passwords differ, correct them before submitting. A missing name, missing email address, short password or unselected consent checkbox prevents registration and produces a message explaining what to correct.')
    p(d, 'UIABO collects the required free-account details in one registration form. The team still needs to supply the final linked privacy notice and terms of use.')
    source(d, 'URS UC-16; current AuthScreen. Privacy/terms document gap: PUM_MISSING_INFORMATION.md, G04.')

    gui(d, '4.3 Incorrect sign-in details', 'Current prototype',
        'An unsuccessful sign-in displays an error on the account form and does not open protected results.',
        ['Check the spelling of the email address and enter the password again.', 'If you cannot remember the password, use Forgot password?.', 'If an account is awaiting verification, open the email link and complete the verification step.', 'If a network or service message appears, restore the connection and try again. An unavailable or suspended account requires assistance from the project team.'],
        refs='URS UC-01 alternative flows; mobile AuthScreen, AuthContext and Firebase error mapping.')
    h(d, '4.4 Password recovery', 2)
    steps(d, ['From Sign in, select Forgot password?.', 'Enter the registered email address and select Send reset email.', 'Read the neutral confirmation and check the inbox for reset instructions.', 'Open the reset link and set the new password using the Firebase page.', 'Return to uiabo and sign in with the new password.'])
    p(d, 'A confirmation message does not reveal whether an account exists for a particular email address. If a link is invalid or expired, request another reset email.')
    source(d, 'URS UC-04; mobile AuthScreen and AuthContext. Current recovery-screen captures are not present in the submission documents.')

    gui(d, '4.5 Free-user Home and allowance', 'Current prototype — TDM design shown',
        'Home greets the signed-in user and shows the free daily allowance. The six cards follow the submitted design; there is no audio card.',
        ['Read the Free badge and remaining allowance.', 'Select Check text to paste an English message.', 'Check webpage link currently offers checking copied text.', 'Image + caption, Read image text, Check image context and Check AI image are premium tools and remain unavailable in this build.', 'Open Results for saved checks, Help for guidance or Profile for account details.'],
        'TDM_pdf057_figure01.png', 'Figure 4.3. Proposed free-user Home design from TDM 6.2.',
        'PTD Sections 1 and 2.4; URS UC-07/08/11; TDM 6.2, PDF p. 57; current HomeScreen.',
        'One completed check is available per Singapore calendar day. A completed Not Enough Information result uses that allowance. Existing results remain readable when no new checks remain.')

    gui(d, '4.6 Premium-user Home and allowance', 'Current prototype — current Android capture shown',
        'A premium account displays a Premium badge and a monthly allowance. Premium status is checked by the backend and cannot be changed through an editable profile field.',
        ['Read the monthly allowance card.', 'Use Check text in the same way as a free user.', 'Recognise that the image cards marked SOON are not yet available, even with premium access.', 'Open Results to revisit saved analyses without submitting the claim again.'],
        'current_home_premium.png', 'Figure 4.4. Android emulator capture, 7 September 2026. The displayed account and remaining count are a snapshot.',
        'URS Section 2.3 and UC-11; TDM 6.3; current HomeScreen and app-integration capture.',
        'Current premium allowance: 60 successful checks per calendar month, resetting at midnight Singapore time on the first day of the month. A project-granted Premium account does not establish a payment subscription.')

    gui(d, '4.7 Text analysis', 'Current prototype — TDM design shown',
        'Text checking accepts a short English message, claim or caption. The current form allows up to 5,000 characters.',
        ['Select Check text from Home.', 'Paste or type the message into Message or caption.', 'Check that the wording includes the important date, amount and context. Do not include passwords or private information.', 'Select Check this text once and keep the screen open while it processes.', 'Read the completed result, explanation, uncertainty and evidence.', 'Open Results to find the saved check again.'],
        'TDM_pdf057_figure02.png', 'Figure 4.5. Proposed text-analysis design from TDM 6.2; the current flow retains these main controls.',
        'URS UC-07; TDM 6.2, PDF p. 57; current TextCheckScreen.',
        'A technical failure does not consume an allowance. After a connection timeout, view saved results before starting a new check; the original request may have completed.')

    gui(d, '4.8 Public webpage-link analysis', 'Proposed design — direct webpage checking not implemented',
        'The agreed design allows a public webpage or article URL to be checked after malicious-link screening. The image below describes that intended feature.',
        ['Select Check webpage link.', 'Paste a supported public webpage or article address.', 'Select the webpage-check action.', 'The proposed service screens the link and extracts accessible article text before assessing the claim.', 'Read the result, source context and citations once processing succeeds.'],
        'TDM_pdf057_figure03.png', 'Figure 4.6. Proposed link-analysis design from TDM 6.2. Malicious-link screening shown here is not currently available.',
        'URS UC-08 and Sections 2.4–2.5; TDM 6.2, PDF p. 57; current TextCheckScreen.',
        'Current alternative: the link tab offers Check copied text. It does not fetch or screen a URL. The design excludes private, paywalled and inaccessible content.')

    gui(d, '4.9 Analysis result', 'Current prototype — current Android capture shown',
        'A completed result presents the concern label and uncertainty alongside the identified claim. Scroll to read the evidence sources and available actions.',
        ['Check that Claim identified represents the statement you intended to check.', 'Read the concern label together with the uncertainty explanation.', 'Read What the evidence suggests and the recommended action.', 'Scroll to Evidence sources and inspect the passages and source links.', 'Use Share only after reviewing what will be sent, or return to Results.'],
        'current_result.png', 'Figure 4.7. Saved-result screen from the Android emulator. This capture illustrates layout, not an independently verified verdict.',
        'URS UC-10; TDM 6.2, PDF p. 58; current ResultScreen and app-integration capture.',
        'The current assessment uses an initial scoring approach and can misinterpret evidence. Read the sources. An image-authenticity score, when implemented, will remain separate from misinformation risk.')
    h(d, '4.9.1 When no assessment can be established', 3)
    p(d, 'Not Enough Information is used when no checkable factual claim or adequate decisive evidence is available. It has no numerical risk score. A technical processing failure is displayed as a failed check, rather than presented as a successful factual assessment.')

    gui(d, '4.10 Saved result history', 'Current prototype — current Android capture shown',
        'Results lists analyses saved to the signed-in account. Reopening an item reads its saved result rather than launching another fact check.',
        ['Select Results in the bottom navigation.', 'Refresh the list if a newly completed item is missing.', 'Select an item to reopen its saved result and citations.', 'Load more items when that option is available.'],
        'current_history.png', 'Figure 4.8. Android result-history capture, 7 September 2026.',
        'URS UC-09/10/12; TDM 6.2, PDF p. 58; current ResultsScreen.',
        'Current history is private to the account. The TDM’s reported-result filters and deletion workflow are proposed features; they are not controls in the current list.')
    h(d, '4.10.1 Planned history deletion', 3)
    p(d, 'URS UC-12 specifies selecting owned results, choosing Delete and confirming permanent removal. That workflow is not implemented in this prototype. Do not assume that the Android app currently provides a deletion button.')

    gui(d, '4.11 Profile and sign-out', 'Current prototype',
        'Profile displays the name, email address and account type for the current individual user.',
        ['Select Profile in the bottom navigation.', 'Edit Name and select Save name.', 'Read the confirmation that the profile has been updated.', 'Select Back to home to return.', 'To leave the account, select Sign out and confirm, or Cancel to remain signed in.'],
        refs='URS UC-02/05/06; current ProfileScreen.',
        note='Email and account type are read-only in the current app. Password recovery starts from the sign-in screen; an in-profile email/password editor is not provided.')
    h(d, '4.11.1 Planned account deletion', 3)
    p(d, 'URS UC-03 specifies a permanent account-deletion flow with a clear warning, confirmation and re-authentication. The account and associated history would be removed. This is a requirement for later implementation; no account-deletion control is available in the current profile.')
    source(d, 'URS UC-03. Final deletion behavior and screen capture are recorded as outstanding work in the companion Markdown file.')

    gui(d, '4.12 Upgrade and subscription', 'Proposed design — payments not implemented',
        'The submitted design describes a premium plan with up to 60 successful submissions per month and the supported image tools. The TDM displays SGD 20 per month as its proposed price.',
        ['Select Upgrade to Premium.', 'Review the proposed plan, price, billing period and allowance.', 'In the intended flow, confirm the upgrade and authorise payment through the selected payment provider.', 'After confirmed payment, the account would receive premium access and subscription details.'],
        'TDM_pdf058_figure03.png', 'Figure 4.9. Proposed upgrade design from TDM 6.2. SGD 20/month is a design price, not a working checkout.',
        'PTD 2.4; URS UC-17/18; TDM 6.2, PDF p. 58; current InfoScreen.',
        'Current behavior: Premium information explains that images and payments are unavailable. The project team can grant an account Premium for demonstration. No payment is collected by this prototype.')
    gui(d, '4.12.1 Renewal and cancellation', 'Proposed design — subscription controls not implemented',
        'The proposed subscription page shows plan status, allowance, billing information and renewal/cancellation actions.',
        ['Open Subscription and review the plan and end date.', 'For renewal, select Renew Subscription and review the price and new period before authorising payment.', 'For cancellation, select Cancel Subscription and read when premium access will end.', 'Confirm the intended action and review the resulting status.'],
        'TDM_pdf060_figure03.png', 'Figure 4.10. Proposed subscription design from TDM 6.3; displayed billing details are illustrative.',
        'URS UC-19/20; TDM 6.3, PDF p. 60.',
        'The URS describes paid access through the current expiry date after cancellation. The payment provider, final billing rules and alignment with the prototype’s calendar-month allowance need confirmation.', level=3)

    gui(d, '4.13 Premium image functions', 'Proposed design — all four image tools remain unimplemented',
        'The required image scope includes image + caption, OCR text correction, caption-image context and deepfake/AI-image likelihood. Free users cannot use these functions. Premium users need an available monthly allowance once the functions are implemented.',
        refs='PTD Sections 1 and 2.1; URS UC-21–24; TDM 6.3; user-confirmed deepfake requirement.')
    gui(d, '4.13.1 Image + caption', 'Proposed design', 'This proposed workflow checks a static image together with the English caption supplied by the user.',
        ['Select Image + caption.', 'Choose a supported static image and review its preview.', 'Enter or correct the accompanying English caption.', 'Select Analyse image and caption.', 'Review the resulting explanation, uncertainty and evidence.'],
        'TDM_pdf059_figure02.png', 'Figure 4.11. Proposed image-and-caption design from TDM 6.3.',
        'URS UC-21; TDM 6.3, PDF p. 59.',
        'The accepted image formats, size limits and final image-processing services have not been established for this manual.', level=3, new_page=False)
    gui(d, '4.13.2 Caption-image context', 'Proposed design',
        'This feature compares the supplied caption with contextual evidence about the image. A real photograph can still be described with an incorrect date, location or event.',
        ['Select Check image context.', 'Choose the image and enter the caption being checked.', 'Review the proposed context cues, such as date, location, event or description.', 'Select Compare image and caption.', 'Read the explanation and evidence; insufficient context should lead to an inconclusive result.'],
        'TDM_pdf059_figure03.png', 'Figure 4.12. Proposed caption-image context design from TDM 6.3.',
        'URS UC-22; TDM 6.3, PDF p. 59.', level=3)
    gui(d, '4.13.3 Read image text (OCR)', 'Proposed design',
        'OCR extracts visible words from an image so the user can check and correct them before submitting a text analysis.',
        ['Select Read image text.', 'Choose a clear image containing readable text.', 'Extract the text and compare it with the original image.', 'Correct any recognition errors and confirm that the text has been reviewed.', 'Select Analyse corrected text, then read the resulting assessment and sources.'],
        'TDM_pdf060_figure01.png', 'Figure 4.13. Proposed OCR review/correction design from TDM 6.3.',
        'URS UC-23; TDM 6.3, PDF p. 60.',
        'Reviewing OCR output is important: a mistaken amount, name or date can change the claim being checked.', level=3)
    gui(d, '4.13.4 Check AI image / deepfake analysis', 'Proposed design — required project feature',
        'This function is intended to assess whether a supported static image may be AI-generated or manipulated. It does not analyse audio or video.',
        ['Select Check AI image.', 'Choose a supported static image and review the preview.', 'Read the notice explaining that the output is a likelihood assessment.', 'Select Analyse image.', 'Review the separate image-authenticity likelihood, uncertainty and explanation of the available detector signals.'],
        'TDM_pdf060_figure02.png', 'Figure 4.14. Proposed deepfake/AI-image design from TDM 6.3.',
        'PTD Sections 1 and 2.1; URS UC-24; TDM 6.3, PDF p. 60.',
        'Only detector signals actually supported by the selected service may be shown. A highlighted region or visual artefact must not be presented as conclusive proof. Detector selection and final screens remain outstanding.', level=3)

    gui(d, '4.14 System administrator dashboard', 'Proposed operational portal — not implemented',
        'The system administrator role manages permitted account access and reviews feedback through the proposed UIABO Backend Portal.',
        ['Sign in with an authorised, active system-administrator account.', 'Review the dashboard summaries and available account/feedback areas.', 'Open Account Management or User Feedback to perform the corresponding task.', 'Use the account/profile area for the authorised user’s own profile, password-reset and sign-out workflows.'],
        'TDM_pdf061_figure01.png', 'Figure 4.15. Proposed administrator dashboard from TDM 6.4. Dashboard numbers are illustrative.',
        'URS UC-25–29; TDM 3.1 and 6.4, PDF p. 61.',
        'No deployed portal address or operational sign-in instructions have been supplied.')
    gui(d, '4.14.1 Account management', 'Proposed operational portal',
        'Account management allows an authorised administrator to locate an account and inspect the information needed for administration.',
        ['Open Account Management and enter a supported query or filter.', 'Select Search and choose a matching account.', 'Review its permitted profile summary, role, tier and status.', 'If appropriate, choose Suspend Account, supply the required reason and confirm.', 'To restore an eligible suspended account, choose Unsuspend Account and confirm.'],
        'TDM_pdf061_figure02.png', 'Figure 4.16. Proposed account-management design from TDM 6.4.',
        'URS UC-30/31/33/34; TDM 6.4, PDF p. 61.',
        'The URS requires role checks, an audit history and safeguards for protected accounts. This manual does not provide a database-editing substitute for those controls.', level=3)
    gui(d, '4.14.2 Create operational account', 'Proposed operational portal',
        'Operational accounts are created by authorised administrators rather than by individual-user registration.',
        ['Open Account Management and select Create Operational Account.', 'Enter the requested account information.', 'Choose the authorised role: System Administrator or Data Engineer.', 'Review the information and confirm creation.', 'The intended system sends secure activation/password-setup instructions and records the action.'],
        'TDM_pdf062_figure01.png', 'Figure 4.17. Proposed operational-account creation design from TDM 6.4.',
        'URS UC-32; TDM 6.4, PDF p. 62.', level=3)
    gui(d, '4.14.3 User feedback management', 'Proposed operational portal',
        'The feedback page presents submitted reviews and feedback for authorised administrative review.',
        ['Select User Feedback.', 'Review the feedback summary and permitted records.', 'Filter or select a feedback item.', 'Read its category, status, submission date and relevant comments.', 'Use only the follow-up actions supported by the completed portal.'],
        'TDM_pdf062_figure02.png', 'Figure 4.18. Proposed administrator feedback design from TDM 6.4.',
        'URS UC-35; TDM 6.4, PDF p. 62.',
        'The available feedback actions must follow the permissions and controls of the completed portal.', level=3)

    gui(d, '4.15 Data engineer dashboard', 'Proposed operational portal — not implemented',
        'The data engineer investigates incorrect-result reports and monitors the data workflows supporting analysis.',
        ['Sign in with an authorised, active data-engineer account.', 'Review available report totals, trends and pipeline summaries.', 'Apply a relevant time range or supported filter.', 'Select a report requiring investigation, or open Pipeline Monitoring.'],
        'TDM_pdf063_figure01.png', 'Figure 4.19. Proposed data-engineer dashboard from TDM 6.5.',
        'URS UC-36–41; TDM 6.5, PDF p. 63.',
        'Operational profile, password-reset and sign-out use the corresponding role-restricted flows specified in the URS.')
    gui(d, '4.15.1 Incorrect-result report details', 'Proposed operational portal',
        'A report detail view brings together the reported problem, submitted content, generated assessment and available provenance.',
        ['Open a report from the dashboard or report list.', 'Read the user’s stated reason and relevant submitted content.', 'Compare the generated claim, concern label and evidence with the report.', 'Inspect available timestamps and pipeline/model information.', 'Record an investigation note or permitted status update and save it.'],
        'TDM_pdf063_figure02.png', 'Figure 4.20. Proposed incorrect-result investigation design from TDM 6.5.',
        'URS UC-42; TDM 6.5, PDF p. 63.',
        'The final investigation statuses and available actions must match the implemented portal; no working update action is implied by this design.', level=3)
    gui(d, '4.15.2 Pipeline monitoring', 'Proposed operational portal',
        'Pipeline monitoring is intended to show data-ingestion runs, source updates, errors and data-quality indicators.',
        ['Open Pipeline Monitoring.', 'Review the latest run states and timestamps.', 'Filter the list or select a run.', 'Inspect available error details, source provenance and the affected data range.', 'Record an issue for follow-up; use a recovery action only when one has been implemented and authorised.'],
        'TDM_pdf064_figure01.png', 'Figure 4.21. Proposed pipeline-monitoring design from TDM 6.5.',
        'URS UC-43; TDM 6.5, PDF p. 64.', level=3)

    gui(d, '4.16 Help, sharing and feedback', 'Current help/sharing; reporting and reviews remain proposed',
        'Help explains how to check text, read uncertainty and understand the allowance. The URS also specifies separate incorrect-result reports and service reviews.',
        refs='URS Section 2.6 and UC-13–15; current InfoScreen and ResultScreen.')
    h(d, '4.16.1 Help and troubleshooting', 3)
    steps(d, ['Select Help in the bottom navigation.', 'Read the text-check instructions and result limitations.', 'Read the allowance and interrupted-request guidance.', 'Select Check some text to return to the text workflow.'])
    table(d, ['Situation', 'Suggested next action'], [
        ('No allowance remains', 'Reopen existing results or wait for the next reset shown by the account tier.'),
        ('No checkable claim / insufficient evidence', 'Read the Not Enough Information explanation. Review the original wording and available authoritative sources.'),
        ('A request times out', 'Check Results before trying again. The original check may already have completed.'),
        ('A source will not open', 'Check the connection and try another citation. Source pages can become unavailable.'),
        ('A result seems wrong', 'Review the evidence rather than relying on the score. The in-app reporting workflow is still planned.'),
    ])
    h(d, '4.16.2 Share a result', 3)
    steps(d, ['Open a completed result and review its explanation and citations.', 'Select Share.', 'Choose the receiving app from the Android share sheet and review the text before sending.'])
    p(d, 'The current share action sends summary text and citation URLs. It does not generate a public UIABO result-page link. The shareable-link workflow in URS UC-15 remains planned.')
    source(d, 'Current InfoScreen and ResultScreen; URS UC-15. Support contact and final legal/help pages are missing from the supplied project material.')
    d.add_page_break()
    h(d, '4.16.3 Report an incorrect result', 3)
    p(d, 'Proposed workflow — not implemented in the Android prototype.')
    steps(d, ['Open the completed result and select Report Incorrect Result.', 'Choose a reason and optionally explain the problem.', 'Submit the report.', 'Read the confirmation; the intended system associates the report with the result for investigation.'])
    source(d, 'URS UC-13; the submitted TDM contains the operational report-detail design, but no separate consumer report-form image.')
    h(d, '4.16.4 Leave a review', 3)
    p(d, 'Proposed workflow — not implemented in the Android prototype.')
    steps(d, ['Open the feedback area and select Leave a Review.', 'Choose a rating and optionally add comments.', 'Submit the review and read the confirmation.'])
    p(d, 'An incorrect-result report concerns a particular assessment; a review concerns the experience of using uiabo.')
    source(d, 'URS UC-14 and PF-A02. A consumer review-form design and final support/contact route are outstanding.')


BUILDERS = [
    ('00_Cover_Document_Control_and_Contents.docx', cover),
    ('01_Introduction.docx', introduction),
    ('02_Initial_Installation_Instructions.docx', installation),
    ('03_Key_Features.docx', features),
    ('04_Initial_GUIs.docx', interfaces),
]


def prepare_assets():
    ASSETS.mkdir(exist_ok=True)
    pdf = PdfReader(ROOT / 'TDM/FYP-26-S3-30_TDM.pdf')
    for n in range(55, 64):
        for i, image in enumerate(pdf.pages[n].images):
            image.image.save(ASSETS / f'TDM_pdf{n+1:03d}_figure{i+1:02d}.png')
    capture_dir = REPO / 'evaluation/reports/app-integration'
    for original, new in [('home-premium.png', 'current_home_premium.png'), ('history.png', 'current_history.png'), ('result.png', 'current_result.png')]:
        copyfile(capture_dir / original, ASSETS / new)


def main():
    PARTS.mkdir(exist_ok=True)
    prepare_assets()
    for name, builder in BUILDERS:
        d = new_doc(); builder(d); d.save(PARTS / name)
        print('Created', name)
    d = new_doc()
    cover(d, combined=True)
    for _, builder in BUILDERS[1:]:
        d.add_page_break(); builder(d)
    output = HERE / f'{CODE}_PrelimUserManual_DRAFT.docx'
    d.save(output)
    print('Created', output.name)


if __name__ == '__main__':
    main()
