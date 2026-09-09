"""Reorganize UIABO's PTD to match the supplied sample's report structure.

Inputs in PTDParts_Updated and the submitted URS are read only. Outputs are the
combined Word document and PTDParts_Reformatted. No PDF is produced. Preserve
manual edits before rebuilding; refresh the contents and page fields in Word.
"""
from pathlib import Path
import re

from docx import Document
from docxcompose.composer import Composer
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
BASE = ROOT / 'PTDParts_Updated'
OUT = ROOT / 'PTDParts_Reformatted'
URS = ROOT.parents[1] / 'URS/FYP-26-S3-30_URS.pdf'
MASTER = ROOT / 'FYP-26-S3-30_Preliminary_Technical_Documentation.docx'
DATE = '8 September 2026'
FILES = {int(p.name[:2]): p for p in BASE.glob('*.docx') if p.name[:2].isdigit()}
PDF = PdfReader(URS)
CHAPTERS = [
    'Introduction', 'Overview', 'Stakeholders', 'Data Collection',
    'Project Timeline', 'Requirement Definition', 'Functional Requirements',
    'Non-functional Requirements', 'Other Requirements', 'Risk Management',
    'Development Methodologies', 'Technical Stack', 'User Stories',
    'Use Case Descriptions', 'Use Case Diagram', 'System Design',
    'Conclusion', 'Glossary', 'Appendix',
]


def text(d, value, style=None):
    return d.add_paragraph(value, style)


def heading(d, value, level=2):
    return d.add_heading(value, level)


def bullets(d, values):
    for value in values:
        text(d, value, 'List Bullet')


def table(d, headers, rows):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for values in rows:
        for cell, value in zip(t.add_row().cells, values):
            cell.text = str(value)
    return t


def source(d, value):
    p = text(d, 'Source: ' + value)
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.italic = True


REFS = {
    '8.6': '19.1.6', '13.2': '19.3.2', '2.10': '3', '5.8': '16.4.6',
    '5.5': '16.3', '1.1': '7.5', '18': '17', '15': '6',
    '13': '19.3', '9': '5.5', '8': '19.1', '5': '16',
}


def repair_references(value):
    # Only copied PTD prose uses this mapping. Source URS/PRD numbering stays intact.
    if value.startswith('Evidence / source: URS Sections'):
        return ('Source: URS Sections 2.2, 3 and 4; PTD Chapters 13–15 contain the '
                'user stories and use cases; Section 19.1 contains verification records.')
    value = re.sub(r'\bSection (\d+(?:\.\d+)*)',
                   lambda m: 'Section ' + REFS.get(m[1], m[1]), value)
    value = value.replace('Figure 5.8a.', 'Figure 16.4a.')
    value = value.replace('Figure 5.8b.', 'Figure 16.4b.')
    value = value.replace('Figure 5.8c.', 'Figure 16.4c.')
    for i in range(1, 5):
        value = value.replace(f'Figure 4.{i}:', f'Figure 11.{i}:')
    return value


def fragment(number, start=0, end=None, headings=None, prefix=None):
    d = Document(FILES[number])
    children = list(d.element.body)
    for i, child in enumerate(children):
        if child.tag == qn('w:sectPr'):
            continue
        if i < start or (end is not None and i >= end):
            d.element.body.remove(child)
    for p in d.paragraphs:
        original = p.text
        if p.style.name.startswith('Heading'):
            if headings and original in headings:
                value, level = headings[original]
                p.text = value
                p.style = d.styles[f'Heading {level}']
            elif prefix:
                old, new = prefix
                if original.startswith(old):
                    p.text = new + original[len(old):]
                    num = re.match(r'^(\d+(?:\.\d+)*)', p.text)
                    if num:
                        p.style = d.styles[f'Heading {min(4, num[1].count(".") + 1)}']
        elif original:
            new = repair_references(original)
            if new != original:
                p.text = new
        # Old fragment-level hard breaks must not split new section headings.
        for br in list(p._p.xpath('.//w:br[@w:type="page"]')):
            br.getparent().remove(br)
        p.paragraph_format.page_break_before = False
    # Discard unused image relationships from sliced input documents.
    used = set()
    for element in d.element.iter():
        for name, value in element.attrib.items():
            if name.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                used.add(value)
    for rid, rel in list(d.part.rels.items()):
        if rel.reltype.endswith(('/image', '/hyperlink')) and rid not in used:
            d.part.drop_rel(rid)
    return d


class Chapter:
    def __init__(self, number):
        self.number = number
        self.d = Document()
        self.composer = Composer(self.d)
        heading(self.d, f'{number}. {CHAPTERS[number-1]}', 1)

    def add(self, number, start=0, end=None, **kwargs):
        self.composer.append(fragment(number, start, end, **kwargs))


def introduction():
    c = Chapter(1); d = c.d
    heading(d, '1.1 Problem Statement')
    c.add(2, 90, 91)
    text(d, 'A message can appear convincing because it is repeatedly forwarded, accompanied by an image, or presented without its original date and source. A useful assessment must identify the factual claim, retrieve relevant evidence and explain any limits. A misleading caption and an AI-manipulated image also require different checks: factual accuracy and image authenticity must be assessed separately.')
    heading(d, '1.2 Purpose')
    text(d, 'This Preliminary Technical Documentation describes the requirements, project planning, technical choices and system design for uiabo, an AI-based misinformation-detection application for short-form content. It brings together the team’s PRD, URS and TDM and records the working September prototype, its validation evidence and the remaining work towards final submission.')
    heading(d, '1.3 Our Vision')
    text(d, 'The team’s vision is an accessible Android application that helps Singapore citizens aged 65 and above check suspicious content before believing or forwarding it. Users should receive a readable assessment, links to relevant evidence and a clear explanation of uncertainty. UIABO supports informed judgement; a model-generated score does not guarantee that content is true or false.')
    heading(d, '1.4 Project Objectives')
    c.add(2, 93, 94)
    text(d, 'The 12 September 2026 milestone is a basic Sprint 1 text demonstration through an Android emulator. The final submission is expected in November 2026; the exact date remains to be confirmed. Static-image deepfake detection remains required for the final product. Audio and video analysis are excluded.')
    heading(d, '1.5 Learning Objectives')
    c.add(2, 11, 17)
    source(d, 'PRD project scope and description; URS introduction; existing PTD overview and project scope clarifications.')
    return c


def overview():
    c = Chapter(2); d = c.d
    text(d, 'UIABO combines an Android application, a staged analysis pipeline and authenticated cloud storage. The following overview explains the product rationale and planned business model. The working text increment and remaining features are identified in Section 7.5.')
    heading(d, '2.1 Competitor Market Research')
    text(d, 'The following research is retained from the team’s PRD. It records the comparison used during product planning; it is not a new market or pricing survey.')
    c.add(3, 1, None, headings={'3.1 AI Misinformation Products on the Market': ('2.1.1 AI Misinformation Products on the Market', 3), '3.2 Research Findings and Analysis': ('2.1.2 Research Findings and Analysis', 3)})
    heading(d, '2.2 Product Comparison Matrix')
    c.add(4, 2, 9)
    heading(d, '2.3 Conceptualization of Ideas')
    c.add(2, 95, 96)
    c.add(4, 10, None)
    heading(d, '2.4 SWOT Analysis of the Product')
    c.add(2, 97, 99)
    heading(d, '2.5 Unique Selling Point')
    for n, (title, body) in enumerate([
        ('Accessible Checking for Older Android Users', 'A guided English-language interface focuses on Singapore users aged 65 and above, with clear input choices, readable labels and a short route from submission to result.'),
        ('Claim-based Evidence and Citations', 'The pipeline identifies a checkable assertion and retrieves source passages. Users can inspect the cited evidence instead of relying solely on an unexplained model verdict.'),
        ('Readable Results and Explicit Uncertainty', 'Concern labels, a plain-language explanation and a Not Enough Information outcome communicate the limits of an assessment. The current risk score is a heuristic indicator, not a calibrated probability.'),
        ('Separate Factual and Image-authenticity Checks', 'The planned premium image functions distinguish caption-image context from AI generation or manipulation. These functions remain required development work and are not part of the completed text flow.'),
    ], 1):
        heading(d, f'2.5.{n} {title}', 3); text(d, body)
    text(d, 'These are proposed product differentiators. Market superiority, improved user outcomes and formal partnerships have not yet been established through independent evaluation.')
    heading(d, '2.6 Target Users')
    c.add(2, 18, 19)
    for n, title, body in [
        (1, 'Unregistered User', 'An unregistered visitor can create an individual account or use the login and password-recovery screens. Analysis and private result history require an authenticated, verified, active account. This is an access state, not an additional operational role.'),
        (2, 'Registered Free User', 'A free individual user receives one successful check per Singapore calendar day. The planned tier supports text and public webpage links. The current prototype supports pasted text, cited results, private history and an allowance display; direct-link analysis and several account/feedback functions remain unfinished.'),
        (3, 'Registered Premium User', 'A premium individual user receives up to 60 successful checks per Singapore calendar month. The planned tier adds image with caption, OCR correction, caption-image context and static-image deepfake analysis. Current project-granted Premium changes the allowance and UI entitlement; it does not implement a payment or subscription lifecycle.'),
        (4, 'System Administrator', 'The planned administrator role queries and views accounts, creates operational accounts, suspends or restores access, and reviews user feedback. The operational portal is not implemented.'),
        (5, 'Data Engineer', 'The planned data-engineering role investigates reported incorrect results and monitors evidence ingestion and data quality. The current request-time retrieval service does not constitute a completed ingestion-monitoring portal.'),
    ]:
        heading(d, f'2.6.{n} {title}', 3); text(d, body)
    heading(d, '2.7 Business Model')
    c.add(2, 20, 21)
    text(d, 'The following model describes the proposed product. Payment collection, renewals, cancellation and organisational API licensing have not been delivered.')
    for title, start, end in [
        ('2.7.1 Target Audience', 86, 89), ('2.7.2 Marketing and Distribution Channels', 66, 74),
        ('2.7.3 Partnerships', 22, 29), ('2.7.4 Revenue Streams', 75, 81),
        ('2.7.5 Key Activities', 30, 39), ('2.7.6 Key Resources', 40, 48),
        ('2.7.7 Cost Structure', 82, 85),
    ]:
        heading(d, title, 3); c.add(2, start, end)
    return c


def stakeholders():
    c = Chapter(3); d = c.d
    text(d, 'UIABO’s stakeholders include its intended users, the development and operational team, academic reviewers and external service or evidence providers. Their needs influence accessibility, access control, evidence quality, project scope and acceptance decisions.')
    heading(d, '3.1 Stakeholder Interests and Responsibilities')
    c.add(2, 106, 107)
    heading(d, '3.2 Project Team Roles and Responsibilities')
    c.add(11, 1, 2)
    heading(d, '3.3 Sprint 1 Component Ownership')
    c.add(11, 3, None)
    text(d, 'Mr Liaw Chun Huei is the recorded supervisor and Mr Terrence Chew is the recorded assessor. Academic review and approval are separate from evidence that a prototype function runs. The communication and escalation plan is presented in Section 5.5.')
    return c


def data_collection():
    c = Chapter(4); d = c.d
    text(d, 'UIABO collects submitted content and claim-specific evidence to produce a traceable assessment. The current pipeline retrieves evidence at request time using existing cloud models. A completed proprietary training dataset and a scheduled ingestion portal are not yet available.')
    c.add(14, 2, None, prefix=('14.', '4.'))
    renames = {'4.1 Data sources': '4.1 Data Sources', '4.2 Types of data and provenance': '4.2 Types of Data', '4.3 Collection methods': '4.3 Collection Methods', '4.4 Preprocessing and quality controls': '4.4 Data Preprocessing', '4.5 Privacy, access and data lifecycle': '4.5 Privacy Considerations', '4.6 Evaluation dataset plan': '4.6 Evaluation Dataset Plan'}
    for p in d.paragraphs:
        if p.text in renames: p.text = renames[p.text]
    return c


def timeline():
    c = Chapter(5); d = c.d
    heading(d, '5.1 Project Milestones'); c.add(10, 4, 5)
    heading(d, '5.2 Gantt Chart'); c.add(10, 1, 3)
    heading(d, '5.3 Work Breakdown Structure (WBS)'); c.add(10, 6, 7)
    for n, role, body in [
        (1, 'Unregistered User', 'Registration, email verification, sign-in and password recovery; validate error handling and access before analysis.'),
        (2, 'Registered Free User', 'Complete text submission, evidence retrieval, result presentation, history and daily allowance; then implement agreed direct-link, account and feedback gaps.'),
        (3, 'Registered Premium User', 'Enforce monthly entitlement; deliver and evaluate image, OCR, context and deepfake functions; implement the agreed subscription lifecycle.'),
        (4, 'System Administrator', 'Deliver role-restricted account queries, account creation, suspension controls and feedback review with access and audit checks.'),
        (5, 'Data Engineer', 'Deliver incorrect-result investigation, source/provenance management and ingestion monitoring, supported by labelled evaluation and recovery records.'),
    ]:
        heading(d, f'5.3.{n} {role}', 3); text(d, body)
    text(d, 'These role-based work packages elaborate the existing requirements. They do not imply that all packages are complete or that new owners or dates have been agreed.')
    heading(d, '5.4 Project Charter'); c.add(16, 1, 3)
    heading(d, '5.5 Communication Management Plan')
    c.add(9, 2, 7)
    heading(d, '5.5.1 Stakeholder Communication Requirements', 3)
    text(d, 'Team members need current contracts, task ownership and blockers. The supervisor and assessor need reviewable requirements, demonstrations and progress evidence. User-facing communication must explain supported functions, uncertainty and data handling in accessible language.')
    heading(d, '5.5.2 Communications Summary', 3); c.add(9, 11, 13)
    heading(d, '5.5.3 Communication Guidelines', 3)
    text(d, 'Use a dated record that identifies the decision or issue, affected component, owner and next action. Keep implementation claims tied to test or demonstration evidence. Share configuration instructions without provider keys or private service-account files.')
    heading(d, '5.5.4 Conflict Escalation Procedure', 3); c.add(9, 14, 17)
    heading(d, '5.5.5 Document Revision Protocol', 3); c.add(9, 17, 18)
    text(d, 'Preserve the previous document, record the reason and source for each revision, reconcile section references and refresh the Word contents and page fields. This revision updates the Word report only; the existing PDF retains the previous structure.')
    heading(d, '5.6 Scope Statement'); c.add(16, 4, 5)
    heading(d, '5.6.1 Deliverables and Dependencies', 3); c.add(16, 6, 10)
    heading(d, '5.6.2 Acceptance and Outstanding Approval', 3); c.add(16, 11, None)
    heading(d, '5.7 Proposed Next Work Order'); c.add(10, 8, None)
    return c


def requirements():
    c = Chapter(6); d = c.d
    heading(d, '6.1 Stakeholder Identification'); c.add(15, 2, 3)
    heading(d, '6.2 Requirement Gathering Methods')
    table(d, ['Method', 'Available evidence and limitation'], [
        ('Document analysis', 'PRD scope and product research, URS requirements/use cases and TDM design.'),
        ('Team discussion and scope clarification', 'Available meeting records and confirmed September prototype scope; deepfake required and audio excluded.'),
        ('Prototype inspection and testing', 'Dated API, pipeline and Android/Firestore records identify integration and usability questions.'),
        ('End-user research', 'No completed interview, survey or focus-group records are supplied. Older-user research and consent materials remain to be prepared.'),
    ])
    heading(d, '6.3 Requirement Analysis Approach'); c.add(15, 4, 6)
    heading(d, '6.4 Summary of User and System Needs'); c.add(15, 7, 8)
    heading(d, '6.5 Use of Tools')
    text(d, 'The project uses shared Word/PDF requirements and design documents, diagrams, GitHub version history, JSON handoff examples, automated backend tests and the Android emulator. These tools support traceability and review. No completed survey-tool deployment or automated requirements-management system is claimed.')
    heading(d, '6.6 Requirement-to-Evidence Traceability'); c.add(15, 9, 10)
    heading(d, '6.7 Acceptance and Change Control'); c.add(15, 11, None)
    return c


def functional():
    c = Chapter(7); d = c.d
    text(d, 'The following requirements describe the intended product. The URS distinguishes base requirements, additional requirements and stretch goals. A base requirement can remain unfinished in the September text increment; implementation status is recorded separately in Section 7.5.')
    heading(d, '7.1 Functional Hierarchy')
    table(d, ['Module', 'Principal functions'], [
        ('User management', 'Registration, login/logout, password recovery, profile and account deletion.'),
        ('Content submission and analysis', 'Text, direct links and premium image/OCR/context/deepfake workflows.'),
        ('Evidence and results', 'Claim extraction, retrieval, assessment, uncertainty, citations, history and sharing.'),
        ('Entitlement and subscriptions', 'Free daily and premium monthly allowances; upgrade, renewal and cancellation.'),
        ('Feedback and administration', 'Incorrect-result reports, reviews and authorised account operations.'),
        ('Data engineering', 'Source quality, reported-result investigation and ingestion monitoring.'),
    ])
    heading(d, '7.1.1 Detailed Functional Requirements', 3); c.add(5, 3, 99)
    heading(d, '7.2 Basic Feature Access Levels')
    table(d, ['Capability', 'Free', 'Premium', 'Operational roles'], [
        ('Account/profile access', 'Own account', 'Own account', 'Assigned role and authorised records'),
        ('Text and webpage-link analysis', '1 successful check/day', 'Up to 60 successful checks/month', 'Not the operational portal purpose'),
        ('Image, OCR, context and deepfake', 'Unavailable', 'Planned premium functions', 'Investigation/maintenance as authorised'),
        ('Result history and evidence', 'Own results', 'Own results', 'Reported records as authorised; portal planned'),
        ('Account management', 'Own account only', 'Own account only', 'Administrator functions planned'),
        ('Ingestion monitoring', 'Unavailable', 'Unavailable', 'Data Engineer portal planned'),
    ])
    heading(d, '7.3 Dependencies')
    c.add(6, 61, 62)
    text(d, 'The current text path is input preparation → claim analysis → evidence retrieval → evidence assessment → authenticated result persistence. Downstream stages depend on validated upstream contracts. Image functions additionally require selected and evaluated OCR, context and authenticity components; billing requires a selected payment integration.')
    heading(d, '7.4 Inputs and Outputs')
    table(d, ['Boundary', 'Input', 'Output'], [
        ('Preparation', 'English text, up to 5,000 characters', 'PreparedText with preserved original text, normalized content and warnings'),
        ('Claim analysis', 'PreparedText', 'ClaimAnalysis with extracted claim, category, checkability and agreement'),
        ('Evidence retrieval', 'Checkable claim', 'RetrievalResult with source passages, URLs, dates where known and warnings'),
        ('Evidence assessment', 'Claim and retrieved evidence', 'AssessmentResult with stance, concern, risk indicator, uncertainty and explanation'),
        ('Application result', 'Authenticated submission and idempotency key', 'Owned TextAnalysisResult, saved evidence and updated allowance'),
        ('Planned image path', 'Supported static image and optional caption/OCR corrections', 'Content/context assessment and a separate image-generation/manipulation likelihood where applicable'),
    ])
    heading(d, '7.5 Current Implementation Status'); c.add(1, 50, None)
    return c


def nonfunctional():
    c = Chapter(8); d = c.d
    c.add(5, 155, 156)
    heading(d, '8.1 Performance Requirements')
    c.add(5, 162, 168, headings={'5.2.2 Load Time': ('8.1.1 Load Time and Interaction Latency', 3)})
    c.add(5, 170, 178, headings={'5.2.4 Result Generation Time': ('8.1.2 Result Generation Time', 3)})
    text(d, 'The current implementation has tighter stage and client time bounds described in Chapter 12. These configured limits do not establish measured response-time guarantees or completion of the URS load targets.')
    heading(d, '8.2 Security Requirements'); c.add(5, 179, 185)
    text(d, 'Current controls include backend verification of Firebase ID tokens, active and verified individual-account checks, server-derived result ownership, private provider credentials and atomic allowance updates. Thirty-minute inactivity expiry and the complete operational-role control model still need validation or implementation.')
    heading(d, '8.3 Interface and Usability Requirements')
    c.add(2, 9, 10)
    text(d, 'UIABO should use readable text, high contrast, large touch targets, explicit input labels and clear unavailable-feature states. Evidence and uncertainty must be understandable to the intended older users. The Android emulator demonstrates the current layout; accessibility conformance and older-user comprehension remain to be evaluated.')
    heading(d, '8.4 Portability and Scalability Requirements')
    c.add(5, 157, 162)
    text(d, 'The agreed initial platform is Android and the working demonstration uses an emulator. Physical-device compatibility, a release build and a capacity test are still needed. The current owned-history query reads and sorts matching records in the backend, which requires optimization as histories grow.')
    heading(d, '8.5 Database Storage Requirements'); c.add(5, 169, 170)
    text(d, 'A defined personal storage allocation and its warning behavior remain outstanding. The report does not infer that the proposed 80% warning is already implemented.')
    source(d, 'PRD non-functional requirements; URS Chapter 4; dated verification and remaining evaluation in Section 19.1.')
    return c


def other_requirements():
    c = Chapter(9); d = c.d
    heading(d, '9.1 User Requirements')
    text(d, 'Individual users need an Android device and an internet connection to access the cloud-assisted text workflow. The prototype supports English. Users should be able to inspect the extracted claim, explanation, uncertainty and citations before deciding whether to share content. Not Enough Information must not be presented as proof that a claim is false.')
    text(d, 'The final product must explain the distinction between factual assessment and image-authenticity likelihood. A supported static-image detector is required; audio and video inputs are outside the agreed scope. User instructions, supported input limits and unavailable functions must remain consistent with the PUM.')
    heading(d, '9.2 Legal and Regulatory Requirements')
    text(d, 'The project still needs a reviewed privacy notice, terms of use, retention and deletion rules, and a documented review of external provider processing and source-reuse conditions. These are unresolved project deliverables; this report does not assert legal approval or compliance certification. Public availability of evidence does not itself grant unrestricted reuse rights.')
    text(d, 'Only content that users are permitted to submit should be processed. Evaluation and demonstrations should use appropriate source permissions and non-sensitive examples. Any participant study requires its consent and data-handling arrangements to be recorded before results are presented as research evidence.')
    heading(d, '9.3 System and Operation Requirements')
    text(d, 'The system depends on Firebase Authentication and Firestore, an available backend, and configured AI/search providers. Provider failures must produce controlled errors or clearly qualified partial evidence. They must not be disguised as a confident result or a successful empty search. Credentials belong in private backend configuration, outside version control and the Android bundle.')
    text(d, 'Production HTTPS, environment separation, monitoring, backup/recovery, operational portals, source-ingestion scheduling and release packaging remain to be completed or verified. The current local demonstration is not evidence of a production deployment.')
    source(d, 'URS operating context and constraints; existing PTD data lifecycle, risk and deployment records.')
    return c


def risks():
    c = Chapter(10)
    c.add(7, 1, None, headings={'7.1 Current risk treatment': ('10.1 Current Risk Treatment', 2)})
    return c


def methodology():
    c = Chapter(11)
    c.add(6, 2, 35, headings={'6.1.1 Development Methodology Chosen: Scrum': ('11.5 Development Methodology Chosen: Scrum', 2)})
    names = {'Waterfall Model': '11.1 Waterfall Model', 'Prototyping Model': '11.2 Prototyping Model', 'Kanban Model': '11.3 Kanban Model', 'Scrum': '11.4 Scrum'}
    for p in c.d.paragraphs:
        if p.text in names:
            p.text = names[p.text]; p.style = c.d.styles['Heading 2']
        if 'AI-authorship scoring' in p.text:
            p.text = p.text.replace('scam detection and AI-authorship scoring', 'URL screening and static-image authenticity assessment')
    return c


def stack():
    c = Chapter(12); d = c.d
    text(d, 'The current prototype uses JavaScript/React Native with Expo, Python/FastAPI/Uvicorn, Firebase Authentication and Cloud Firestore. The tables retain the reasons for choosing the stack. Planned link, image, payment and portal capabilities must not be inferred from a framework’s general capabilities.')
    heading(d, '12.1 Frontend Framework'); c.add(6, 39, 41)
    heading(d, '12.2 Backend Framework'); c.add(6, 42, 44)
    heading(d, '12.3 Database'); c.add(6, 45, 47)
    text(d, 'Current protected account/result operations go through FastAPI and the backend Admin SDK. Only the implemented collections should be treated as operational; Section 16.3 distinguishes proposed schemas.')
    heading(d, '12.4 Machine Learning Libraries and Analysis Components')
    text(d, 'The claim stage currently calls pretrained Ollama Cloud models rather than training a local misinformation model. It uses gpt-oss:120b, gemma4:31b and nemotron-3-super for classification, with gemma4:31b for claim extraction. OCR and static-image deepfake libraries or services still need selection and evaluation.')
    c.add(6, 54, 59, headings={'6.3.1 Model voting, validation and prompt handling': ('12.4.1 Model Voting, Response Validation and Prompt Handling', 3)})
    heading(d, '12.5 APIs'); c.add(6, 48, 50)
    heading(d, '12.5.1 Authentication', 3); c.add(6, 51, 53)
    heading(d, '12.5.2 Backend Endpoints and Provider Interfaces', 3); c.add(6, 60, 64)
    heading(d, '12.6 Hosting'); c.add(6, 65, 66)
    heading(d, '12.7 Server')
    text(d, 'Uvicorn runs the FastAPI application on the development computer. The Android emulator reaches the host through 10.0.2.2:8000. Firebase and the AI/search services remain external cloud dependencies; there is no verified public backend endpoint or completed production server configuration.')
    heading(d, '12.8 Deployment'); c.add(6, 66, 75)
    heading(d, '12.8.1 Release Work Still Required', 3); c.add(6, 76, None)
    return c


def clean_pdf_text(value):
    return re.sub(r'\s+', ' ', value).strip()


def stories():
    c = Chapter(13); d = c.d
    text(d, 'The following user stories are transcribed from the submitted UIABO URS. UC identifiers are unchanged. Free and premium users share UC-01–UC-15; operational accounts use their own role-specific use cases. The stories describe required behavior, not a claim that each function is complete.')
    heading(d, '13.1 Unregistered User')
    text(d, 'The URS groups registration and password recovery under the destination account roles rather than defining a separate unregistered-user story set. UC-16 covers free registration; UC-18 covers planned premium registration; UC-04 covers individual password recovery. Existing sign-in stories lead to the relevant authenticated role.')
    count = 0
    for n, label, page in [(2, 'Registered Free User', 15), (3, 'Registered Premium User', 34), (4, 'System Administrator', 44), (5, 'Data Engineer', 57)]:
        heading(d, f'13.{n} {label}')
        raw = PDF.pages[page-1].extract_text()
        raw = re.split(r'3\.\d\.2 Use Case Description', raw)[0]
        found = re.findall(r'•\s*(UC-\d{2}:.*?)(?=•|\Z)', raw, re.S)
        assert found, page
        for item in found:
            text(d, clean_pdf_text(item).replace('data- engineering', 'data-engineering'), 'List Bullet')
        count += len(found)
        source(d, f'URS, page {page}; original use-case identifiers retained.')
    assert count == 58, count
    return c


CASE_FIELDS = ['Name:', 'ID:', 'Stakeholders & Goals:', 'Description:', 'Actor(s):', 'Trigger:', 'Pre-condition:', 'Main Flow:', 'Sub-flow:', 'Alternative flow:']


def case_records():
    records = {}
    pattern = re.compile('|'.join(re.escape(s) for s in CASE_FIELDS))
    for page, obj in enumerate(PDF.pages, 1):
        raw = obj.extract_text() or ''
        m = re.search(r'Use Case (\d{2})\s*[–-]\s*(.*?)\nName:', raw, re.S)
        if not m:
            continue
        number = int(m[1])
        body = raw[raw.index('Name:', m.start()):].strip()
        labels = list(pattern.finditer(body))
        assert [m[0] for m in labels] == CASE_FIELDS, (page, [m[0] for m in labels])
        rows = []
        for i, field in enumerate(labels):
            value = body[field.end(): labels[i+1].start() if i+1 < len(labels) else len(body)]
            value = clean_pdf_text(value)
            # Restore logical steps, not the PDF's physical line wraps.
            value = re.sub(r'\s+(?=(?:\d+[a-z]?\.|S\d+\.)\s)', '\n', value)
            rows.append((field[0][:-1], value))
        records[number] = (clean_pdf_text(m[2]), rows, page)
    assert set(records) == set(range(1, 44)), records.keys()
    return records


def usecases():
    c = Chapter(14); d = c.d
    text(d, 'These editable descriptions retain all 43 use cases from the submitted URS, including their preconditions, main flows, sub-flows and alternatives. They specify intended behavior. Current completion and scope differences are documented in Sections 7.5 and 19.1; a requirement such as direct-link analysis or a 20-minute cancellation rule is not evidence that it is implemented.')
    records = case_records()
    groups = [(1, 'Shared Across Individual Users', range(1,16)), (2, 'Registered Free User', range(16,18)), (3, 'Registered Premium User', range(18,25)), (4, 'System Administrator', range(25,36)), (5, 'Data Engineer', range(36,44))]
    for group, label, numbers in groups:
        group_heading = heading(d, f'14.{group} {label}')
        group_heading.paragraph_format.page_break_before = group > 1
        for offset, number in enumerate(numbers, 1):
            name, rows, page = records[number]
            h = heading(d, f'14.{group}.{offset} UC-{number:02d} — {name}', 3)
            h.paragraph_format.page_break_before = offset > 1
            t = table(d, ['Use Case Field', 'Description'], rows)
            t.autofit = False
            for row in t.rows:
                row.cells[0].width = Cm(3.1); row.cells[1].width = Cm(13.7)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.line_spacing = 1.0
                        p.paragraph_format.space_after = Pt(3)
                        for r in p.runs: r.font.size = Pt(9.5)
            source(d, f'URS, page {page}; UC-{number:02d}.')
    return c


def diagrams():
    c = Chapter(15); d = c.d
    text(d, 'The four use-case diagrams are retained from the UIABO URS. Their original figure identifiers belong to that source document. They describe the planned role boundaries and complement the editable stories and descriptions in Chapters 13 and 14. Unregistered access is represented through registration and recovery, without inventing an additional role diagram.')
    for n, title, page in [(1, 'Registered Free User', 14), (2, 'Registered Premium User', 33), (3, 'System Administrator', 43), (4, 'Data Engineer', 56)]:
        heading(d, f'15.{n} {title}')
        index = 101 + page - 13
        c.add(5, index, index+1)
        source(d, f'URS, page {page}; source diagram retained.')
    return c


def design():
    c = Chapter(16); d = c.d
    text(d, 'This chapter retains UIABO’s TDM system designs and adds the current Android and persistence records. Proposed diagrams, database collections and wireframes are separated from observed implementation. Example screen values in the source designs are illustrative.')
    heading(d, '16.1 Data Flow Diagram'); c.add(5, 346, 347)
    heading(d, '16.2 System Architecture Design'); c.add(5, 186, 187)
    heading(d, '16.3 Database Design'); c.add(5, 291, 345, prefix=('5.5.', '16.3.'))
    heading(d, '16.4 Wireframe Design'); c.add(5, 348, 375, prefix=('5.7.', '16.4.'))
    heading(d, '16.4.6 Current Android Implementation', 3); c.add(5, 376, 384)
    heading(d, '16.5 Sequence Diagrams'); c.add(5, 189, 281, prefix=('5.4.1.', '16.5.'))
    heading(d, '16.6 Activity Diagrams'); c.add(5, 282, 290, prefix=('5.4.2.', '16.6.'))
    heading(d, '16.7 Current Authentication, Persistence and Retry Design'); c.add(5, 385, None)
    return c


def conclusion():
    c = Chapter(17); d = c.d
    c.add(18, 1, 3)
    heading(d, '17.1 Remaining Deliverables and Acceptance Evidence'); c.add(18, 4, 5)
    text(d, 'The report follows the sample’s progression from project rationale through requirements to design. Supporting test records, actual meeting notes and references are collected in Chapter 19. PTD_MISSING_INFORMATION.md remains the companion checklist for unresolved information and implementation work.')
    return c


def glossary():
    c = Chapter(18); c.add(17, 1, None); return c


def appendix():
    c = Chapter(19); d = c.d
    text(d, 'The appendix preserves supporting progress evidence, meeting records and references. These records remain dated; reformatting the report does not create new test results, meeting approvals or completion claims.')
    heading(d, '19.1 System Test Summary'); c.add(8, 1, None, prefix=('8.', '19.1.'))
    heading(d, '19.2 Meeting Minutes'); c.add(12, 1, None, prefix=('12.', '19.2.'))
    heading(d, '19.3 References'); c.add(13, 1, None, prefix=('13.', '19.3.'))
    heading(d, '19.4 Project Website and Submission Status'); c.add(9, 8, 9)
    text(d, 'This Word revision is dated 8 September 2026 and follows the supplied sample’s report structure. The previous PDF has intentionally not been regenerated. The structural crosswalk is in PTD_UPDATE_SUMMARY.md, and remaining evidence is recorded in PTD_MISSING_INFORMATION.md.')
    return c


def add_field(p, command):
    r = p.add_run()
    for kind in ('begin', 'separate', 'end'):
        if kind == 'separate':
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = command; r._r.append(e)
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind); r._r.append(e)


def cover():
    d = fragment(0, 0, 11)
    for p in d.paragraphs:
        if p.text.startswith('Version '): p.text = f'Version 0.3 — Working draft — {DATE}'
        if p.text and not p._p.xpath('.//w:drawing'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()
    p = text(d, 'Page of Contents'); p.style = d.styles['Title']
    add_field(d.add_paragraph(), ' TOC \\o "1-3" \\h \\z \\u ')
    d.add_page_break()
    p = text(d, 'Document Control'); p.style = d.styles['Heading 1']
    p.paragraph_format.page_break_before = False
    table(d, ['Field', 'Record'], [
        ('Title', 'Preliminary Technical Documentation'),
        ('Document name', 'FYP-26-S3-30 Preliminary Technical Documentation'),
        ('Project', 'uiabo — AI-Based Misinformation Detection for Short-Form Content'),
        ('Document owner', 'FYP-26-S3-30 project team'),
        ('Record of revision', '0.3 — 8 September 2026 — working draft for team review'),
        ('Evidence baseline', 'Existing PRD/URS/TDM and dated implementation/test records through 7 September 2026'),
    ])
    heading(d, 'Record of Revision', 2)
    table(d, ['Date', 'Version', 'Description', 'Section affected'], [
        ('4 September 2026', 'Initial draft', 'Existing UIABO section drafts.', 'Original section files'),
        ('7 September 2026', '0.2', 'Compiled UIABO report and added current pipeline, Android, Firestore and missing coverage.', 'Whole report'),
        ('8 September 2026', '0.3', 'Reordered chapters and subsections to follow the supplied sample; converted stories and 43 use cases to editable text/tables; refreshed Word contents.', 'Whole Word report'),
    ])
    text(d, 'Revision ownership and individual contribution attribution are subject to team review; no approval or signature is inferred. Implemented features, planned requirements and pending evidence are distinguished throughout. The exact November deadline and other missing information remain in the companion checklist.')
    return d


def format_document(d, is_master=False):
    normal = d.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.widow_control = True
    for level in range(1, 5):
        s = d.styles[f'Heading {level}']; s.font.name = 'Arial'; s.font.bold = True
        s.font.size = Pt(15 if level == 1 else 12 if level == 2 else 11)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(6)
    for p in d.paragraphs:
        if p.style.name.startswith('Heading'):
            if p.style.name == 'Heading 1' and re.match(r'^\d+\.', p.text):
                p.paragraph_format.page_break_before = True
            for run in p.runs:
                run.font.name = 'Arial'; run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.bold = True; run.font.size = d.styles[p.style.name].font.size
        elif p.text and p.style.name == 'Normal' and p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if p.style.name == 'Frontmatter Heading':
            p.style = d.styles['Normal']
            p.paragraph_format.keep_with_next = True
            for run in p.runs: run.bold = True
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(3)
        if not p.text and not p._p.xpath('.//w:drawing|.//w:fldChar|.//w:br'):
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)
            for run in p.runs: run.font.size = Pt(1)
    for t in d.tables:
        for i, row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = False
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.name = 'Arial'
                        if run.font.size is None or run.font.size.pt > 10:
                            run.font.size = Pt(10)
                        if i == 0: run.bold = True
    # Keep retained diagrams within a page; never stretch them.
    for image in d.inline_shapes:
        scale = min(1, Cm(16.8)/image.width, Cm(21.0)/image.height)
        if scale < 1:
            image.width = int(image.width*scale); image.height = int(image.height*scale)
    for section in d.sections:
        section.page_width = Cm(21); section.page_height = Cm(29.7)
        section.top_margin = Cm(2); section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.1); section.right_margin = Cm(2.1)
        section.header_distance = Cm(.8); section.footer_distance = Cm(.8)
        section.different_first_page_header_footer = bool(is_master)
        hp = section.header.paragraphs[0]
        hp.text = 'FYP-26-S3-30  |  uiabo  |  Preliminary Technical Documentation'
        for run in hp.runs: run.font.name = 'Arial'; run.font.size = Pt(8)
        fp = section.footer.paragraphs[0]; fp.clear(); fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(fp, ' PAGE ')
    d.core_properties.title = 'UIABO Preliminary Technical Documentation'
    d.core_properties.subject = 'FYP-26-S3-30 — sample-aligned report structure'
    d.core_properties.author = 'FYP-26-S3-30 project team'
    d.core_properties.comments = 'Version 0.3, 8 September 2026. Word-only structural revision.'


def main():
    OUT.mkdir(exist_ok=True)
    builders = [introduction, overview, stakeholders, data_collection, timeline,
                requirements, functional, nonfunctional, other_requirements, risks,
                methodology, stack, stories, usecases, diagrams, design, conclusion,
                glossary, appendix]
    front = cover(); format_document(front, True)
    front.save(OUT/'00_Cover_Contents_and_Document_Control.docx')
    master = Document(OUT/'00_Cover_Contents_and_Document_Control.docx')
    composer = Composer(master)
    for number, builder in enumerate(builders, 1):
        c = builder(); format_document(c.d)
        name = f'{number:02d}_{CHAPTERS[number-1].replace(" ", "_").replace("-", "_")}.docx'
        c.composer.save(OUT/name)
        composer.append(Document(OUT/name))
        print('Reformatted', name)
    format_document(master, True)
    composer.save(MASTER)
    print('Saved', MASTER.name, 'with', len(master.inline_shapes), 'inline figures.')


if __name__ == '__main__':
    main()
