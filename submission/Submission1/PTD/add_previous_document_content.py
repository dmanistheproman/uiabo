"""Add reusable PRD/URS/TDM content to the current UIABO PTD.

The script edits only the consolidated Word document. It does not generate a
PDF. Content imported from planning documents is labelled as planned so it is
not mistaken for implementation evidence.
"""

from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
import re
import shutil
import tempfile

import fitz
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pypdf import PdfReader


SCRIPT_ROOT = Path(__file__).resolve().parent
WORKSPACE = SCRIPT_ROOT.parents[3]
ROOT = WORKSPACE / "submission/Submission1/PTD"
DOCX_PATH = ROOT / "FYP-26-S3-30_Preliminary_Technical_Documentation.docx"
URS_PDF = WORKSPACE / "submission/URS/FYP-26-S3-30_URS.pdf"
PRD_DOCX = WORKSPACE / "submission/PRD/FYP-26-S3-30_PRD.docx"
TDM_ROOT = WORKSPACE / "TDM_Diagrams"
UI_ROOT = WORKSPACE / "app"
BACKUP_ROOT = ROOT / "backups"


CASE_FIELDS = [
    "Name:",
    "ID:",
    "Stakeholders & Goals:",
    "Description:",
    "Actor(s):",
    "Trigger:",
    "Pre-condition:",
    "Main Flow:",
    "Sub-flow:",
    "Alternative flow:",
]

IMPLEMENTED_CASES = {1, 2, 4, 5, 6, 7, 9, 10, 11, 16}
PARTIAL_CASES = {15}


def find_paragraph(document: Document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def optional_paragraph(document: Document, text_prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(text_prefix):
            return paragraph
    return None


def clear_between(start, end) -> None:
    node = start._p.getnext()
    while node is not None and node is not end._p:
        following = node.getnext()
        node.getparent().remove(node)
        node = following


def remove_from_to(start, end) -> None:
    node = start._p
    while node is not None and node is not end._p:
        following = node.getnext()
        node.getparent().remove(node)
        node = following


def move_before(marker, element) -> None:
    marker._p.addprevious(element)


def add_paragraph_before(document, marker, text="", style=None):
    paragraph = document.add_paragraph(text, style=style)
    move_before(marker, paragraph._p)
    return paragraph


def add_heading_before(document, marker, text, level=2):
    paragraph = add_paragraph_before(document, marker, text, f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_bullets_before(document, marker, values) -> None:
    for value in values:
        add_paragraph_before(document, marker, value, "List Bullet")


def set_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def add_table_before(document, marker, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
    for row in table.rows:
        set_cant_split(row)
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                cell.width = Cm(widths[index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5)
    move_before(marker, table._tbl)
    add_paragraph_before(document, marker, "")
    return table


def add_gap_box(document, marker, code, label) -> None:
    add_paragraph_before(document, marker, f"{code} — {label}")
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.cell(0, 0).text = ""
    table.rows[0].height = Cm(2.1)
    move_before(marker, table._tbl)
    add_paragraph_before(document, marker, "")


def add_picture_before(document, marker, path: Path, caption: str, width=16.0) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width))
    move_before(marker, paragraph._p)
    caption_paragraph = add_paragraph_before(document, marker, caption, "Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_picture_grid(document, marker, entries) -> None:
    rows = (len(entries) + 1) // 2
    table = document.add_table(rows=rows, cols=2)
    table.autofit = False
    for row in table.rows:
        set_cant_split(row)
    for index, (path, caption) in enumerate(entries):
        cell = table.cell(index // 2, index % 2)
        cell.width = Cm(8.2)
        picture_paragraph = cell.paragraphs[0]
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.add_run().add_picture(str(path), width=Cm(7.35))
        caption_paragraph = cell.add_paragraph(caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption_paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
    move_before(marker, table._tbl)
    add_paragraph_before(document, marker, "")


def clean_pdf_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def parse_use_cases() -> dict[int, tuple[str, list[tuple[str, str]], int]]:
    reader = PdfReader(URS_PDF)
    records = {}
    field_pattern = re.compile("|".join(re.escape(item) for item in CASE_FIELDS))
    for page_number, page in enumerate(reader.pages, 1):
        raw = page.extract_text() or ""
        match = re.search(r"Use Case (\d{2})\s*[–-]\s*(.*?)\nName:", raw, re.S)
        if not match:
            continue
        number = int(match.group(1))
        body = raw[raw.index("Name:", match.start()) :].strip()
        labels = list(field_pattern.finditer(body))
        if [item.group(0) for item in labels] != CASE_FIELDS:
            raise ValueError(f"Unexpected use-case layout on URS page {page_number}")
        rows = []
        for index, field in enumerate(labels):
            end = labels[index + 1].start() if index + 1 < len(labels) else len(body)
            value = clean_pdf_text(body[field.end() : end])
            value = re.sub(r"\s+(?=(?:\d+[a-z]?\.|S\d+\.)\s)", "\n", value)
            rows.append((field.group(0)[:-1], value))
        records[number] = (clean_pdf_text(match.group(2)), rows, page_number)
    if set(records) != set(range(1, 44)):
        raise ValueError(f"Expected UC01–UC43; found {sorted(records)}")
    return records


def parse_user_stories() -> list[tuple[str, list[str], int]]:
    reader = PdfReader(URS_PDF)
    groups = [
        ("Free User", 15),
        ("Premium User", 34),
        ("System Administrator", 44),
        ("Data Engineer", 57),
    ]
    output = []
    for label, page_number in groups:
        raw = reader.pages[page_number - 1].extract_text() or ""
        raw = re.split(r"3\.\d\.2 Use Case Description", raw)[0]
        stories = [
            clean_pdf_text(item)
            for item in re.findall(r"•\s*(UC-\d{2}:.*?)(?=•|\Z)", raw, re.S)
        ]
        if not stories:
            raise ValueError(f"No stories extracted from URS page {page_number}")
        output.append((label, stories, page_number))
    return output


def case_status(number: int) -> str:
    if number in IMPLEMENTED_CASES:
        return "Implemented in the current Android text prototype"
    if number in PARTIAL_CASES:
        return "Partially implemented: native summary sharing; no public result link"
    return "Planned requirement from the URS; not yet implemented"


def replace_introduction(document: Document) -> None:
    purpose = find_paragraph(document, "1.2 Purpose of the Project")
    vision = find_paragraph(document, "1.3 Product Vision")
    clear_between(purpose, vision)
    add_paragraph_before(
        document,
        vision,
        "This Preliminary Technical Documentation describes UIABO's project rationale, "
        "requirements, planning, technical choices and system design. It also records the "
        "working Android text-checking prototype, its supporting evidence and the remaining "
        "work towards the complete approved system.",
    )
    add_paragraph_before(
        document,
        vision,
        "The project itself aims to help older Singapore users assess suspicious short-form "
        "content through readable explanations, explicit uncertainty and inspectable evidence.",
    )

    objectives = find_paragraph(document, "1.4 Implemented Prototype Objectives")
    learning = find_paragraph(document, "1.5 Learning Objectives")
    objectives.text = "1.4 Project Objectives and Implemented Prototype"
    clear_between(objectives, learning)
    add_paragraph_before(
        document,
        learning,
        "The approved project scope is broader than the current implementation. The objectives are:",
    )
    add_bullets_before(
        document,
        learning,
        [
            "Provide an elderly-friendly Android application for English short-form text, public webpage links and supported static images.",
            "Give Free users one successful text or link check per Singapore day and Premium users up to 60 successful checks per month.",
            "Identify factual claims, retrieve supporting or contradicting evidence, communicate uncertainty and provide citations.",
            "Provide Premium image upload, OCR correction, caption-image context checking and a separate probabilistic image-manipulation assessment.",
            "Provide role-restricted System Administrator and Data Engineer functions for account management, feedback investigation and pipeline monitoring.",
            "Evaluate accuracy, relevance, usability, security and performance without presenting an AI score as proof of truth.",
        ],
    )
    add_paragraph_before(document, learning, "The implemented September prototype currently provides:")
    add_bullets_before(
        document,
        learning,
        [
            "Verified account access, profile editing and server-controlled Free/Premium allowances.",
            "English text submission through Android to the FastAPI pipeline.",
            "Claim extraction, evidence retrieval, assessment, citations and explicit Not Enough Information handling.",
            "Owned Firestore results, history, result reopening and native summary sharing.",
        ],
    )


def replace_overview(document: Document) -> None:
    target = find_paragraph(document, "2.6 Target Users and Access Model")
    business = find_paragraph(document, "2.7 Operating Costs")
    clear_between(target, business)
    add_paragraph_before(
        document,
        business,
        "Unauthenticated access is a pre-login state rather than a separate long-term user role. "
        "Registration, login and password recovery lead to one of the approved account roles.",
    )
    add_table_before(
        document,
        business,
        ["Role", "Planned access", "Current position"],
        [
            ("Free User", "Text and public-link checking; history, sharing, reporting, review and Premium upgrade; one successful check per day.", "Text checking, history, allowance and native summary sharing are implemented."),
            ("Premium User", "All Free functions; 60 successful checks per month; image, OCR, caption-context and image-authenticity functions.", "Premium allowance exists; Premium image and payment functions are planned."),
            ("System Administrator", "Role-restricted account query, operational-account creation, suspension controls and feedback review.", "Planned; no operational portal is implemented."),
            ("Data Engineer", "Role-restricted incorrect-result dashboard, report investigation and ingestion-pipeline monitoring.", "Planned; no operational portal is implemented."),
        ],
        widths=[3.2, 7.6, 5.8],
    )

    business.text = "2.7 Business Model and Operating Costs"
    stakeholders = find_paragraph(document, "3 Stakeholders")
    clear_between(business, stakeholders)
    add_paragraph_before(
        document,
        stakeholders,
        "The PRD defines a social-impact freemium model. Free access supports occasional "
        "checking, while Premium access supports more frequent use and the planned image functions. "
        "The model remains a proposal until pricing and demand are validated.",
    )
    add_heading_before(document, stakeholders, "2.7.1 Value Proposition and Customer Segments", 3)
    add_bullets_before(
        document,
        stakeholders,
        [
            "Older Android users receive a guided evidence-checking flow with plain-language explanations and uncertainty.",
            "Free users receive one successful text or webpage-link submission per day.",
            "Premium users receive up to 60 successful submissions per month and planned image, OCR, context and authenticity functions.",
            "Future organisational users may access selected verification functions through approved API licensing.",
        ],
    )
    add_heading_before(document, stakeholders, "2.7.2 Partners, Activities and Channels", 3)
    add_table_before(
        document,
        stakeholders,
        ["Area", "Planned approach from the PRD"],
        [
            ("Potential partners", "Public agencies, fact-checking/news organisations, community and senior-support organisations, universities and research organisations. No formal partnership is asserted."),
            ("Key activities", "Maintain retrieval sources, evaluate claim/evidence behaviour, validate citations, manage accounts and allowances, test usability and protect user data."),
            ("Channels", "Android application, community demonstrations, digital-literacy outreach and possible future organisational integration."),
            ("Revenue", "Free introductory access, monthly Premium access and possible future API licensing. Earlier interface samples show SGD20, but the PRD states that the price requires validation."),
        ],
        widths=[3.5, 13.0],
    )
    add_heading_before(document, stakeholders, "2.7.3 Proposed Cost Structure", 3)
    add_bullets_before(
        document,
        stakeholders,
        [
            "Cloud hosting, database, authentication and storage usage.",
            "AI inference, fact-check search, web search/extraction, OCR and image-analysis services.",
            "Android development, testing devices and release/distribution costs.",
            "Dataset preparation, source maintenance, security review and usability evaluation.",
        ],
    )
    add_heading_before(document, stakeholders, "2.7.4 Actual Project Expenditure", 3)
    add_gap_box(document, stakeholders, "G03", "Actual project expenditure")


def extract_prd_gantt(temp_dir: Path) -> list[Path]:
    prd = Document(PRD_DOCX)
    heading = find_paragraph(prd, "4.5 Gantt Chart")
    paragraph = heading._p.getnext()
    while paragraph is not None and not paragraph.xpath(".//a:blip"):
        paragraph = paragraph.getnext()
    if paragraph is None:
        return []
    outputs = []
    for index, blip in enumerate(paragraph.xpath(".//a:blip"), 1):
        relationship_id = blip.get(qn("r:embed"))
        part = prd.part.related_parts[relationship_id]
        extension = part.content_type.split("/")[-1].replace("jpeg", "jpg")
        path = temp_dir / f"prd_gantt_{index}.{extension}"
        path.write_bytes(part.blob)
        outputs.append(path)
    return outputs


def replace_timeline(document: Document, temp_dir: Path) -> None:
    gantt = find_paragraph(document, "5.2 Actual Gantt Chart")
    wbs = find_paragraph(document, "5.3 Overall Product Hierarchy / Functionality WBS")
    clear_between(gantt, wbs)
    add_paragraph_before(
        document,
        wbs,
        "The following Gantt chart is the proposed schedule retained from the PRD. It must not be "
        "presented as the actual completion record; the actual task dates are still required below.",
    )
    for index, image_path in enumerate(extract_prd_gantt(temp_dir), 1):
        add_picture_before(document, wbs, image_path, f"Figure 5.{index}. Proposed PRD Gantt chart, part {index}.")
    add_gap_box(document, wbs, "G05", "Actual task timeline")

    communication = find_paragraph(document, "5.5 Communication and Scope Management")
    requirements = find_paragraph(document, "6 Requirement Definition")
    clear_between(communication, requirements)
    add_paragraph_before(
        document,
        requirements,
        "The team uses GitHub for source control and shared implementation. Sprint 1 work was "
        "divided around explicit PreparedText, ClaimAnalysis, RetrievalResult and AssessmentResult "
        "handoffs so members could develop against common samples. Meeting notes and task records "
        "were used to review progress and identify decisions still requiring agreement.",
    )
    add_bullets_before(
        document,
        requirements,
        [
            "Scrum provides two-week iteration structure, while a Kanban-style task view tracks work in progress.",
            "Git branches, commits and pull/fetch workflows provide implementation history and integration control.",
            "Shared JSON/Pydantic interfaces define component inputs, outputs and controlled failure behaviour.",
            "The 13 August implementation-planning record identifies the first pipeline goal, component owners and review checklist.",
        ],
    )
    add_gap_box(document, requirements, "G07", "Formal communication schedule, review decisions and scope approvals")
    add_heading_before(document, requirements, "5.6 Project Scope Statement", 2)
    add_table_before(
        document,
        requirements,
        ["Scope area", "Approved project position"],
        [
            ("In scope", "Android; English short-form text; supported public webpage links; supported static images; evidence-backed results; Free/Premium accounts; System Administrator and Data Engineer functions."),
            ("Premium scope", "Image upload with caption, OCR review/correction, caption-image context analysis and a separate probabilistic AI-generation/manipulation likelihood assessment."),
            ("Operational scope", "Account management, feedback/incorrect-result review, data-quality investigation and ingestion-pipeline monitoring."),
            ("Out of scope", "Video and speech analysis, private/paywalled/inaccessible pages, non-English analysis and any claim that a model score proves truth or manipulation."),
            ("Current increment", "Authenticated Android English-text checking, evidence retrieval, assessment, Firestore history and server-controlled allowances."),
        ],
        widths=[3.5, 13.0],
    )


def replace_requirements(document: Document) -> None:
    chapter = find_paragraph(document, "6 Requirement Definition")
    functional = find_paragraph(document, "7 Functional Requirements")
    clear_between(chapter, functional)
    add_heading_before(document, functional, "6.1 Stakeholder Identification and Analysis", 2)
    add_paragraph_before(
        document,
        functional,
        "Requirements were organised around intended older individual users, family/community "
        "helpers, operational users, the project team, supervisor/assessor and external service "
        "dependencies. Free and Premium users need a readable checking journey; System "
        "Administrators and Data Engineers require separate role-restricted operational workflows.",
    )
    add_heading_before(document, functional, "6.2 Requirements Gathering", 2)
    add_bullets_before(
        document,
        functional,
        [
            "Desk research into misinformation tools, older-user accessibility and comparable products.",
            "Project scope and feature analysis recorded in the PRD.",
            "User stories, use cases, operating constraints and dependencies recorded in the URS.",
            "Architecture, database, activity, sequence and interface design recorded in the TDM.",
            "Sprint planning, shared interface samples, implementation reviews and automated/live testing.",
        ],
    )
    add_paragraph_before(
        document,
        functional,
        "The available records do not show a completed older-user interview, survey or formal stakeholder acceptance exercise.",
    )
    add_gap_box(document, functional, "G08", "User consultation or requirements-review record")
    add_heading_before(document, functional, "6.3 Summary of User and System Needs", 2)
    add_table_before(
        document,
        functional,
        ["Need", "Required response"],
        [
            ("Simple access", "Registration, email verification, login, password recovery and clear role/tier access."),
            ("Supported checking", "English text and public links for individual users; supported image functions for Premium users."),
            ("Understandable result", "Concern label, evidence, citations, explanation, uncertainty and a safe next action."),
            ("Honest limitation", "Not Enough Information where evidence is insufficient; separate factual and image-authenticity assessments."),
            ("Private history", "Authenticated ownership checks, saved results, deletion controls and protected account data."),
            ("Operational oversight", "Role-restricted account, feedback, incorrect-result and ingestion-monitoring functions."),
        ],
        widths=[4.2, 12.3],
    )
    add_heading_before(document, functional, "6.4 Requirements Analysis and Prioritisation", 2)
    add_paragraph_before(
        document,
        functional,
        "The team prioritised a vertical text-checking increment because it exercises authentication, "
        "input preparation, claim analysis, evidence retrieval, assessment, persistence and Android "
        "presentation. Link, image, payment and operational-portal requirements remain in the approved "
        "scope but are separate planned increments.",
    )
    add_heading_before(document, functional, "6.5 Acceptance Criteria and Traceability", 2)
    add_table_before(
        document,
        functional,
        ["Area", "Documented acceptance condition", "Evidence status"],
        [
            ("Access", "Protected functions require a verified, active account and correct role.", "Developer tests exist; stakeholder sign-off is missing."),
            ("Text input", "Accept valid English text up to 5,000 characters and reject invalid input safely.", "Implemented and tested."),
            ("Allowance", "Free: one successful check/day; Premium: 60/month; failed work does not consume allowance.", "Implemented baseline and tested."),
            ("Result", "Return claim, concern, uncertainty, explanation and inspectable evidence; NEI has no fabricated score.", "Implemented baseline; general accuracy not established."),
            ("Ownership", "Users can access only their own saved results.", "Implemented and tested."),
            ("Full scope", "Complete link, image, payment and operational use cases defined in Chapters 7 and 14.", "Planned; not accepted as implemented."),
        ],
        widths=[2.8, 8.3, 5.4],
    )
    add_gap_box(document, functional, "G09", "Approved acceptance criteria and stakeholder sign-off")
    add_heading_before(document, functional, "6.6 Requirements Tools and Records", 2)
    add_paragraph_before(
        document,
        functional,
        "The team uses requirements and design documents, role-based use cases, shared JSON/Pydantic "
        "contracts, GitHub history, automated tests, evaluation datasets and meeting records to trace "
        "requirements into implementation and testing.",
    )


def replace_functional_requirements(document: Document) -> None:
    chapter = find_paragraph(document, "7 Functional Requirements")
    next_chapter = find_paragraph(document, "8 Non-Functional Requirements")
    clear_between(chapter, next_chapter)
    add_heading_before(document, next_chapter, "7.1 Functional Hierarchy", 2)
    add_table_before(
        document,
        next_chapter,
        ["Module", "Functions", "Status"],
        [
            ("Account and access", "Register, verify, login/logout, reset password, profile, deletion and role/tier checks.", "Core access implemented; deletion and operational role flows planned."),
            ("Text and link checking", "Text submission, public-link retrieval, malicious-link screening and evidence-backed result.", "Text implemented; direct link and URL-risk checking planned."),
            ("Result management", "View evidence, history, deletion, incorrect-result reporting, review and sharing.", "View/history/native sharing implemented; other functions planned."),
            ("Premium image", "Image/caption upload, OCR correction, context analysis and image-authenticity likelihood.", "Planned."),
            ("Subscription", "Upgrade, Premium registration, renewal, cancellation and entitlement confirmation.", "Planned; demonstration entitlement is server-controlled."),
            ("System administration", "Account search/view/create, suspension controls and feedback review.", "Planned."),
            ("Data engineering", "Incorrect-result dashboard/report investigation and ingestion monitoring.", "Planned."),
        ],
        widths=[3.6, 8.3, 4.6],
    )
    add_heading_before(document, next_chapter, "7.2 Access Levels", 2)
    add_table_before(
        document,
        next_chapter,
        ["Function", "Free", "Premium", "System Admin", "Data Engineer"],
        [
            ("Text/link assessment", "Yes", "Yes", "No", "No"),
            ("Image/OCR/context/authenticity", "Locked", "Planned", "No", "No"),
            ("Personal result history", "Yes", "Yes", "No", "No"),
            ("Account administration", "No", "No", "Planned", "No"),
            ("Feedback/incorrect-result operations", "Submit planned", "Submit planned", "Review planned", "Investigate planned"),
            ("Ingestion monitoring", "No", "No", "No", "Planned"),
        ],
        widths=[5.4, 2.3, 2.5, 3.1, 3.2],
    )
    add_heading_before(document, next_chapter, "7.3 Implemented Text-Checking Functions", 2)
    add_bullets_before(
        document,
        next_chapter,
        [
            "Register a Free account, verify email, sign in/out, request a password reset, view a profile and update the display name.",
            "Submit up to 5,000 characters of English text through an authenticated Android session.",
            "Prepare input, classify/extract a claim, retrieve evidence, assess passages and return a concern result with uncertainty and citations.",
            "Store owned completed/failed results in Firestore and reopen personal history.",
            "Apply daily/monthly successful-check allowances, idempotent replay and a per-user processing lock.",
            "Share a text summary and citation URLs through the device share sheet.",
        ],
    )
    add_heading_before(document, next_chapter, "7.4 Planned Individual-User Functions", 2)
    add_bullets_before(
        document,
        next_chapter,
        [
            "Submit supported public article/webpage links and check malicious-link indicators.",
            "Delete an account or selected result history with confirmation and ownership checks.",
            "Report an incorrect result, leave a review and share a public result link where approved.",
            "Upgrade, register for Premium, renew or cancel through a selected payment/subscription mechanism.",
            "Upload supported static images with captions, review OCR text, analyse context and request an image-authenticity likelihood assessment.",
        ],
    )
    add_heading_before(document, next_chapter, "7.5 Planned Operational Functions", 2)
    add_bullets_before(
        document,
        next_chapter,
        [
            "System Administrators authenticate, manage their profiles, query/view accounts, create operational accounts, suspend/unsuspend eligible accounts and review feedback.",
            "Data Engineers authenticate, manage their profiles, view incorrect-result dashboards and reports, inspect provenance and monitor ingestion runs.",
        ],
    )
    add_heading_before(document, next_chapter, "7.6 Dependencies, Inputs and Outputs", 2)
    add_table_before(
        document,
        next_chapter,
        ["Area", "Input/dependency", "Output"],
        [
            ("Authentication", "Email/password, Firebase availability and verified identity token.", "Authenticated role/tier context or controlled denial."),
            ("Text analysis", "English text, allowance, AI/search credentials and accessible sources.", "Owned result containing claim, concern, uncertainty, explanation and evidence."),
            ("Link analysis", "Supported public URL, safe-fetch rules and webpage access.", "Extracted content, URL-risk indicators and assessment; planned."),
            ("Image analysis", "Supported static image, caption, OCR/authenticity providers and Premium entitlement.", "OCR/context/authenticity outputs with limitations; planned."),
            ("Subscription", "Payment provider confirmation.", "Server-controlled Premium entitlement; planned."),
            ("Operations", "Authorised operational role and stored feedback/pipeline records.", "Account actions, investigation records and monitoring status; planned."),
        ],
        widths=[3.0, 7.0, 6.5],
    )


def add_nfr_targets(document: Document) -> None:
    next_chapter = find_paragraph(document, "9 Other Requirements")
    existing = optional_paragraph(document, "8.6 Documented Non-Functional Targets")
    if existing:
        remove_from_to(existing, next_chapter)
    add_heading_before(document, next_chapter, "8.6 Documented Non-Functional Targets", 2)
    add_paragraph_before(
        document,
        next_chapter,
        "The following values come from the PRD/URS and remain targets unless a measured result is stated in the appendix.",
    )
    add_table_before(
        document,
        next_chapter,
        ["Category", "Documented target", "Current evidence"],
        [
            ("Capacity", "At least 100 concurrent users and 100 requests per minute; architecture should support 20% annual user growth.", "Not load-tested."),
            ("Interface", "Initial login/home/register load within 2 seconds; taps respond within 1 second; assets within 2.5 seconds on a stable connection.", "No formal device measurement."),
            ("Result time", "Text and link results within 5 minutes; image functions within 10 minutes; cancel work beyond 20 minutes.", "Selected retrieval latency exists; end-to-end/load compliance is not established."),
            ("Storage", "Notify a user at 80% of their allocated personal storage quota.", "Fields exist; warning enforcement is not implemented."),
            ("Security", "Protect sensitive data in transit and at rest; authenticated protected access; 30-minute inactivity expiry.", "Authentication/ownership controls exist; session-expiry target is not verified."),
            ("Accessibility", "Readable typography, high contrast, large targets, simple navigation and plain-language feedback for older users.", "Designed, but no completed older-user study."),
        ],
        widths=[2.8, 8.4, 5.3],
    )


def replace_other_requirements(document: Document) -> None:
    chapter = find_paragraph(document, "9 Other Requirements")
    next_chapter = find_paragraph(document, "10 Risk Management")
    clear_between(chapter, next_chapter)
    add_heading_before(document, next_chapter, "9.1 User and Operating Requirements", 2)
    add_paragraph_before(
        document,
        next_chapter,
        "Individual users require a supported Android device, registered email address and stable "
        "internet connection. The initial system supports English text, public webpage links and "
        "supported static images; video, audio, inaccessible pages and non-English analysis are outside scope.",
    )
    add_paragraph_before(
        document,
        next_chapter,
        "The complete interface must explain supported inputs, access tiers, allowance behaviour, "
        "result labels, uncertainty, reporting, history controls and safe next actions. A user should "
        "not rely on UIABO as the sole basis for a high-impact decision.",
    )
    add_heading_before(document, next_chapter, "9.2 Privacy, Data Lifecycle and Legal Records", 2)
    add_bullets_before(
        document,
        next_chapter,
        [
            "Registration and Premium enrolment flows require terms/privacy notice presentation and recorded user consent.",
            "Users must submit only content they are permitted to process; screenshots may contain third-party personal data.",
            "Protected account and result data require authenticated ownership or authorised role checks.",
            "Provider credentials remain on the backend; the final privacy notice must disclose relevant cloud/external processing.",
            "Planned account deletion removes the authentication account, profile and associated history after confirmation and re-authentication.",
            "Planned result deletion verifies ownership before permanently removing selected records.",
            "Retention periods, deletion completion behaviour and provider data handling still require formal approval and documentation.",
        ],
    )
    add_gap_box(document, next_chapter, "G13", "Approved privacy notice, consent basis and data-lifecycle documentation")
    add_heading_before(document, next_chapter, "9.3 Assumptions and Dependencies", 2)
    add_table_before(
        document,
        next_chapter,
        ["Assumption/dependency", "Project position"],
        [
            ("User access", "A supported Android device, registered email and sufficiently stable connection are available; a helper may assist onboarding."),
            ("Submitted content", "Content is permitted, sufficiently clear and within supported formats."),
            ("Evidence", "Relevant credible evidence may be absent or conflicting; Not Enough Information is an expected result."),
            ("Cloud platform", "Authentication, profiles, allowances and results depend on Firebase configuration and availability."),
            ("Analysis services", "AI, search, OCR, URL screening and image-authenticity functions depend on provider limits and supported outputs."),
            ("Subscription", "Premium activation/renewal/cancellation depends on a selected payment or app-store mechanism."),
        ],
        widths=[4.1, 12.4],
    )
    add_heading_before(document, next_chapter, "9.4 Operational Recovery", 2)
    add_paragraph_before(
        document,
        next_chapter,
        "The implemented request lease and idempotency key support recovery from interrupted text checks. "
        "A repeated completed request returns its existing saved result without another deduction. Provider "
        "failures return controlled errors. Production monitoring, backup, restore testing and disaster-recovery "
        "records remain unfinished.",
    )
    add_gap_box(document, next_chapter, "G14", "Tested backup and recovery procedure")


def replace_user_stories(document: Document) -> None:
    chapter = find_paragraph(document, "13 User Stories")
    next_chapter = find_paragraph(document, "14 Use Case Descriptions")
    clear_between(chapter, next_chapter)
    add_paragraph_before(
        document,
        next_chapter,
        "The following stories are retained from the approved URS. Their presence records the full "
        "project scope; it does not mean every story is implemented.",
    )
    for group_index, (label, stories, page_number) in enumerate(parse_user_stories(), 1):
        add_heading_before(document, next_chapter, f"13.{group_index} {label}", 2)
        rows = []
        for story in stories:
            match = re.match(r"UC-(\d{2}):\s*(.*)", story)
            if not match:
                continue
            number = int(match.group(1))
            rows.append((f"UC-{number:02d}", match.group(2), case_status(number)))
        add_table_before(
            document,
            next_chapter,
            ["Use case", "User story", "Implementation status"],
            rows,
            widths=[2.0, 9.4, 5.1],
        )
        add_paragraph_before(document, next_chapter, f"Source: UIABO URS, page {page_number}.")


def replace_use_cases(document: Document) -> None:
    chapter = find_paragraph(document, "14 Use Case Descriptions")
    next_chapter = find_paragraph(document, "15 Use Case Diagram")
    clear_between(chapter, next_chapter)
    add_paragraph_before(
        document,
        next_chapter,
        "These editable descriptions retain UC-01 to UC-43 from the approved URS. Status rows "
        "distinguish the current prototype from planned behaviour.",
    )
    records = parse_use_cases()
    groups = [
        (1, "Shared Across Individual Users", range(1, 16)),
        (2, "Free User", range(16, 18)),
        (3, "Premium User", range(18, 25)),
        (4, "System Administrator", range(25, 36)),
        (5, "Data Engineer", range(36, 44)),
    ]
    for group_number, label, numbers in groups:
        group_heading = add_heading_before(document, next_chapter, f"14.{group_number} {label}", 2)
        group_heading.paragraph_format.page_break_before = group_number > 1
        for offset, number in enumerate(numbers, 1):
            name, rows, page_number = records[number]
            case_heading = add_heading_before(
                document,
                next_chapter,
                f"14.{group_number}.{offset} UC-{number:02d} — {name}",
                3,
            )
            case_heading.paragraph_format.page_break_before = True
            output_rows = []
            for field, value in rows:
                output_rows.append((field, value))
                if field == "ID":
                    output_rows.append(("Implementation status", case_status(number)))
            add_table_before(
                document,
                next_chapter,
                ["Use Case Field", "Description"],
                output_rows,
                widths=[3.2, 13.3],
            )
            add_paragraph_before(document, next_chapter, f"Source: UIABO URS, page {page_number}.")


def extract_use_case_images(temp_dir: Path) -> list[tuple[str, Path]]:
    pdf = fitz.open(URS_PDF)
    entries = []
    for label, page_number in [
        ("Free User", 14),
        ("Premium User", 33),
        ("System Administrator", 43),
        ("Data Engineer", 56),
    ]:
        page = pdf[page_number - 1]
        images = page.get_images(full=True)
        if not images:
            raise ValueError(f"No use-case diagram found on URS page {page_number}")
        extracted = pdf.extract_image(images[0][0])
        path = temp_dir / f"use_case_{label.lower().replace(' ', '_')}.{extracted['ext']}"
        path.write_bytes(extracted["image"])
        entries.append((label, path))
    pdf.close()
    return entries


def replace_use_case_diagrams(document: Document, temp_dir: Path) -> None:
    chapter = find_paragraph(document, "15 Use Case Diagram")
    next_chapter = find_paragraph(document, "16 System Design")
    clear_between(chapter, next_chapter)
    add_paragraph_before(
        document,
        next_chapter,
        "The diagrams below are retained from the URS and show the complete planned role boundaries. "
        "Registration and recovery represent access before authentication; no separate permanent "
        "unregistered-user role is required.",
    )
    for index, (label, path) in enumerate(extract_use_case_images(temp_dir), 1):
        add_heading_before(document, next_chapter, f"15.{index} {label}", 2)
        add_picture_before(
            document,
            next_chapter,
            path,
            f"Figure 15.{index}. {label} use-case diagram retained from the UIABO URS.",
            width=15.5,
        )


def add_full_system_design(document: Document) -> None:
    next_chapter = find_paragraph(document, "17 Conclusion")
    existing = optional_paragraph(document, "16.10 Full Planned-System Diagrams")
    if existing:
        remove_from_to(existing, next_chapter)
    add_heading_before(document, next_chapter, "16.10 Full Planned-System Diagrams", 2)
    add_paragraph_before(
        document,
        next_chapter,
        "These TDM diagrams describe the planned complete system. The current runtime implements the "
        "Android text path and a subset of the shown Firestore collections.",
    )
    diagrams = [
        (TDM_ROOT / "design_approach/UIABO_Architecture_Design.png", "Figure 16.6. Planned UIABO system architecture."),
        (TDM_ROOT / "design_approach/UIABO_Data_Flow.png", "Figure 16.7. Planned UIABO data flow."),
        (TDM_ROOT / "database_design/UIABO_Firestore_Database_Diagram.png", "Figure 16.8. Planned Firestore database design."),
    ]
    for path, caption in diagrams:
        add_picture_before(document, next_chapter, path, caption, width=16.0)

    add_heading_before(document, next_chapter, "16.11 Proposed Interface Samples", 2)
    add_paragraph_before(
        document,
        next_chapter,
        "The following static samples were prepared during earlier interface-design work. They are "
        "proposed screens, not emulator evidence. The obsolete public landing page is intentionally excluded.",
    )
    groups = [
        (
            "16.11.1 Pre-authentication",
            [
                (UI_ROOT / "unregistered_user/screenshots/login.png", "Proposed login screen"),
                (UI_ROOT / "unregistered_user/screenshots/signup.png", "Proposed registration screen"),
            ],
        ),
        (
            "16.11.2 Free User",
            [
                (UI_ROOT / "free_user/screenshots/dashboard.png", "Free Home and feature menu"),
                (UI_ROOT / "free_user/screenshots/text_analysis.png", "Text analysis input"),
                (UI_ROOT / "free_user/screenshots/link_analysis.png", "Planned public-link analysis"),
                (UI_ROOT / "free_user/screenshots/result.png", "Result and evidence"),
                (UI_ROOT / "free_user/screenshots/result_history.png", "Result history"),
                (UI_ROOT / "free_user/screenshots/upgrade.png", "Planned Premium upgrade"),
            ],
        ),
        (
            "16.11.3 Premium User",
            [
                (UI_ROOT / "premium_user/screenshots/dashboard.png", "Premium Home and feature menu"),
                (UI_ROOT / "premium_user/screenshots/subscription.png", "Planned subscription screen"),
                (UI_ROOT / "premium_user/screenshots/image_caption.png", "Planned image and caption input"),
                (UI_ROOT / "premium_user/screenshots/ocr_analysis.png", "Planned OCR review"),
                (UI_ROOT / "premium_user/screenshots/context_analysis.png", "Planned caption-image context result"),
                (UI_ROOT / "premium_user/screenshots/deepfake_analysis.png", "Planned image-authenticity result"),
            ],
        ),
        (
            "16.11.4 System Administrator",
            [
                (UI_ROOT / "system_admin/screenshots/dashboard.png", "Planned administration dashboard"),
                (UI_ROOT / "system_admin/screenshots/account_management.png", "Planned account management"),
                (UI_ROOT / "system_admin/screenshots/create_operational_account.png", "Planned operational-account creation"),
                (UI_ROOT / "system_admin/screenshots/feedback.png", "Planned feedback review"),
            ],
        ),
        (
            "16.11.5 Data Engineer",
            [
                (UI_ROOT / "data_engineer/screenshots/dashboard.png", "Planned incorrect-result dashboard"),
                (UI_ROOT / "data_engineer/screenshots/incorrect_result_report.png", "Planned report investigation"),
                (UI_ROOT / "data_engineer/screenshots/pipeline_monitoring.png", "Planned ingestion monitoring"),
            ],
        ),
    ]
    for heading, entries in groups:
        add_heading_before(document, next_chapter, heading, 3)
        missing = [str(path) for path, _ in entries if not path.exists()]
        if missing:
            raise FileNotFoundError("Missing interface samples: " + ", ".join(missing))
        add_picture_grid(document, next_chapter, entries)


def update_test_record(document: Document) -> None:
    target = None
    for table in document.tables:
        if table.rows and table.cell(0, 0).text.strip() == "Record":
            if any("Backend suite" in row.cells[0].text for row in table.rows[1:]):
                target = table
                break
    if target is None:
        raise ValueError("Appendix test-result table not found")
    for row in target.rows[1:]:
        if "Backend suite" in row.cells[0].text:
            row.cells[0].text = "Recorded backend baseline, 9 September"
            row.cells[1].text = "400 tests passed; one Starlette/HTTPX deprecation warning."
            row.cells[2].text = "Previously recorded baseline; the exact environment must remain identifiable."
    cells = target.add_row().cells
    cells[0].text = "PTD review rerun, 9 September"
    cells[1].text = "398 passed, 2 failed; one existing deprecation warning."
    cells[2].text = (
        "Failures concerned real retrieval/API wiring under local web mode and the expected "
        "catalogue-mode default. Resolve configuration isolation before claiming a current clean suite."
    )
    for row in target.rows:
        set_cant_split(row)


def update_document_control(document: Document) -> None:
    for paragraph in document.paragraphs[:20]:
        if paragraph.text.startswith("Version 0.5"):
            paragraph.text = "Version 0.6 — 9 September 2026"
    for table in document.tables:
        if table.rows and table.cell(0, 0).text.strip() == "Field" and table.cell(0, 1).text.strip() == "Record":
            for row in table.rows[1:]:
                if row.cells[0].text.strip() == "Version":
                    row.cells[1].text = "0.6 — working draft"
                    break
            break
    revision_table = document.tables[2]
    cells = revision_table.add_row().cells
    cells[0].text = "9 September 2026"
    cells[1].text = "0.6"
    cells[2].text = (
        "Added reusable full-scope PRD, URS and TDM content; restored all role stories/use cases "
        "and clearly separated planned functions from implemented evidence."
    )


def ensure_update_fields_on_open(document: Document) -> None:
    settings = document.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def update_word_fields(path: Path) -> tuple[bool, int | None, str | None]:
    try:
        import win32com.client

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        opened = word.Documents.Open(str(path), ReadOnly=False, AddToRecentFiles=False)
        try:
            for toc in opened.TablesOfContents:
                toc.Update()
            opened.Fields.Update()
            pages = opened.ComputeStatistics(2)
            opened.Save()
        finally:
            opened.Close(False)
            word.Quit()
        return True, pages, None
    except Exception as error:  # Word field refresh is optional; the DOCX remains usable.
        return False, None, str(error)


def validate(document_path: Path) -> dict[str, int]:
    document = Document(document_path)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    use_cases = [h for h in headings if re.match(r"^14\.\d+\.\d+ UC-\d{2}", h)]
    required = [
        "2.7 Business Model and Operating Costs",
        "5.6 Project Scope Statement",
        "7.6 Dependencies, Inputs and Outputs",
        "8.6 Documented Non-Functional Targets",
        "9.2 Privacy, Data Lifecycle and Legal Records",
        "16.10 Full Planned-System Diagrams",
        "16.11 Proposed Interface Samples",
    ]
    missing = [value for value in required if value not in headings]
    if missing:
        raise ValueError(f"Required headings missing after update: {missing}")
    if len(use_cases) != 43:
        raise ValueError(f"Expected 43 detailed use cases; found {len(use_cases)}")
    text_content = "\n".join(p.text for p in document.paragraphs)
    if re.search(r"AIza[0-9A-Za-z_-]{20,}|-----BEGIN PRIVATE KEY-----", text_content):
        raise ValueError("Credential-like content detected")
    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "headings": len(headings),
        "use_cases": len(use_cases),
        "figures": len(document.inline_shapes),
    }


def main() -> None:
    for path in [DOCX_PATH, URS_PDF, PRD_DOCX]:
        if not path.exists():
            raise FileNotFoundError(path)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = BACKUP_ROOT / f"before_previous_docs_{stamp}"
    backup_dir.mkdir(parents=True, exist_ok=False)
    shutil.copy2(DOCX_PATH, backup_dir / DOCX_PATH.name)

    document = Document(DOCX_PATH)
    with tempfile.TemporaryDirectory(prefix="uiabo_ptd_sources_") as temporary:
        temp_dir = Path(temporary)
        replace_introduction(document)
        replace_overview(document)
        replace_timeline(document, temp_dir)
        replace_requirements(document)
        replace_functional_requirements(document)
        add_nfr_targets(document)
        replace_other_requirements(document)
        replace_user_stories(document)
        replace_use_cases(document)
        replace_use_case_diagrams(document, temp_dir)
        add_full_system_design(document)
        update_test_record(document)
        update_document_control(document)
        ensure_update_fields_on_open(document)

        temporary_docx = temp_dir / DOCX_PATH.name
        document.save(temporary_docx)
        validate(temporary_docx)
        shutil.copy2(temporary_docx, DOCX_PATH)

    refreshed, pages, refresh_error = update_word_fields(DOCX_PATH)
    result = validate(DOCX_PATH)
    print(f"Saved: {DOCX_PATH}")
    print(f"Backup: {backup_dir / DOCX_PATH.name}")
    print(f"Validation: {result}")
    print(f"Word fields refreshed: {refreshed}; pages: {pages}")
    if refresh_error:
        print(f"Word refresh note: {refresh_error}")


if __name__ == "__main__":
    main()
