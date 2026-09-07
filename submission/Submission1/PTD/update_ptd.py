"""Compile the existing UIABO PTD sections with a dated progress update.

Original PTDParts are read only. Generated copies go to PTDParts_Updated.
Export the combined DOCX with Word to the requested PDF after field updates.
Rebuilding overwrites generated copies; preserve manual edits before running again.
"""
from pathlib import Path
from copy import deepcopy
import re

from docx import Document
from docxcompose.composer import Composer
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
BASE = ROOT / 'PTDParts'
OUT = ROOT / 'PTDParts_Updated'
WORKSPACE = ROOT.parents[2]
REPO = WORKSPACE if (WORKSPACE / 'backend').is_dir() else WORKSPACE / 'uiabo'
DATE = '7 September 2026'
CODE = 'FYP-26-S3-30'


def p(d, text, style=None): return d.add_paragraph(text, style)
def h(d, text, level=1): return d.add_heading(text, level)
def bullets(d, values):
    for value in values: p(d, value, 'List Bullet')
def src(d, text):
    para = p(d, 'Evidence / source: ' + text)
    for run in para.runs:
        run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string('536779')
def tbl(d, headers, rows):
    t = d.add_table(rows=1, cols=len(headers)); t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers): cell.text = text
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = str(text)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(5)
                for run in para.runs: run.font.size = Pt(9.5)
    return t
def field(para, value):
    run = para.add_run()
    a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin')
    b=OxmlElement('w:instrText');b.set(qn('xml:space'),'preserve');b.text=value
    c=OxmlElement('w:fldChar');c.set(qn('w:fldCharType'),'end')
    run._r.extend([a,b,c])
def after(para, text, style=None):
    from docx.text.paragraph import Paragraph
    element=OxmlElement('w:p');para._p.addnext(element)
    out=Paragraph(element,para._parent);out.text=text
    if style:out.style=style
    return out
def insert_note(d, heading, text):
    para=next(p for p in d.paragraphs if p.text==heading)
    return after(para,text)
def load(name):
    d=Document(BASE/name)
    template=Document()
    for style_name in ['Title','Heading 1','Heading 2','Heading 3','List Bullet','Caption','Light Shading Accent 1']:
        if style_name not in d.styles:
            d.styles.element.append(deepcopy(template.styles[style_name].element))
    return d
def code(d, text):
    for line in text.strip().splitlines():
        para=p(d,line);para.paragraph_format.space_after=Pt(2)
        for run in para.runs:run.font.name='Consolas';run.font.size=Pt(9)


def style(d):
    for name,size in [('Normal',11),('Title',23),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
        s=d.styles[name];s.font.name='Arial';s.font.size=Pt(size)
        s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.08
        if name.startswith('Heading'):
            s.font.color.rgb=RGBColor.from_string('173F56');s.paragraph_format.keep_with_next=True
            s.paragraph_format.page_break_before=False
    if 'Frontmatter Heading' not in d.styles:
        s=d.styles.add_style('Frontmatter Heading',WD_STYLE_TYPE.PARAGRAPH)
        s.base_style=d.styles['Heading 2']
        outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'9')
        s.element.get_or_add_pPr().append(outline)
    for para in d.paragraphs:
        if para.style.name.startswith('Heading') and not re.match(r'^\d+(?:\.|\s)',para.text):
            para.style='Frontmatter Heading'
    # Original extraction sections sometimes ended in an explicit page break.
    # Collapse only consecutive empty break paragraphs added at composition boundaries.
    last_break=False
    for element in list(d.element.body):
        if element.tag!=qn('w:p'):
            last_break=False
            continue
        has_text=bool(''.join(element.itertext()).strip())
        has_drawing=bool(element.findall('.//'+qn('w:drawing')))
        page_breaks=[b for b in element.findall('.//'+qn('w:br')) if b.get(qn('w:type'))=='page']
        if not has_text and not has_drawing and page_breaks:
            if last_break:d.element.body.remove(element)
            last_break=True
        elif has_text or has_drawing:
            last_break=False
    # Use a break-before on the next content paragraph. A standalone break can
    # spill after a full table and create an otherwise empty page.
    from docx.text.paragraph import Paragraph
    for element in list(d.element.body):
        if element.tag!=qn('w:p') or ''.join(element.itertext()).strip():continue
        if element.findall('.//'+qn('w:drawing')):continue
        breaks=[b for b in element.findall('.//'+qn('w:br')) if b.get(qn('w:type'))=='page']
        if not breaks:continue
        following=element.getnext();empty=[]
        while following is not None and following.tag==qn('w:p') and not ''.join(following.itertext()).strip() and not following.findall('.//'+qn('w:drawing')) and not following.findall('.//'+qn('w:fldChar')):
            empty.append(following);following=following.getnext()
        if following is not None and following.tag==qn('w:p'):
            Paragraph(following,d._body).paragraph_format.page_break_before=True
            d.element.body.remove(element)
            for blank in empty:d.element.body.remove(blank)
    for s in d.sections:
        s.page_width=Cm(21);s.page_height=Cm(29.7)
        s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7)
        s.left_margin=Cm(2);s.right_margin=Cm(2)
        s.header_distance=Cm(.7);s.footer_distance=Cm(.7)
        s.different_first_page_header_footer=False
        num=s._sectPr.find(qn('w:pgNumType'))
        if num is not None:s._sectPr.remove(num)
        header=s.header.paragraphs[0];header.clear()
        header.add_run(f'{CODE}  |  Preliminary Technical Documentation  |  {DATE}').font.size=Pt(8)
        footer=s.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run('Working draft 0.2  |  Page ').font.size=Pt(9);field(footer,' PAGE ')
    for shape in d.inline_shapes:
        factor=min(1,Cm(16.7)/shape.width,Cm(22.2)/shape.height)
        if factor<1:shape.width=int(shape.width*factor);shape.height=int(shape.height*factor)
    d.core_properties.author='FYP-26-S3-30 project team'
    d.core_properties.title='uiabo — Preliminary Technical Documentation'
    return d


def cover():
    d=Document();style(d)
    p(d,'School of Computing and Information Technology');p(d,'CSIT321 — Project')
    h(d,'uiabo.',0);h(d,'Preliminary Technical Documentation',0)
    p(d,'Project topic: CSIT-26-S3-30 — AI-Based Misinformation Detection for Short-Form Content')
    p(d,'Group: FYP-26-S3-30');p(d,'Supervisor: Mr Liaw Chun Huei');p(d,'Assessor: Mr Terrence Chew')
    p(d,'Version 0.2 — Working draft — '+DATE)
    original=load('00_Cover_and_Contents.docx')
    d.element.body.insert(len(d.element.body)-1,deepcopy(original.tables[0]._tbl))
    p(d,'The report retains the agreed product requirements and design, and records the implemented text prototype separately. Current and proposed functions are distinguished throughout.')
    d.add_page_break();h(d,'Document Control')
    tbl(d,['Field','Record'],[
        ('Document owner','FYP-26-S3-30 project team'),('Source baseline','UIABO PTDParts, PRD, URS and TDM in the submission folder'),
        ('Comparison reference','FYP-26-S2-XX_PrelimTechDocs.pdf; structure and coverage only'),
        ('Progress snapshot',DATE+'; live text pipeline, Android app and Firestore integration'),
        ('Review status','Working draft for team review; no supervisor sign-off asserted'),
    ])
    p(d,'Record of Revision','Frontmatter Heading')
    tbl(d,['Date','Version','Recorded change'],[
        ('4 September 2026','Initial section draft','Date carried by the existing UIABO PTD section files.'),
        (DATE,'0.2','Compiled UIABO sections; added coverage absent from the new sample comparison; refreshed implementation, testing, scope and remaining-work records.'),
    ])
    p(d,'Evidence categories','Frontmatter Heading')
    bullets(d,['Implemented: observed in the current workspace and/or a dated test record.',
        'Planned requirement: retained from PRD/URS/TDM; its presence in a design does not establish completion.',
        'Proposed addition: a draft analysis or management process prepared for team review, with no invented agreement or sign-off.',
        'Pending evidence: information or validation not yet available. See Section 18 and PTD_MISSING_INFORMATION.md.'])
    d.add_page_break();p(d,'Table of Contents','Frontmatter Heading');field(p(d,''),' TOC \\o "1-2" \\h \\z \\u ')
    return d


STATUS_ROWS=[
    ('Account access','Available','Free registration, verification-link gating, sign-in, password reset and sign-out; profile name editing.'),
    ('Input preparation','Available','English detection/validation, meaning-preserving normalization, length checks and warnings.'),
    ('Claim analysis','Available with limits','Ollama Cloud classification ensemble and source-span extraction; satire/disagreement limits remain.'),
    ('Evidence retrieval','Available with limits','Google Fact Check discovery, Tavily extraction and scoped search, citations and controlled failures.'),
    ('Assessment','Available baseline','Lexical evidence stance and heuristic scores; semantic mistakes observed, accuracy not established.'),
    ('Android text flow','Available','TDM-based Home, text submission, progress, result detail, source links and personal history.'),
    ('Firestore/allowance','Available','Owner-bound saved results; atomic completion/charge; idempotent retries; free daily and premium monthly counters.'),
    ('Webpage submission','Planned','Direct fetching and malicious-link screening are not implemented; UI offers copied-text checking.'),
    ('Image/OCR/context/deepfake','Required/planned','TDM designs exist; image pipelines remain unimplemented. Deepfake remains required.'),
    ('Payments','Planned','Premium may be granted for demonstration; no checkout, paid renewal or cancellation service.'),
    ('Feedback/deletion/public sharing','Partial/planned','Native text/source sharing exists. Public result links, reviews, reports and account/history deletion remain planned.'),
    ('Operational portal','Planned','System-administrator and data-engineer designs exist; working portals are not implemented.'),
]


def functionality():
    d=load('01_Product_Functionality_Overview.docx')
    d.paragraphs[1].text='uiabo is an elderly-friendly Android application for assessing potentially misleading English short-form content. The agreed product scope covers text, public webpage links and static images, with free and premium access. The September text prototype implements a narrower working slice; the scope and status table below distinguishes it from the full planned product.'
    note=insert_note(d,'Core Functions','The following functional descriptions retain the planned product scope from the PRD. They are not a claim that every function has been implemented. Audio and video analysis are excluded; required deepfake detection concerns supported static images.')
    d.add_page_break();h(d,'1.1 Current implementation status',2)
    tbl(d,['Area','Status','Evidence / limit'],STATUS_ROWS)
    src(d,'PRD 1.2; URS 2.2–2.3; current source files and evaluation/reports/app-integration/README.md, dated 7 September 2026.')
    return d


def overview():
    d=load('02_Project_Overview.docx')
    para=insert_note(d,'2.2 Learning Objectives','Proposed learning objectives for team review. These are goals aligned with the work, not retrospective claims that every member has demonstrated each outcome.')
    for text in ['Design and integrate an Android client, authenticated Python API and cloud persistence using explicit component contracts.',
        'Compare claim-analysis, retrieval and evidence-assessment methods using labelled examples and documented error analysis.',
        'Communicate uncertainty and evidence in an interface suitable for older Android users, then validate it through usability testing.',
        'Apply credential separation, user ownership, failure handling and reproducible testing in a multi-service application.',
        'Coordinate a five-person project using version control, documented handoffs, scope decisions and reviewable progress evidence.']:
        para=after(para,text,'List Bullet')
    d.add_page_break();h(d,'2.5 Problem statement, purpose and vision',2)
    p(d,'Short forwarded messages, captions and screenshots can circulate without their original context. Checking them requires users to identify the claim, find relevant sources and interpret competing evidence. UIABO focuses on older Singapore Android users who may benefit from a simpler guided workflow and readable explanations.')
    p(d,'The project purpose is to support evidence review before sharing. Its proposed vision is a clear, accessible checking assistant that communicates what is known, what is uncertain and where the evidence can be inspected. It does not promise an automatic guarantee of truth.')
    h(d,'2.6 Project objectives and measurable evidence',2)
    tbl(d,['Objective','How completion is assessed','Current position'],[
        ('Complete the basic text journey','Verified account → English text → cited result → owned history in Android.','Demonstrated in the local emulator and live Firestore checks.'),
        ('Preserve evidence and ownership','Save user_id, timestamps, passages and URLs; deny other users access.','Implemented; automated and live checks recorded.'),
        ('Improve assessment quality','Labelled held-out claims, evidence relevance and stance/verdict metrics.','Dataset and acceptance thresholds pending; no accuracy percentage claimed.'),
        ('Deliver required image features','OCR, caption/context and separate deepfake likelihood with limits.','Designs exist; detector selection and implementation pending.'),
        ('Validate usability and performance','Older-user task study, response-time distributions and load checks against URS targets.','Not yet established.'),
    ])
    h(d,'2.7 Conceptualization and product rationale',2)
    p(d,'The concept combines a guided Android input flow with claim-level evidence and plain-language uncertainty. The PRD compares existing monitoring, source-rating, fact-check retrieval and multimedia tools. UIABO’s proposed differentiator is their selected combination for older users, rather than a claim that it outperforms those tools. The staged pipeline enables individual components to be evaluated and replaced without redesigning the entire app.')
    h(d,'2.8 SWOT analysis',2)
    p(d,'Draft qualitative analysis derived from the agreed design and current implementation; review with the team.')
    tbl(d,['Area','Analysis'],[
        ('Strengths','Focused user group; readable TDM-based flow; source citations and uncertainty; explicit contracts; working authenticated text-to-history integration.'),
        ('Weaknesses','Lexical stance errors; small synthetic evaluation sets; restricted English source catalogue; unfinished image, payment and portal work.'),
        ('Opportunities','Evaluate with older users/community helpers; improve Singapore-context examples; replace weak assessment rules while retaining the app/API contracts.'),
        ('Threats','Provider outages/quotas; missing or outdated evidence; prompt injection; user over-trust; limited time before submission; competing verification tools.'),
    ])
    h(d,'2.9 Proposed unique selling points',2)
    bullets(d,['Guided English Android checking for Singapore users aged 65 and above.',
        'Readable concern labels, uncertainty, claim-level explanations and inspectable evidence in one workflow.',
        'Free daily access with a defined premium allowance; image-context and authenticity functions remain part of the planned product.',
        'A distinction between factual misinformation risk and image-manipulation likelihood.'])
    p(d,'These are design propositions, not validated market advantages, formal partnerships or demonstrated accuracy superiority.')
    h(d,'2.10 Stakeholders',2)
    tbl(d,['Stakeholder','Need / responsibility'],[
        ('Older free/premium users','Understand submissions, evidence, uncertainty and allowance through accessible interaction.'),
        ('Family/community helpers','Assist onboarding and interpretation; not a separate account role.'),
        ('Supervisor: Mr Liaw Chun Huei','Academic and technical guidance; scope and progress review.'),
        ('Assessor: Mr Terrence Chew','Review the project deliverables and demonstrated outcomes.'),
        ('Five-person project team','Implement, evaluate, integrate and document the system; roles in Section 11.'),
        ('System administrator / data engineer','Planned operational access, feedback investigation and data-quality monitoring.'),
        ('Service and evidence providers','External dependencies with access, availability and usage constraints; no partnership asserted.'),
    ])
    src(d,'PRD 1–2; URS 1.2/2.3; existing PTD 2.1–2.4; project discussion. SWOT, explicit learning objectives and objective evidence criteria are proposed additions.')
    return d


def design():
    d=load('05_Project_Design.docx')
    insert_note(d,'5. Project Design','This chapter retains the submitted URS/TDM design, including facsimile use cases and proposed diagrams. Requirements, example screen values and future collections shown in those figures are not evidence of implementation. Section 1.1 and the additions below give the current state; Section 15 maps requirements to implementation and verification.')
    insert_note(d,'5.2 Non-Functional Requirements','The following numbers are requirements/targets, not achieved measurements. In particular, concurrent capacity, requests per minute, UI load time, inactivity expiry and storage warnings remain unverified. Existing source diagrams may retain earlier wording; current authentication uses email/password.')
    insert_note(d,'5.5 Database Design','Current consumer persistence uses users, usage_allowances, analysis_results and an implementation-level analysis_locks collection. Subscription, report, feedback, source-catalogue, ingestion-run and operational-audit schemas below remain proposed. The runtime source catalogue is currently a code allowlist, not an operational Firestore sources editor.')
    insert_note(d,'5.7 User Interface Design','These figures are the submitted TDM design mock-ups. Section 5.8 shows the available Android captures; unfinished image, payment and portal screens must not be described as deployed.')
    for para in d.paragraphs:
        if para.text=='Authentication: The system must implement username and passwords for all user accounts.':
            para.text='Authentication: individual accounts use an email address and password through Firebase Authentication. The backend verifies identity and account access for protected operations.'
    d.add_page_break();h(d,'5.8 Current Android implementation',2)
    p(d,'The current Home follows the TDM grid: Check text, Check webpage link, Image + caption, Read image text, Check image context and Check AI image. There is no audio placeholder. Only text analysis is connected; the link tab offers copied text and the image cards remain unavailable.')
    for name,label in [('home-premium.png','Figure 5.8a. Premium Home and monthly allowance.'),('history.png','Figure 5.8b. Private result history.'),('result.png','Figure 5.8c. Saved assessment view; this is a UI example, not a validated verdict.')]:
        para=p(d,'');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(REPO/'evaluation/reports/app-integration'/name),height=Cm(16))
        p(d,label,'Caption')
        if name!='result.png':d.add_page_break()
    src(d,'Existing Android emulator captures from 7 September 2026. Development overlay and account values are snapshot details; Premium grant did not create billing.')
    d.add_page_break();h(d,'5.9 Current authentication, persistence and retry design',2)
    p(d,'The Android client obtains a Firebase ID token and sends it to FastAPI. The backend verifies the token, email-verification state and active account, then derives the owner and tier from server-side records. It does not accept a client-selected user_id, role or allowance. Result/history reads use authenticated API routes.')
    tbl(d,['Step','Current behavior'],[
        ('Reserve','An idempotency key and text fingerprint identify the submission. A bounded per-user lease prevents concurrent requests from consuming the same allowance.'),
        ('Analyse','The pipeline runs with request-local intermediate storage and validated component handoffs.'),
        ('Commit','A transaction saves the authoritative owned result and decrements allowance for a completed analysis together. Failed analyses are saved without a charge.'),
        ('Retry','A successful repeated request with the same key/text returns its existing result. Reusing a key for different text is rejected.'),
        ('Read history','Query records belonging to the user, sort by creation time and result ID, then paginate. Other users cannot retrieve the record through the API.'),
    ])
    p(d,'Free counters reset each Singapore calendar day; premium counters reset each Singapore calendar month. A completed Not Enough Information result consumes one check. A client connection timeout may happen after server completion, so the app directs users to inspect saved history before starting a new request.')
    p(d,'Current Firestore additions include analysis_locks/{uid}, request_fingerprint and owner metadata on analysis_results. The history implementation reads the user’s matching records and sorts them in the backend. A proposed composite index exists in firestore.indexes.json but was not deployed because the current service account lacks index-administration permission; large histories need query optimization.')
    src(d,'backend/app/analyses/store.py; accounts/repository.py; routers/analysis.py; Firebase ID-token verification and transaction documentation listed in Section 13.')
    return d


def implementation():
    d=load('06_Project_Implementation.docx')
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.text=cell.text.replace('Axios','fetch')
                if 'The React Native application will access the database through the Firebase JavaScript SDK.' in cell.text:
                    cell.text=cell.text.replace('The React Native application will access the database through the Firebase JavaScript SDK.','The current Android client accesses protected account and result data through FastAPI; the backend Firebase Admin SDK performs Firestore operations.')
                if 'fetch reduces the networking code' in cell.text:
                    cell.text='The current app uses fetch for JSON requests and Firebase bearer headers, with request timeouts and friendly error handling. Server-only credentials and analysis logic remain outside the Android client.'
    insert_note(d,'6.2 Development Tools','Current implementation clarification: JavaScript/React Native with Expo 57, Python/FastAPI/Uvicorn, Firebase Authentication and Cloud Firestore. The mobile API wrapper uses fetch rather than the Axios library originally proposed. Database descriptions below retain the wider planned storage scope; see Section 5.5 for implemented collections.')
    d.add_page_break();h(d,'6.3 Current analysis components',2)
    tbl(d,['Component','Implemented method','Important limitation'],[
        ('Input preparation','Validate at most 5,000 characters; normalize Unicode/whitespace without discarding meaning; verify English; warn about instruction-like content.','A warning is not a complete prompt-injection defence.'),
        ('Claim classification','Concurrent Ollama Cloud calls to gpt-oss:120b, gemma4:31b and nemotron-3-super. At least two valid replies are needed; valid votes select a majority category.','A split returns unverifiable; agreement is not a truth probability.'),
        ('Claim extraction','gemma4:31b extracts a principal factual source span; validate JSON and occurrence in the submitted text.','A span match cannot prove all decisive context was retained.'),
        ('Evidence retrieval','Google published-check discovery → Tavily page extraction → scoped Tavily search when needed; validate, rank and deduplicate citations.','Lexical relevance can miss paraphrases or admit related but inconclusive sources.'),
        ('Evidence assessment','Supporting/contradicting/neutral stance using lexical overlap, negation, amounts and dates; heuristic quality and risk rules.','Known semantic stance errors; no calibrated probability or trained assessment model.'),
        ('Orchestration','Validate handoffs, skip unnecessary stages, distinguish failures from insufficient evidence, return structured results.','One principal claim is processed; multiple-claim coverage needs evaluation.'),
    ])
    h(d,'6.3.1 Model voting, validation and prompt handling',3)
    p(d,'Invalid responses and provider failures do not vote. Fewer than two valid classifiers produces a controlled failure. A winning category uses agreement divided by the three configured models, so 0.67 is agreement rather than a 67% likelihood of truth. Submitted content is placed in a separate untrusted data message; system instructions define the task. Locally enforced schemas and span checks reject malformed or fabricated extraction output.')
    p(d,'The configured claim stage uses a 90-second overall deadline and bounded provider calls. Retrieval uses a 65-second stage deadline; the current mobile text request waits up to 180 seconds. These bounds limit waiting but do not establish a response-time service guarantee. There are no automatic provider retries in these stages.')
    src(d,'Current claim_analysis/categories.py and service.py; evidence_retrieval service/providers; mobile/src/services/api.js. Ollama Cloud reference in Section 13.')
    d.add_page_break();h(d,'6.4 APIs, interfaces and dependencies',2)
    tbl(d,['Endpoint / boundary','Purpose'],[
        ('GET / and GET /health','Basic process reachability; does not prove provider readiness.'),
        ('PUT /account/me','Create/refresh the authenticated profile; verification status is enforced by subsequent protected operations.'),
        ('GET /account/me','Return the verified active user’s own profile and allowance.'),
        ('POST /analysis/text','Verified, active account; English text payload; optional Idempotency-Key; controlled errors and owned completed/failed persistence.'),
        ('GET /analysis/results','Owned result list, default page limit 20 and maximum 50; optional result-ID cursor.'),
        ('GET /analysis/results/{result_id}','Owned result detail; missing/other-user records return 404.'),
        ('Pipeline models','PreparedText → ClaimAnalysis → RetrievalResult → AssessmentResult → TextAnalysisResult, plus FailedAnalysisRecord.'),
    ])
    tbl(d,['Dependency','Purpose / credential boundary'],[
        ('Firebase Authentication / Firestore','Client identity and backend-owned storage. Admin JSON stays outside the repository; no password is stored in a user-profile document.'),
        ('Ollama Cloud','Claim-stage inference using a server-only API key; no local model download required.'),
        ('Google Fact Check Tools API','Search existing ClaimReview records; no claim-verification conclusion is implied by a search match alone.'),
        ('Tavily Search / Extract','Fetch source passages and search the scoped catalogue; provider keys remain backend-only.'),
        ('External source pages','Evidence/provenance; availability, publication date and meaning must be evaluated for each claim.'),
    ])
    p(d,'The app displays controlled errors for invalid input, authentication/access failure, depleted allowance, overlapping/reused requests and unavailable services. Empty successful searches can produce Not Enough Information; technical failures are not converted into a confident assessment.')
    src(d,'backend/app/routers, pipeline/shared/models.py; provider APIs in Section 13.')
    d.add_page_break();h(d,'6.5 Hosting, deployment and configuration',2)
    p(d,'Current delivery is a local Windows/Android-emulator demonstration with a cloud Firebase project and cloud AI/search providers. There is no verified public API deployment, installed release APK, Play Store listing or operational portal. Source-control use does not establish a CI/CD pipeline.')
    p(d,'Project-root .env supplies OLLAMA_API_KEY, GOOGLE_FACT_CHECK_API_KEY, TAVILY_API_KEY and GOOGLE_APPLICATION_CREDENTIALS. The credential path points outside Git. The mobile environment contains Firebase Web configuration and EXPO_PUBLIC_API_URL; it must not contain provider keys or the Admin JSON. Real credentials are deliberately omitted from this report.')
    code(d,r'''
cd C:\Dev\uiabo\backend
.\.venv\Scripts\python.exe -m uvicorn app.main:app --host 127.0.0.1 --port 8000

cd C:\Dev\uiabo\mobile
npm.cmd ci
npx.cmd expo start --lan
''')
    p(d,'The existing virtual environment and private configuration must be prepared first. The standard Android emulator uses http://10.0.2.2:8000 to reach this backend. This is local development HTTP; the submitted deployment design requires HTTPS. The PUM Section 2 provides the longer setup procedure.')
    p(d,'The current workspace uses Expo 57, React Native 0.86.3, React 19.2.3, Node v24.14.0 and Python 3.14.3. Expo’s SDK 57 reference lists Node 22.13.x and Android 7+ as framework minimums; these are not a UIABO device-compatibility certification. A clean-machine and physical-device matrix remains pending (Expo, SDK 57 reference).')
    h(d,'6.6 Release work still required',2)
    bullets(d,['Choose and document deployment targets, HTTPS endpoints, environment separation and release artifact/version.',
        'Provide reproducible Firebase rules/IAM configuration and a tested recovery/backup procedure.',
        'Implement and evaluate required static-image/deepfake functions, then payments and operational workflows as agreed.',
        'Validate assessment accuracy, latency, capacity, accessibility and failure behavior before stronger release claims.'])
    src(d,'Current READMEs, configuration templates, package files and PUM; official Expo SDK 57 reference accessed 7 September 2026.')
    return d


def risks():
    d=load('07_Risk_Analysis.docx')
    for para in list(d.paragraphs):
        if not para.text and not para._p.findall('.//'+qn('w:drawing')) and not para._p.findall('.//'+qn('w:br')):
            para._p.getparent().remove(para._p)
    insert_note(d,'7. Risk Analysis','The original PRD risk ratings below are retained as the planning baseline. They are qualitative estimates, not current measured likelihoods or completed compliance assessments. The update after the table records observed risks and remaining actions.')
    d.add_page_break();h(d,'7.1 Current risk treatment',2)
    tbl(d,['Risk / evidence','Treatment in place','Remaining action / responsible area'],[
        ('Assessment misreads debunking or unrelated negation.','Uncertainty, citations and prototype limitations are displayed.','Assessment owner + integration: evaluate semantic stance with quoted evidence and held-out labels.'),
        ('Cloud/search provider failure or quota.','Bounded calls, fallback retrieval, sanitized errors; failed runs do not charge.','Integration: record provider usage and representative outage/recovery behavior.'),
        ('Prompt injection or malformed AI output.','Untrusted-data prompts, local schema validation, source-span checks.','Claim/assessment owners: broader adversarial evaluation; no immunity claim.'),
        ('Another user reads a saved result or double charge occurs.','Server ownership, transactions, idempotency, per-user lease; focused tests passed.','Backend: broader review of rules/IAM, recovery and deployment controls.'),
        ('Evidence is outdated, related or duplicated.','Scoped sources, URL/content deduplication and dates retained when known.','Retrieval/assessment: temporal context and independent-source evaluation.'),
        ('Required features exceed available schedule.','September prototype narrowed to the basic text workflow; image tools labelled unavailable.','Project leader/team: confirm owners and dates for required deepfake/image delivery.'),
        ('User over-trust or difficult navigation.','Readable TDM-based screen layout and evidence/uncertainty messaging.','UI/testing: older-user usability study and false-reassurance analysis.'),
        ('Privacy/compliance documentation incomplete.','Provider secrets separated; owned data access checked; no payment-card storage.','Team: final privacy/terms, retention/deletion, processor and policy review. No legal-compliance certification claimed.'),
    ])
    src(d,'PRD 4.3; observed assessment and app-integration reports; current implementation. Action allocation is a proposed follow-up by responsibility area.')
    return d


def tests():
    d=Document();h(d,'8. System Test Summary')
    h(d,'8.1 Overview',2)
    p(d,'The former PTD reported 154 tests and disconnected claim/retrieval stages. The latest recorded integration snapshot reports 262 passing backend tests, a successful Android export and live authenticated Firestore checks. Claim analysis and retrieval are now connected. These results verify behavior and integration; they do not establish misinformation-detection accuracy.')
    h(d,'8.2 Scope and methodology',2)
    bullets(d,['Unit/contract tests cover input preparation, language checks, AI response validation, retrieval filtering/fallback, lexical assessment and orchestrator stopping rules.',
        'FastAPI tests verify authenticated access, owner-only history/detail, quota, retries, conflicting keys and controlled errors.',
        'Transaction/lease tests exercise duplicate completion, concurrent requests, stale leases, rollback, calendar resets and denial of client-selected Premium.',
        'Offline tests use mocked providers and in-memory/fake stores. Separate opt-in live smoke tests consume provider quota and use isolated integration accounts.',
        'Android export verifies bundling; emulator captures and API observations establish selected user flows. They are not physical-device or elderly-user acceptance tests.'])
    h(d,'8.3 Dated result summary',2)
    tbl(d,['Date / record','Observed result','Limit'],[
        ('6 Sep — Matthew claim report','207 backend tests at that point; 7/8 synthetic claim-stage examples matched; satire disagreed.','Not a held-out benchmark or factual-veracity accuracy.'),
        ('7 Sep — retrieval report','247 backend tests at that point; 45 retrieval tests. Three live text-pipeline requests completed with real providers.','Operational success includes an observed incorrect stance on a myth passage.'),
        ('7 Sep — app integration','262 backend tests passed; one existing Starlette/httpx deprecation warning. Android production export succeeded.','Dated recorded run; no new full test execution was needed for this documentation update.'),
        ('7 Sep — live Firestore report','9 recorded checks: replay, owned history/detail, other-user denial, quota, stored evidence and direct-read denial.','Focused checks with isolated accounts, not a complete security audit.'),
        ('7 Sep — Android captures','Premium Home, private history and saved result opened in emulator.','No paid subscription, physical-device matrix or older-user study proven.'),
    ])
    d.add_page_break();h(d,'8.4 Live integration observations',2)
    tbl(d,['Input / experiment','Observed output','Interpretation'],[
        ('Great Wall visible from the Moon','Needs Caution, score 67; Snopes and two NASA sources; recorded 6.36 s.','A NASA myth-description passage was wrongly treated as supporting. Operational pass; assessment error remains.'),
        ('Singapore independence on 9 August 1965','Low Concern, score 25; NUS/NLB sources; recorded 14.90 s in the pipeline check.','One factual example, not a general accuracy or latency measure.'),
        ('Chicken rice preference','Not Enough Information, null score; retrieval skipped; recorded 12.89 s.','Non-checkable stopping path observed.'),
        ('Authenticated Firestore smoke flow','Saved independent-Singapore claim with evidence; repeat returned same result; next free submission rejected with 429.','Checks ownership and charging; 11.72 s report duration covers the smoke sequence, not a single inference latency.'),
    ])
    h(d,'8.5 Evaluation still required',2)
    tbl(d,['Area','Evidence needed'],[
        ('Assessment accuracy','Team-labelled, held-out claims/passages; per-class precision/recall/F1, confusion matrix and false Low Concern rate.'),
        ('Retrieval quality','Relevant decisive evidence found, passage context retained, date/entity alignment and source independence.'),
        ('Performance/capacity','Repeated latency measurements including p50/p95, concurrent-load testing against 100 users / 100 requests per minute targets.'),
        ('Usability/accessibility','Task completion, comprehension of uncertainty, errors and navigation with older users or justified representatives.'),
        ('Security/privacy/recovery','Broader rules/IAM review, adversarial prompts, deployment transport, backup/recovery and retention/deletion validation.'),
        ('Image and release acceptance','Selected detector evaluation, OCR/context tests, physical devices and signed release package.'),
    ])
    h(d,'8.6 Proposed assessment-improvement experiment',2)
    p(d,'Compare the current lexical baseline with an evidence-grounded LLM or evaluated entailment model. Require per-source stance, an exact evidence quote and a short justification; validate the quote against the passage. Check entities, dates, amounts and eligibility, and retain a neutral/insufficient-evidence outcome. Deduplicate dependent sources before aggregation. This is a proposed experiment, not an implemented fix.')
    p(d,'Start with 60–100 balanced claims including myths, unrelated negation, dated announcements and missing evidence. Two teammates should label and reconcile disagreements. Separate development examples from held-out evaluation and report coverage as well as errors, so returning Not Enough Information for everything is not mistaken for good accuracy. No final numeric acceptance target has been approved.')
    src(d,'evaluation/reports/matthew_claim_evaluation.md; chu_retrieval_evaluation.md; text_pipeline_live_2026-09-07.json; app-integration/README.md and firestore-live.json. Existing synthetic samples are not independent ground truth.')
    return d


def communication():
    d=load('09_Project_Communication.docx')
    insert_note(d,'9.2 Project Website','No deployed UIABO project website or public operational portal has been supplied. The current deliverables are the local Android prototype, source repository and submission documents. A website URL, hosting record and screenshots remain pending if a website is required.')
    h(d,'9.3 Communication management plan',2)
    p(d,'Proposed additions for team review. Existing evidence supports GitHub, shared documents and meeting notes; the cadence and escalation rules below are not claimed as previously approved practices.')
    tbl(d,['Audience / purpose','Channel / artifact','Proposed trigger'],[
        ('Team implementation coordination','Repository changes, shared contracts, task list and blocker notes.','Each handoff or blocker; include owner, next action and evidence.'),
        ('Supervisor progress/scope review','Dated progress summary and demonstrable build.','Agreed review meetings and material scope changes.'),
        ('Assessor delivery','Submission PDFs, user manual and reproducible demonstration instructions.','Required submission milestones.'),
        ('User feedback','Planned usability notes and feedback/report workflows.','Each approved study or implemented feedback submission.'),
    ])
    h(d,'9.4 Conflict escalation and document revision',2)
    bullets(d,['Record the disputed requirement or technical decision, alternatives and impact in a shared note.',
        'The affected component owners and project leader review it against the URS/TDM and prototype goal.',
        'Escalate unresolved scope or academic-requirement questions to the supervisor; record the decision without inventing approval.',
        'For interface changes, update schema, examples, tests and dependent components together.',
        'For document revisions, retain the original, record the date/reason/source, refresh contents and review the exported PDF.'])
    src(d,'Existing Section 9, MeetingNotes and Sprint1/UIABO_SPRINT_1_TEAM_TASKS.md. The PUM and this PTD update provide dated documentation artifacts.')
    return d


def schedule():
    d=load('10_Project_Schedule.docx')
    insert_note(d,'10. Project Schedule','The Gantt images below are the original PRD planning baseline. They do not reflect a newly approved rebaseline or establish actual completion. The dated milestones and work breakdown after the figures show the evidence available for this update.')
    d.add_page_break();h(d,'10.1 Milestones and current progress',2)
    tbl(d,['Milestone','Date / state','Evidence or action'],[
        ('Initial backend/Firebase foundation','12–13 Aug 2026, recorded','Meeting brief records basic API tests and Firestore connectivity.'),
        ('Claim analysis integration','6 Sep 2026, recorded','Ollama live report and validation changes.'),
        ('Retrieval integration','7 Sep 2026, recorded','Google/Tavily live requests and pipeline results.'),
        ('Android and owned persistence','7 Sep 2026, recorded','Emulator, Firestore history/detail and quota checks.'),
        ('Preliminary documentation','7 Sep 2026, draft','PUM section files/PDF and this updated PTD.'),
        ('Prototype submission','12 Sep 2026, user-provided deadline','Basic Sprint 1 text demonstration; scope clarity, known limits and reproducible setup needed.'),
        ('Final submission','November 2026, exact date pending','Complete required image/deepfake scope and evaluation; confirm official deadline and rebaseline.'),
    ])
    h(d,'10.2 Work breakdown structure',2)
    tbl(d,['Work package','Deliverable / responsibility area','State'],[
        ('W1 Requirements/design','PRD, URS, TDM; team and documentation.','Existing baseline; scope differences recorded.'),
        ('W2 Input preparation','English validation and preparation; Yi Da stage.','Integrated.'),
        ('W3 Claim analysis','Classification/extraction; Matthew stage.','Integrated; additional evaluation needed.'),
        ('W4 Evidence retrieval','Provider discovery/search/citations; Chu stage.','Live completion exists in the current workspace; confirm contribution attribution.'),
        ('W5 Assessment','Evidence stance/risk/explanation; Poon stage.','Integrated baseline; semantic accuracy improvement outstanding.'),
        ('W6 App/API/persistence','Android flow, auth, results and quota; Donovan-led integration.','Basic text journey working.'),
        ('W7 Required image scope','OCR, image/caption/context, deepfake.','Not implemented; confirm named owners and dates.'),
        ('W8 Remaining product functions','Payments, deletion, reports, public sharing and portals.','Planned; prioritization and delivery dates needed.'),
        ('W9 Evaluation/release','Dataset, usability, load/security, packaging and documentation.','Partial tests/drafts; acceptance evidence incomplete.'),
    ])
    h(d,'10.3 Proposed next work order',2)
    p(d,'Before 12 September, prioritize reproducible text demonstration, clear unavailable-feature labels, meaningful accuracy examples and consistent documentation. For the remaining project, prioritize required image/deepfake implementation and assessment evaluation, then agreed operational/subscription functions and release validation. These priorities do not assign invented sprint dates or certify supervisor acceptance.')
    src(d,'PRD 4.5 Gantt; dated meeting/integration reports; user-provided September and November milestones.')
    return d


def roles():
    d=load('11_Roles_and_Responsibilities.docx')
    h(d,'11.1 Sprint 1 component ownership and contribution record',2)
    tbl(d,['Member','Agreed component handoff','Current evidence / attribution limit'],[
        ('Wong Yi Da','PreparedText: input preparation','Input validation and English gate are connected.'),
        ('Matthew Alexander Peeris','ClaimAnalysis: extraction/classification','User confirmed his work was ready; cloud integration/validation results are recorded.'),
        ('Chu Wai Chung','RetrievalResult: evidence retrieval','Original filtering/test hook retained; unfinished live provider work was completed in the current integration workspace. Do not equate ownership with sole authorship.'),
        ('Poon Chun Ping','AssessmentResult: evidence assessment','Lexical baseline integrated; observed stance errors remain.'),
        ('Ho Sze Wei, Donovan','Orchestration, API, app and persistence integration','Coordinates completed pipeline handoffs, Android results, Firestore ownership and allowances, and documentation updates.'),
    ])
    p(d,'The original role table records planned responsibilities. Before final submission, the team should reconcile actual contributions with repository history and reassigned tasks; no unverified percentage contribution or individual sign-off is included.')
    src(d,'Sprint 1 team-task document; existing PRD role table; dated integration reports and project discussion.')
    return d


def minutes():
    d=load('12_Meeting_Minutes.docx')
    insert_note(d,'12. Meeting Minutes','Historical meeting records are retained as recorded; they are not rewritten to reflect later progress. The earlier note treating deepfake as scope-sensitive is superseded for current planning by the user’s confirmation that deepfake is required and audio excluded. This later clarification is not claimed as a recorded supervisor approval at the August meeting.')
    h(d,'12.3 Later progress records and missing minute fields',2)
    p(d,'The September provider and app-integration reports are technical progress evidence, not formal meeting minutes. No meeting time, attendance, minute taker or approval has been invented from those reports. Complete these fields from actual team records before final submission.')
    return d


def references():
    d=load('13_References.docx')
    insert_note(d,'13. References','The existing PRD market-research references are retained with their original recorded access dates; that historical comparison was not re-audited during this documentation update. Technical references below were consulted for the current integration description.')
    h(d,'13.1 Implementation references',2)
    for title,url in [
        ('Expo. SDK 57 reference.','https://docs.expo.dev/versions/v57.0.0/'),
        ('Firebase. Verify ID tokens.','https://firebase.google.com/docs/auth/admin/verify-id-tokens'),
        ('Firebase. Transactions and batched writes.','https://firebase.google.com/docs/firestore/manage-data/transactions'),
        ('Ollama. Cloud.','https://docs.ollama.com/cloud'),
        ('Google. Fact Check Tools API: claims.search.','https://developers.google.com/fact-check/tools/api/reference/rest/v1alpha1/claims/search'),
        ('Tavily. Search API.','https://docs.tavily.com/documentation/api-reference/endpoint/search'),
        ('Tavily. Extract API.','https://docs.tavily.com/documentation/api-reference/endpoint/extract'),
    ]:p(d,f'{title} Accessed 7 September 2026. {url}')
    h(d,'13.2 Internal evidence and submission references',2)
    for value in ['PRD/FYP-26-S3-30_PRD.docx; URS/FYP-26-S3-30_URS.pdf; TDM/FYP-26-S3-30_TDM.pdf.',
        'MeetingNotes/team_meeting_brief_2026-08-13.md and Sprint1/UIABO_SPRINT_1_TEAM_TASKS.md.',
        'uiabo/evaluation/reports/matthew_claim_evaluation.md and matthew_claim_live_2026-09-06.json.',
        'uiabo/evaluation/reports/chu_retrieval_evaluation.md and text_pipeline_live_2026-09-07.json.',
        'uiabo/evaluation/reports/app-integration/README.md, firestore-live.json and Android captures.',
        'uiabo/evaluation/trusted_sources/README.md and pipeline component READMEs.',
        'Submission1/PUM/FYP-26-S3-30_PrelimUserManual_DRAFT.docx and PUM_MISSING_INFORMATION.md.']:
        p(d,value)
    return d


def data_collection():
    d=Document();h(d,'14. Data Collection and Evidence Management')
    p(d,'This chapter adds the data-collection coverage present in the comparison sample. UIABO currently retrieves evidence at request time and uses existing cloud models. It has not trained a misinformation model on a completed proprietary dataset, and the planned data-ingestion portal is not implemented.')
    h(d,'14.1 Data sources',2)
    tbl(d,['Source','Data / role','Limit'],[
        ('User submission','English text and original wording.','May contain personal information, false claims or instructions; treat as untrusted.'),
        ('Google Fact Check Tools','Previously reviewed claim metadata and review URLs.','A reviewed claim match is discovery, not sufficient evidence by itself.'),
        ('Tavily extraction/search','Page passages, titles, URLs and available dates from scoped web sources.','Pages may be unavailable, truncated or out of context.'),
        ('Runtime source catalogue','gov.sg, nasa.gov, snopes.com, fullfact.org, factcheck.org, reuters.com, apnews.com, channelnewsasia.com, nus.edu.sg.','Initial English allowlist; inclusion does not guarantee truth, independence or peer review.'),
        ('Team evaluation samples','Fixed handoff fixtures and synthetic claim/assessment examples.','Development tests; insufficient as independent evaluation or a training corpus.'),
    ])
    h(d,'14.2 Types of data and provenance',2)
    tbl(d,['Data group','Fields / treatment'],[
        ('Identity and access','Firebase UID, name/email, verification, account status and role; passwords remain in Firebase Authentication.'),
        ('Claim','Original/normalized input, extracted source span, category, checkability, classification reason and agreement.'),
        ('Evidence','Evidence ID, title, source URL, publisher/domain, optional author/date, passage, source type, relevance and retrieval timestamp.'),
        ('Assessment/result','Stance, quality, concern label, heuristic score or null, uncertainty, explanation, next action, warnings and pipeline version.'),
        ('Usage and recovery','Owner, result/request identifiers, fingerprint, timestamps, lease and applicable allowance period.'),
    ])
    h(d,'14.3 Collection methods',2)
    p(d,'For checkable content, search up to five English reviewed claims; accept catalogue URLs and extract up to three review pages. Preserve a labelled claim/verdict block when it agrees with provider metadata. If fewer than two distinct passages remain, run scoped advanced search requesting up to eight results. Validate and rank up to six final evidence items. Provider-generated answers are disabled in that search request.')
    p(d,'Google failure or unusable page extraction can trigger the search fallback. A successful empty search is distinguished from provider failure. Partial failure with usable evidence retains warnings. These are current bounded request-time operations, not a scheduled crawling or ingestion service.')
    h(d,'14.4 Preprocessing and quality controls',2)
    bullets(d,['Preserve original input; normalize Unicode composition and whitespace while preserving meaning-bearing wording and punctuation.',
        'Validate English, input length and usable content. Flag instruction-like content instead of deleting arbitrary claim text.',
        'Validate source URL scheme/domain and citation fields; retain meaningful URL query identifiers while removing tracking fragments/parameters.',
        'Remove duplicate URLs/content. This reduces duplication but does not establish independent corroboration.',
        'Keep unknown publication dates null; do not substitute event dates or claim dates.',
        'Retain evidence snapshots with the result so later source changes do not silently rewrite the original explanation.'])
    p(d,'Current relevance thresholds use lexical coverage and provider relevance. These scores are not truth probabilities. Related passages may omit a decisive condition, and older sources may remain useful for historical claims; recency must be interpreted in context.')
    h(d,'14.5 Privacy, access and data lifecycle',2)
    p(d,'Analysis text is sent to the configured cloud claim service, and the extracted claim is used in external evidence searches. The system therefore has external processing dependencies. The final user-facing privacy notice, retention schedule, deletion workflow and review of provider terms/processing arrangements are outstanding. Public accessibility of a source does not establish unrestricted reuse rights.')
    p(d,'Current safeguards include backend-only provider secrets, verified/active-account checks, owned result access and no payment-card data collection. Do not present these focused measures as a complete privacy or legal-compliance assessment. Demonstrations should use dedicated accounts and non-sensitive example text.')
    h(d,'14.6 Evaluation dataset plan',2)
    p(d,'Build a versioned claim-and-evidence dataset with the original claim, date/context, expected verdict, decisive cited passages and annotation rationale. Record provenance and usage conditions. Label/reconcile independently, separate development from held-out evaluation, and retain failure cases. The proposed 60–100-case starting set in Section 8.6 is not a completed dataset or final sample-size justification.')
    src(d,'Current retrieval/input implementations, shared models, source catalogue and saved integration reports; provider documentation in Section 13.')
    return d


TRACE=[
    ('PF-B01; UC-01/02/04–06','Individual account access/profile','AuthScreen, VerifyEmailScreen, ProfileScreen; account/auth tests.','Available subset; operational equivalents remain planned.'),
    ('PF-B02; UC-03/16/18','Registration and account deletion','Free registration/verification exists.','Deletion and paid registration not implemented.'),
    ('PF-B03; UC-07','English text analysis','All text stages, POST /analysis/text, emulator flow.','Operationally demonstrated; accuracy not established.'),
    ('PF-B04; UC-08','Direct public-link analysis','TDM link screen; copied-text alternative in app.','Direct fetching/screening pending.'),
    ('PF-B05; UC-09–11/13','History, details, allowance and reports','Owned GET routes, quota tests, live Firestore.','History/detail/allowance available; reporting pending.'),
    ('PF-B06/07; UC-17/19/20','Paid upgrade, renewal, cancellation','TDM designs and Premium server role.','Payment subscription flow pending.'),
    ('PF-B08/09; UC-21/23','Image+caption and OCR','TDM designs and shared intended results.','Pipelines/format limits/evaluation pending.'),
    ('PF-B10/11; UC-22/24','Context and deepfake/AI image','Required static-image scope in URS/TDM.','Detector and context pipeline pending.'),
    ('PF-B12–14; UC-30–34','Operational account management','URS cases and TDM portal designs.','Role-restricted portal pending.'),
    ('PF-B15/16; UC-41–43','Reports and ingestion monitoring','URS cases and TDM portal designs.','Dashboard, investigation and ingestion workflows pending.'),
    ('PF-A01/02/04; UC-12/14/35','Delete history and review feedback','URS additional requirements.','Not implemented.'),
    ('PF-A03; UC-15','Public shareable result link','Native summary/citation sharing exists.','Public result-page links pending.'),
    ('URS 4.1–4.4','Capacity, UI speed, storage and result-time targets','Selected live timings and bundle success.','No load/physical-device/storage-warning certification.'),
    ('URS 4.5 / 2.5','Access, privacy and security','Focused auth/ownership/idempotency tests.','Inactivity expiry, broader security and compliance review pending.'),
]


def requirements():
    d=Document();h(d,'15. Requirement Definition and Traceability')
    h(d,'15.1 Stakeholder identification and evidence gathering',2)
    p(d,'Requirements evidence consists of the PRD’s documented product research, the URS user stories/use cases, TDM diagrams and screens, available meeting notes, and subsequent user scope clarifications. The primary audience is older Singapore Android users; operational roles and academic reviewers are identified in Section 2.10. No completed user interview, survey or focus-group study is claimed without its records.')
    h(d,'15.2 Requirement analysis and priority',2)
    p(d,'The URS distinguishes base requirements PF-B01–B16 and additional requirements PF-A01–A04. The team’s Sprint 1 goal is the basic text pipeline. The user clarified a September Android-emulator prototype and confirmed deepfake as required while excluding audio. These statements narrow the immediate increment; they do not silently remove required final-project features.')
    p(d,'The analysis approach maps each user action to a use case, component boundary, data record and verification artifact. Feasibility is judged against evidence quality, external dependencies, UI effort and schedule. Changes affecting scope or shared field names should be recorded and reviewed with affected owners; unresolved academic requirements need supervisor clarification.')
    h(d,'15.3 Functional hierarchy and access',2)
    tbl(d,['Module','Functions / roles'],[
        ('Individual account management','Free/premium registration/access, verification, profile, allowance; deletion and billing planned.'),
        ('Content analysis','Text now; public-link and premium image/OCR/context/deepfake planned.'),
        ('Results and evidence','Owned history/detail, citations, uncertainty and native share now; deletion/public links/reports planned.'),
        ('Administration','Planned authorised account queries, creation, suspension and feedback review.'),
        ('Data engineering','Planned incorrect-result investigation and ingestion monitoring.'),
    ])
    h(d,'15.4 Requirement-to-evidence matrix',2)
    tbl(d,['Requirement','Purpose','Implementation / evidence','Status / gap'],TRACE)
    h(d,'15.5 Acceptance and change control',2)
    p(d,'A feature is operationally complete only when its intended user flow, authorization, data persistence and failure behavior are demonstrated. Quality claims require their own evaluation. The 262-test record cannot substitute for labelled fact-check accuracy or older-user usability. Acceptance thresholds and final sign-off remain to be agreed and documented.')
    src(d,'URS Sections 2.2, 3 and 4; existing Section 5 facsimiles contain detailed use cases; Section 8 supplies dated verification records.')
    return d


def charter():
    d=Document();h(d,'16. Project Charter, Scope and Delivery Controls')
    p(d,'This charter summary is a proposed consolidation of existing requirements and schedule information. It is not a signed authorization or newly approved scope baseline.')
    tbl(d,['Charter field','Draft record'],[
        ('Project','uiabo — AI-Based Misinformation Detection for Short-Form Content; FYP-26-S3-30.'),
        ('Purpose','Help older Singapore Android users review suspicious short-form content using evidence and clear uncertainty.'),
        ('Team / oversight','Five members listed on the cover; supervisor Mr Liaw Chun Huei; assessor Mr Terrence Chew.'),
        ('Immediate milestone','12 September 2026 prototype, as stated by the user; basic Sprint 1 text demonstration.'),
        ('Final milestone','November 2026; exact official deadline not supplied.'),
        ('Resources','Existing development machines, Firebase project and provider accounts. No agreed funded budget or committed provider spend documented.'),
        ('Constraints','English Android scope, external service access/quotas, small evaluation base and limited remaining development time.'),
        ('Decision / approval','Team review and supervisor confirmation of unresolved acceptance/scope items required; no signatures inferred.'),
    ])
    h(d,'16.1 Scope statement',2)
    tbl(d,['Category','Scope'],[
        ('Implemented increment','Individual authentication, English text pipeline, cited result, Android Home/text/history/detail and server-enforced allowances.'),
        ('Required/planned product','Public links with screening; image+caption, OCR, context and deepfake; agreed accounts/subscriptions and operational functions.'),
        ('Additional features','History deletion, reviews, public result links and operational feedback as identified by the URS.'),
        ('Excluded','Audio/speech, video and non-English analysis; private/paywalled inaccessible content. No iOS release or public API licensing is part of the current working increment.'),
    ])
    h(d,'16.2 Deliverables and dependencies',2)
    bullets(d,['Product: Android application, authenticated API, Firebase configuration, evidence pipeline, source code and reproducible setup.',
        'Documentation: PRD, URS, TDM, this PTD, PUM, source/evaluation records and actual meeting minutes.',
        'Verification: automated results, labelled evaluation, usability/performance/security records and final demonstration evidence.',
        'Dependencies: provider availability, appropriate credentials, source access, detector selection and approved final deadlines.'])
    h(d,'16.3 Acceptance and unresolved approval',2)
    p(d,'The local text demonstration has evidence of completion, but the full planned product is incomplete. Final acceptance must explicitly address accuracy limits, required image/deepfake delivery, unfinished user/operational flows and non-functional targets. Confirm whether a public project website is an assessed deliverable. Approval dates, signatures and exact final submission dates remain pending rather than fabricated.')
    src(d,'PRD scope/business model/schedule; URS; team-task document; user-provided prototype scope and dates.')
    return d


def glossary():
    d=Document();h(d,'17. Glossary')
    tbl(d,['Term','Meaning in uiabo'],[
        ('API','The interface used by the Android client to request backend operations.'),('Authentication / authorization','Establishing identity / deciding what that identity is allowed to access.'),
        ('Firebase UID / ID token','Account identifier / signed client identity token verified by the backend.'),('Firestore','Document database used for account, allowance and owned result records.'),
        ('Claim','The factual assertion selected from submitted content for evidence checking.'),('ClaimReview','Published fact-check metadata used to discover existing reviews.'),
        ('Evidence stance','Whether a passage supports, contradicts or is neutral toward the specific claim.'),('Provenance','Where evidence came from and the source/time/context needed to inspect it.'),
        ('Not Enough Information','No suitable checkable claim or decisive evidence; not equivalent to false.'),('Risk score','Current heuristic concern indicator, not a calibrated probability of falsehood.'),
        ('Uncertainty','A statement of limitations, weak evidence or disagreement in an assessment.'),('OCR','Extraction of visible text from an image, with user correction planned.'),
        ('Deepfake / AI-image analysis','Planned static-image generation/manipulation likelihood assessment, separate from factual truth.'),('Idempotency','Repeating the same successful request returns the same result without another charge.'),
        ('Transaction / lease','Atomic group of storage changes / temporary per-user reservation preventing overlapping work.'),('Ground truth / held-out set','Reviewed reference labels / evaluation examples excluded from development tuning.'),
        ('WBS','Work breakdown structure: deliverables divided into manageable work packages.'),('Traceability','Mapping a requirement to design, implementation and verification evidence.'),
    ])
    return d


def conclusion():
    d=Document();h(d,'18. Conclusion and Remaining Deliverables')
    p(d,'UIABO has progressed from a mock API and disconnected components to a working authenticated text-check journey in an Android emulator. Claim analysis, evidence retrieval, baseline assessment, cited results, owned Firestore history and server-enforced free/premium allowances are connected. Dated tests support those operational claims.')
    p(d,'The most important remaining quality gap is evidence interpretation: the current lexical stance component can misread myth quotations and unrelated negation. Required static-image/deepfake work, direct links, payments and operational portals remain unfinished. Documentation now describes these as planned, and it does not equate integration with validated factual accuracy.')
    h(d,'18.1 Remaining information and acceptance evidence',2)
    tbl(d,['Item','Required next evidence / decision'],[
        ('Assessment evaluation','Labelled dataset, independent review, held-out metrics and agreed acceptance criteria.'),
        ('Image/deepfake delivery','Named owners, detector/services, input limits, implementation and evaluation.'),
        ('Remaining user/operational functions','Link screening, subscriptions, deletion, reports, public result links and portals according to agreed priority.'),
        ('Deployment/privacy','Release package, public addresses if required, configuration/rules/IAM, privacy/terms, retention/deletion and recovery procedures.'),
        ('Non-functional validation','Load, response-time distribution, physical-device testing, inactivity rules, accessibility and older-user comprehension.'),
        ('Management records','Official November deadline, rebaselined Gantt/WBS dates, approved learning objectives/charter, complete meeting minutes and contribution attribution.'),
    ])
    h(d,'18.2 Appendix and evidence locator',2)
    p(d,'Detailed use cases, diagrams and proposed database/UI designs are retained in Section 5. Dated test records and their limits are summarized in Section 8 and located in Section 13.2. PTD_UPDATE_SUMMARY.md maps the comparison sample to the revised chapters. PTD_MISSING_INFORMATION.md is the living checklist. The original PDF and prior checklist are preserved in backups/.')
    return d


def main():
    OUT.mkdir(exist_ok=True)
    builders=[('00_Cover_and_Contents.docx',cover),('01_Product_Functionality_Overview.docx',functionality),
        ('02_Project_Overview.docx',overview),('03_Similar_Existing_Products.docx',lambda:load('03_Similar_Existing_Products.docx')),
        ('04_Product_Comparison.docx',lambda:load('04_Product_Comparison.docx')),('05_Project_Design.docx',design),
        ('06_Project_Implementation.docx',implementation),('07_Risk_Analysis.docx',risks),('08_System_Test_Summary.docx',tests),
        ('09_Project_Communication.docx',communication),('10_Project_Schedule.docx',schedule),('11_Roles_and_Responsibilities.docx',roles),
        ('12_Meeting_Minutes.docx',minutes),('13_References.docx',references),('14_Data_Collection.docx',data_collection),
        ('15_Requirements_Traceability.docx',requirements),('16_Project_Charter_and_Scope.docx',charter),
        ('17_Glossary.docx',glossary),('18_Conclusion_and_Remaining_Deliverables.docx',conclusion)]
    for name,builder in builders:
        d=style(builder());d.save(OUT/name);print('Updated',name)
    master=Document(OUT/builders[0][0]);composer=Composer(master)
    for name,_ in builders[1:]:
        master.add_page_break();composer.append(Document(OUT/name))
    style(master)
    path=ROOT/'FYP-26-S3-30_Preliminary_Technical_Documentation.docx'
    composer.save(path)
    print('Combined',path.name,'images',len(master.inline_shapes))


if __name__=='__main__':main()
